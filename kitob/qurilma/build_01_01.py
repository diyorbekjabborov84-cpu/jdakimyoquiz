#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "chiqish" / "1.1 Kislota tushunchasining kelib chiqishi.docx"
IMG_TIMELINE = ROOT / "rasmlar" / "01_timeline.png"
IMG_ACIDUS = ROOT / "rasmlar" / "01_acidus.png"
OUT.parent.mkdir(parents=True, exist_ok=True)

BLACK = RGBColor(0x1A,0x1A,0x1A)
NAVY = RGBColor(0x1A,0x2E,0x4A)
GRAY = RGBColor(0x4A,0x4A,0x4A)
LIGHT = RGBColor(0x6B,0x6B,0x6B)
BRONZE = RGBColor(0x8C,0x6A,0x2A)

def set_margins(sec):
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.5); sec.right_margin = Cm(2)
    sec.header_distance = Cm(1); sec.footer_distance = Cm(1)

def p_add(doc, text, size=10, bold=False, italic=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, indent=0, first_indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.15
    if indent: pf.left_indent = Cm(indent)
    if first_indent is not None: pf.first_line_indent = Cm(first_indent)
    pf.widow_control = True
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = "Times New Roman"
    rPr = r._r.get_or_add_rPr()
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rf.set(qn('w:eastAsia'), 'Times New Roman'); rf.set(qn('w:cs'), 'Times New Roman'); rPr.append(rf)
    # kerning
    return p

def p_mixed(doc, segments, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, indent=0, first_indent=None):
    # segments: list of (text, {size,bold,italic,color})
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.15
    if indent: pf.left_indent = Cm(indent)
    if first_indent is not None: pf.first_line_indent = Cm(first_indent)
    for text, st in segments:
        r = p.add_run(text)
        r.font.size = Pt(st.get('size',10)); r.bold = st.get('bold',False); r.italic = st.get('italic',False); r.font.color.rgb = st.get('color',BLACK); r.font.name = "Times New Roman"
        rPr = r._r.get_or_add_rPr()
        rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
    return p

def add_heading(doc, num, title, level=1):
    if level==1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format; pf.space_before = Pt(14); pf.space_after = Pt(4)
        r = p.add_run(num); r.font.size = Pt(11); r.bold = True; r.font.color.rgb = NAVY; r.font.name = "Times New Roman"
        r2 = p.add_run("  " + title); r2.font.size = Pt(11); r2.bold = True; r2.font.color.rgb = NAVY; r2.font.name = "Times New Roman"
        for rr in p.runs:
            rPr = rr._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
        # line
        p2 = doc.add_paragraph(); p2Pr = p2._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:color'), '8C6A2A'); b.set(qn('w:space'), '4'); pBdr.append(b); p2Pr.append(pBdr)
        pf2 = p2.paragraph_format; pf2.space_before = Pt(1); pf2.space_after = Pt(8)
        return p
    else:
        p = doc.add_paragraph()
        pf = p.paragraph_format; pf.space_before = Pt(10); pf.space_after = Pt(4)
        r = p.add_run(num + "  " + title); r.font.size = Pt(10); r.bold = True; r.font.color.rgb = BLACK; r.font.name = "Times New Roman"
        rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
        return p

def add_image(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format; pf.space_before = Pt(8); pf.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15))
    # caption
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf2 = p2.paragraph_format; pf2.space_before = Pt(2); pf2.space_after = Pt(8)
    r = p2.add_run(caption); r.font.size = Pt(8); r.italic = True; r.font.color.rgb = GRAY; r.font.name = "Times New Roman"
    rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
    return p

def add_table(doc, headers, rows, caption=None):
    if caption:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.space_before = Pt(8); pf.space_after = Pt(4)
        r = p.add_run(caption); r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = NAVY; r.font.name = "Times New Roman"
        rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER if False else 1
    table.autofit = True
    # header
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); r.font.name = "Times New Roman"
        shading = OxmlElement('w:shd'); shading.set(qn('w:val'), 'clear'); shading.set(qn('w:fill'), '1A2E4A'); c._tc.get_or_add_tcPr().append(shading)
        # padding
        tcPr = c._tc.get_or_add_tcPr()
        for m in ['top','bottom','left','right']:
            mar = OxmlElement(f'w:tcMar')
            # Actually use w:tcMar with w:top etc inside — simpler: use w:shd already
            pass
        c.vertical_alignment = docx.enum.table.WD_ALIGN_VERTICAL.CENTER if False else 1
        rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = table.rows[ri+1].cells[ci]
            c.text = ""
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci>0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val)); r.font.size = Pt(8.5); r.font.color.rgb = BLACK; r.font.name = "Times New Roman"
            if ci==0: r.bold = True
            # zebra
            bg = "E8EEF5" if ri%2==0 else "FFFFFF"
            shading = OxmlElement('w:shd'); shading.set(qn('w:val'), 'clear'); shading.set(qn('w:fill'), bg); c._tc.get_or_add_tcPr().append(shading)
            rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)
    # borders
    tbl = table._tbl; tblPr = tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '1A2E4A'); el.set(qn('w:space'), '0'); tblBorders.append(el)
    tblPr.append(tblBorders)
    # spacing after
    p = doc.add_paragraph(); pf = p.paragraph_format; pf.space_after = Pt(4); pf.space_before = Pt(2)
    return table

import docx

doc = Document()
sec = doc.sections[0]
set_margins(sec)
sec.page_height = Cm(29.7); sec.page_width = Cm(21)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(10); st.font.color.rgb = BLACK
st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.15

# Sarlavha bloki
p_add = p_add
p_add(doc, "NOORGANIK KISLOTALAR", 9, False, False, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 0, 2)
p_add(doc, "I QISM. NAZARIY POYDEVOR — 1-BOB", 7, True, False, BRONZE, WD_ALIGN_PARAGRAPH.CENTER, 0, 4)
p_add(doc, "1.1. Kislota tushunchasining kelib chiqishi", 16, True, False, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 0, 4)
p_add(doc, "«Acidus» so'zidan zamonaviy ta'rifgacha — kislotalar qanday kashf etilgan?", 9, False, True, GRAY, WD_ALIGN_PARAGRAPH.CENTER, 0, 10)
# chiziq
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:color'), '8C6A2A'); b.set(qn('w:space'), '4'); pBdr.append(b); pPr.append(pBdr)
pf = p.paragraph_format; pf.space_before = Pt(2); pf.space_after = Pt(10)

# Kirish
p_add(doc, "Kirish", 11, True, False, NAVY, WD_ALIGN_PARAGRAPH.LEFT, 8, 4)
p_add(doc, "Biz kundalik hayotda «kislota» so'zini ko'p eshitamiz: limon nordon, sirka nordon, oshqozonda xlorid kislota pepsin faoliyati uchun kislotali muhit yaratadi, qo‘rg‘oshinli akkumulyatorda sulfat kislota elektrolit vazifasini bajaradi. Lekin «kislota nima?» degan savolga kimyogarlar 300 yil davomida bir xil javob bermagan. Avvaliga «nordon ta'mli narsa», keyin «tarkibida kislorod bor modda», so'ngra «suvda vodorod ioni beradigan modda» deyilgan. Hozir esa bu tushuncha ancha kengaygan.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, first_indent=1)
p_add(doc, "Ushbu mavzuda biz kislota tushunchasi qanday paydo bo'lganini, qaysi moddalar birinchi bo'lib kislota deb tan olinganini va bu so'zning o'zi qaerdan kelganini ko'rib chiqamiz. Maqsad — keyingi mavzularda Arrhenius, Brønsted–Lowry va Lewis nazariyalarini o'rganish uchun tarixiy zamin yaratish.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 8, first_indent=1)

add_heading(doc, "1.1.1.", "So'zning kelib chiqishi: lotincha acidus", level=2)
p_mixed(doc, [
    ("Lotin tilida ", {"size":10, "color":BLACK}),
    ("acidus", {"size":10, "italic":True, "color":BLACK}),
    (" — «nordon», «o'tkir», «achchiq» degan ma'noni bildiradi. Shu so'zdan inglizcha ", {"size":10, "color":BLACK}),
    ("acid", {"size":10, "italic":True}),
    (", fransuzcha ", {"size":10}),
    ("acide", {"size":10, "italic":True}),
    (" kabi atamalar shakllangan. Ruscha ", {"size":10}),
    ("кислота", {"size":10, "italic":True}),
    (" esa slavyan tillaridagi «nordon» ma’noli ildiz bilan bog‘liq; o‘zbekcha «kislota» rus tili orqali o‘zlashgan. Qadimgi rimliklar sirka va nordon sharobni ", {"size":10}),
    ("acetum", {"size":10, "italic":True}),
    (" deb atagan, bu so'z ham ", {"size":10}),
    ("acere", {"size":10, "italic":True}),
    (" — «nordon bo'lmoq» fe'lidan olingan.", {"size":10}),
], before=0, after=6, first_indent=1)
p_add(doc, "E'tibor bering: dastlab «kislota» kimyoviy ta'rif emas, balki ta'm bilish a'zosi orqali berilgan tavsif bo'lgan. Nordon ta'm — kislotalarning eng birinchi belgisi. Keyinchalik bu belgi yetarli emasligi ma'lum bo'ldi: masalan, oshqozon xlorid kislotasi (HCl) ham nordon, lekin uni tatib ko'rib bo'lmaydi — juda xavfli.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, first_indent=1)

p_add(doc, "Nordon ta’m kimyoviy ta’rif o‘rnini bosa olmaydi. Laboratoriya moddalarini tatib ko‘rish qat’iyan mumkin emas. Kislota xossalari indikatorlar, elektr o‘tkazuvchanlik, titrlash va boshqa o‘lchash usullari orqali tekshiriladi.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 8, first_indent=1)

# Rasm 1 — Acidus
add_table(doc, ["Kuzatilgan belgi", "Ilmiy savol"], [["Nordon ta’m", "Qaysi zarrachalar nordonlikni yuzaga keltiradi?"], ["Indikator rangining o‘zgarishi", "Rang eritma muhitiga qanday bog‘liq?"], ["Asos bilan neytrallanish", "Reaksiyada qanday zarrachalar almashinadi?"]], caption="1.1-sxema. Kundalik kuzatishdan ilmiy tushunchaga o‘tish. Moddalarni tatib aniqlash mumkin emas.")

add_heading(doc, "1.1.2.", "Qadimdan ma’lum kislotalar va mineral kislotalarning olinishi", level=2)
p_add(doc, "Kislotalar laboratoriyada ixtiro qilingan emas — tabiatda topilgan. Tarixda muhim o‘rin tutgan misollar quyidagilar; ularni insoniyat bilgan aynan birinchi uchta kislota deb hisoblash to‘g‘ri emas:", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.LEFT, 0, 6)

add_table(doc,
    ["Nomi (tarixiy)", "Tarkibi", "Qaerdan olingan", "Qachon ta'riflangan"],
    [
        ["Sirka kislotasi", "CH₃COOH", "Nordon vinodan, uzum sharobini achitib", "Qadimgi davrlardan ma’lum"],
        ["Vitriol moyi", "H₂SO₄", "Temir kuporosi (FeSO₄) ni qizdirib", "O‘rta asr vitriol distillash an’analari"],
        ["Kuchli suv (aqua fortis)", "HNO₃", "Selitrani kislotali sulfatlar bilan qayta ishlab", "O‘rta asr alkimyo manbalari"],
    ],
    caption="1.1-jadval. Tarixiy nomlar va ularning zamonaviy kimyoviy talqini."
)
p_add(doc, "Sirka etanolning mikrobiologik oksidlanishi natijasida hosil bo‘ladi; uning olinishi uchala misol ham distillashga asoslangan degani emas. Mineral kislotalarni ajratishda esa distillash muhim rol o‘ynagan. Vitriollarni qizdirish va selitrani qayta ishlash haqidagi tarixiy retseptlarning tarkibi turlicha bo‘lgan. Shu sababli mineral kislotalarning kashf etilishini aniq bir sana yoki yagona shaxsga qat’iy bog‘lash ehtiyotkorlikni talab qiladi.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 8, first_indent=1)

add_heading(doc, "1.1.3.", "Alkimyo davrida kislotalar qanday tushunilgan?", level=2)
p_add(doc, "Alkimyo manbalarida mineral kislotalar ko‘pincha «o‘tkir suvlar» kabi nomlar bilan tavsiflangan; zamonaviy yagona tasnif hali shakllanmagan edi. Xossalarni izohlashda turli sifat tavsiflari qo‘llangan, ammo ularni bugungi kislota kuchi yoki konsentratsiyasi bilan tenglashtirish mumkin emas.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, first_indent=1)
p_add(doc, "Ular kislotalarning ikkita muhim xossasini topgan:", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.LEFT, 0, 4)
for txt in [
    "1) Ayrim metallar suyultirilgan, oksidlovchi bo‘lmagan kislotalar bilan vodorod ajratadi. Bu barcha kislota–metall juftlari uchun umumiy qoida emas.",
    "2) Kislotalar ishqorlar bilan qo'shilsa, neytrallanadi — «o'tkirlik» yo'qoladi va tuz hosil bo'ladi."
]:
    p_add(doc, txt, 9, False, False, BLACK, WD_ALIGN_PARAGRAPH.LEFT, 0, 2, indent=0.8)
p_add(doc, "Bu kuzatishlar to'g'ri edi, lekin nazariya yo'q edi. Alkimyogarlar «nima uchun?» degan savolga javob bera olmagan — faqat «qanday?»ni tasvirlagan.", 10, False, False, BLACK, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 8, first_indent=1)

add_heading(doc, "1.1.4.", "Boyl va Lavuazyedan keyingi burilish", level=2)
p_mixed(doc, [
    ("XVII asrda Robert Boyl kislotalarni aniqlashning birinchi obyektiv usulini taklif qildi: ", {"size":10}),
    ("lakmus qog'ozi.", {"size":10, "italic":True}),
    (" Lakmusning qizarishi eritma kislotali ekanini ko‘rsatadi, lekin aynan qaysi kislota mavjudligini aniqlamaydi. Bu ta'm bilishdan ko'ra xavfsiz va aniq edi.", {"size":10}),
], before=0, after=6, first_indent=1)
p_mixed(doc, [
    ("XVIII asr oxirida Antuan Lavuazye barcha kislotalar tarkibida kislorod bor deb xato qildi va kislorodga ", {"size":10}),
    ("oxygen", {"size":10, "italic":True}),
    (" — «kislota hosil qiluvchi» deb nom berdi (yun. ", {"size":10}),
    ("oxys", {"size":10, "italic":True}),
    (" — kislota, ", {"size":10}),
    ("gennan", {"size":10, "italic":True}),
    (" — tug'moq). Keyinroq Gemfri Devi vodorod xlorid (HCl) da kislorod yo'qligini isbotlab, bu nazariyani rad etdi. Shundan keyin olimlar kislotalikning haqiqiy sababini vodoroddan izlay boshladi — bu yo'l Arrheniusgacha olib bordi.", {"size":10}),
], before=0, after=8, first_indent=1)

# Rasm 2 — Timeline
add_image(doc, ROOT / "rasmlar" / "01_timeline_uz.png", "1.1-rasm. Kislota tushunchasining rivojlanish ketma-ketligi.")
add_table(doc, ["Davr", "Asosiy qarash", "Chegarasi yoki yangiligi"], [["XVII asr — Robert Boyl", "Kislotalarni kuzatiladigan xossalar bilan tavsiflash", "Indikatorlar, neytrallanish; zarracha modeli yo‘q"], ["XVIII asr oxiri — Lavuazye", "Kislorodni kislotalilik manbai deb hisoblash", "HCl kabi kislorodsiz kislotalarni tushuntirmaydi"], ["XIX asr boshi — Devi", "Kislorodsiz kislotalarni tan olish", "Vodorodning rolini o‘rganishga yo‘l ochildi"], ["1884 — Arrhenius", "Suvda vodorod ionlari hosil bo‘lishi", "Suvli eritmalar doirasidagi model"], ["1923 — Brønsted–Lowry", "Proton berish va qabul qilish", "Suvsiz muhitlarga ham tatbiq qilinadi"], ["1923 — Lewis", "Elektron juftini qabul qilish va berish", "Protonsiz kislota–asos reaksiyalarini ham qamraydi"]], caption="1.2-sxema. Kislota tushunchasining rivojlanishi. Bosqichlar teng vaqt oralig‘ida joylashmagan.")

add_heading(doc, "", "Xulosa", level=1)
bullets = [
    "Inglizcha acid lotincha acidus bilan bog‘liq; o‘zbekcha kislota atamasi esa rus tili orqali o‘zlashgan.",
    "Tarixiy misollar: sirka kislotasi (CH₃COOH), sulfat kislota (H₂SO₄) va nitrat kislota (HNO₃).",
    "Alkimyogarlar kislotalarning metallar va ishqorlar bilan reaksiyasini topgan, ammo nazariya yarata olmagan.",
    "Boyl lakmus sinovini, Lavuazye kislorod nazariyasini taklif qildi — ikkinchisi xato bo'lib chiqdi va vodorodning rolini o'rganishga turtki berdi.",
    "Tarix shuni ko'rsatadi: bitta belgi (ta'm, rang, tarkib) yetarli emas — kislotalarni tushunish uchun bir nechta nazariya kerak. Shu nazariyalarni keyingi mavzularda o'rganamiz."
]
for b in bullets:
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.3)
    r = p.add_run("•  " + b); r.font.size = Pt(9); r.font.color.rgb = BLACK; r.font.name = "Times New Roman"
    rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman'); rPr.append(rf)

p_add(doc, "Tayanch so'zlar: acidus, atsetum, vitriol moyi, aqua fortis, lakmus, oxygen.", 8, False, True, GRAY, WD_ALIGN_PARAGRAPH.LEFT, 8, 4)
# Savollar
p_add(doc, "Mustaqil ishlash uchun", 10, True, False, NAVY, WD_ALIGN_PARAGRAPH.LEFT, 8, 4)
qs = [
    "1. Nima uchun dastlab kislotalar ta'mi bo'yicha ta'riflangan? Bu usulning ikkita kamchiligini ayting.",
    "2. Sirka, H₂SO₄ va HNO₃ ning tarixiy nomlarini zamonaviy formulalari bilan yozing.",
    "3. Lavuazye nazariyasi nima uchun xato bo'lib chiqdi? HCl misolida tushuntiring.",
    "4. Boyl lakmus sinovi nima uchun ta'm bilishdan ko'ra ilmiyroq edi?",
]
for q in qs:
    p_add(doc, q, 9, False, False, BLACK, WD_ALIGN_PARAGRAPH.LEFT, 0, 2, indent=0.5)

# footer/header
for sec in doc.sections:
    hdr = sec.header; hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NOORGANIK KISLOTALAR  •  1-BOB  •  1.1-MAVZU"); r.font.size = Pt(7); r.font.color.rgb = LIGHT; r.font.name = "Times New Roman"; r.italic = True
    rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rPr.append(rf)
    ftr = sec.footer; ftr.is_linked_to_previous = False
    p = ftr.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); t = OxmlElement('w:top'); t.set(qn('w:val'), 'single'); t.set(qn('w:sz'), '4'); t.set(qn('w:color'), 'D9D9D9'); t.set(qn('w:space'), '4'); pBdr.append(t); pPr.append(pBdr)
    r = p.add_run("Noorganik kislotalar  •  1.1"); r.font.size = Pt(7); r.font.color.rgb = LIGHT; r.font.name = "Times New Roman"
    rPr = r._r.get_or_add_rPr(); rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), 'Times New Roman'); rPr.append(rf)

doc.save(str(OUT))
print(f"Saved: {OUT} ({OUT.stat().st_size/1024:.0f} KB)")
