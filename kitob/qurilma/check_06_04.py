#!/usr/bin/env python3
"""6.4 moduli uchun QA: model hisoblarining matn va hujjat bilan mosligi."""
from pathlib import Path
import csv, math, re, json
from docx import Document
import prep_06_04 as M

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'chiqish' / 'Noorganik-kislotalar-6.4-HBr-bromid-kislota.docx'
ASSETS = ROOT / 'rasmlar' / 'hbr'
MS = ROOT / 'matn' / '06_04.md'
fails = []


def ck(cond, msg):
    if cond:
        print('  OK  ', msg)
    else:
        print('  FAIL', msg); fails.append(msg)


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol


print('1) Model: spetsiatsiya va tribromid')
s = M.speciation_series()
ck(close(s[0]['ph'], 1.000, 1e-3) and close(s[0]['w_br'], 100.0, 0.001),
   '0,100 M: Br⁻ 100,000 %; pH 1,000')
ck(close(s[7]['ph'], -0.943, 1e-3) and close(s[7]['w_br'], 100.0, 0.001),
   '8,770 M: Br⁻ 100,000 %; pH −0,943')
ck(close(s[9]['ph'], -1.079, 1e-3) and close(s[9]['w_hbr2'], 0.0001, 1e-4),
   '12,000 M: pH −1,079; HBr₂⁻ 0,0001 %')
ck(close(M.KA_HBR, 6.3e8, 1e6) and close(M.K_HBR2, 10 ** 0.7, 1e-6),
   'Ka = 6,3×10⁸; lgK(HBr₂⁻) = 0,7')
tb = M.tribromide_series()
ck(close(tb[0]['w_br3'], 60.40, 0.05), '0,10 M Br⁻: Br₃⁻ ulushi 60,40 %')
ck(close(tb[6]['w_br3'], 99.29, 0.05), '8,77 M Br⁻: Br₃⁻ ulushi 99,29 %')
ck(close(tb[8]['w_br3'], 99.48, 0.05), '12,0 M Br⁻: Br₃⁻ ulushi 99,48 %')

print('2) Model: ishlab chiqarish, redoks, metallar, chiqindi')
sy = M.synthesis(1000.0)
ck(close(sy['m_h2'], 12.46, 0.01) and close(sy['m_br2'], 987.5, 0.5)
   and close(sy['dh_mj'], 639.5, 0.5), 'sintetik: H₂ 12,46 kg; Br₂ 987,5 kg; ΔrH −639,5 MJ')
lb = M.lab_method(1000.0)
ck(close(lb['m_br2'], 987.5, 0.5) and close(lb['m_so2'], 395.9, 0.5)
   and close(lb['m_h2so4'], 606.1, 0.5), 'laboratoriya: Br₂ 987,5; SO₂ 395,9; H₂SO₄ 606,1 kg')
bo = M.bromide_oxidation(1000.0)
ck(close(bo['m_cl2'], 344.6, 0.2) and close(bo['m_br2'], 776.6, 0.2)
   and close(bo['m_nacl'], 568.0, 0.2), 'bromin sikli: Cl₂ 344,6; Br₂ 776,6; NaCl 568,0 kg')
ck(close(bo['m_input'], bo['m_output'], 0.2), 'bromin sikli mass balansi: 1344,6 ≈ 1344,5 kg')
h2 = M.h2o2_bromination(1000.0); o2 = M.oxidation_by_oxygen(1000.0)
ck(close(h2['m_h2o2'], 210.2, 0.2) and close(h2['m_br2'], 987.5, 0.5),
   'H₂O₂ bilan: 210,2 kg H₂O₂; Br₂ 987,5 kg')
ck(close(o2['m_o2'], 98.9, 0.2) and close(o2['m_br2'], 987.5, 0.5),
   'O₂ bilan: O₂ 98,9 kg; Br₂ 987,5 kg')
md = {x['metal']: x for x in M.metal_dissolution(1.0)}
ck(close(md['Fe']['m_salt'], 3.862, 1e-3) and close(md['Fe']['m_hbr'], 2.898, 1e-3)
   and close(md['Fe']['m_h2'], 0.036, 1e-3), '1 g Fe → 3,862 g FeBr₂; HBr 2,898 g; H₂ 0,036 g')
ck(close(md['Al']['m_salt'], 9.884, 1e-3) and close(md['Al']['m_hbr'], 8.996, 1e-3),
   '1 g Al → 9,884 g AlBr₃; HBr 8,996 g')
ck(close(md['Zn']['m_hbr'], 2.475, 1e-3) and close(md['Mg']['m_hbr'], 6.658, 1e-3),
   '1 g Zn → HBr 2,475 g; 1 g Mg → HBr 6,658 g')
ck(close(md['Cu']['m_hbr'], 2.547, 1e-3) and close(md['Cu']['m_o2'], 0.252, 1e-3),
   '1 g Cu → HBr 2,547 g + O₂ 0,252 g')
n1 = M.neutralisation(1000.0, 10.0, 'Na2CO3')
ck(close(n1['m_agent'], 65.5, 0.1) and close(n1['m_salt'], 127.2, 0.1)
   and close(n1['m_gas'], 27.2, 0.1) and close(n1['m_h2o'], 11.1, 0.1),
   'neytralizatsiya: Na₂CO₃ 65,5; NaBr 127,2; CO₂ 27,2; H₂O 11,1 kg')
ck(close(n1['m_input'], n1['m_output'], 0.1), 'neytralizatsiya mass balansi: 165,5 = 165,5 kg')
n2 = M.neutralisation(1000.0, 10.0, 'Ca(OH)2')
ck(close(n2['m_agent'], 45.8, 0.1) and close(n2['m_salt'], 123.5, 0.1)
   and close(n2['m_h2o'], 22.3, 0.1) and close(n2['m_input'], n2['m_output'], 0.1),
   'Ca(OH)₂ bilan: 45,8 / 123,5 / 22,3 kg; kirish = chiqish 145,8 kg')
t = M.titration_standardisation(0.2000, 30.00)
ck(close(t['c_hbr'], 0.12580, 1e-5), 'standartlash: 0,2000 g Na₂CO₃; 30,00 mL → 0,12580 mol/L')
ck(close(M.PHYS['tb'], -66.8) and close(M.PHYS['tm'], -86.9) and close(M.PHYS['tc'], 89.8)
   and close(M.PHYS['pc'], 85.0) and close(M.PHYS['aze_pct'], 47.6)
   and close(M.PHYS['aze_tb'], 124.3) and close(M.PHYS['aze_c'], 8.77)
   and close(M.PHYS['sat'], 68.85) and close(M.PHYS['sol0'], 221.0)
   and close(M.PHYS['sol15'], 204.0) and close(M.PHYS['sol100'], 130.0),
   'fizik konstantalar: −66,8 / −86,9 °C; 89,8 °C; 85,0 bar; azeotrop 47,6 % / 124,3 °C')
g = {x[0]: x for x in M.GROUP}
ck(close(g['HBr'][1], 366.0) and close(g['HBr'][2], -8.8) and close(g['HBr'][3], -6.8)
   and close(g['HF'][1], 569.0) and close(g['HI'][1], 295.0),
   'guruh: D 569/432/366/295; pK_a (suv) va DMSO to‘g‘ri')

print('3) CSV fayllar va rasmlar')
for name, n in [('fizik_hbr', 22), ('eritma', 4), ('spetsiatsiya', 10), ('guruh', 4),
                ('ishlabchiqarish', 12), ('metallar', 5), ('chiqindi', 13), ('tribromid', 9)]:
    rows = list(csv.DictReader((ASSETS / (name + '.csv')).open()))
    ck(len(rows) == n, '%s.csv: %d qator' % (name, len(rows)))
pngs = sorted(p.stem for p in ASSETS.glob('*.png'))
ck(len(pngs) == 17, 'rasmlar soni: %d' % len(pngs))
for n in ['molekulax', 'fizikx', 'guruh', 'olishsx', 'bromokomplekssx', 'metalsx',
          'redoks', 'organiksx', 'qollanishsx', 'tahlilsx', 'hisob',
          'ishlabchiqarish', 'metallar', 'teshikx', 'bromsx', 'idishx', 'chiqindi']:
    ck(n in pngs, 'rasm %s.png bor' % n)

print('4) Matn: model raqamlari va atamalar')
ms = MS.read_text()
tables_txt = json.dumps(json.loads((ROOT / 'matn' / '06_04_jadvallar.json').read_text()), ensure_ascii=False)
notes_txt = json.dumps(json.loads((ROOT / 'matn' / '06_04_yonlavhalar.json').read_text()), ensure_ascii=False)
ms_all = ms + tables_txt + notes_txt
need = ['1,4144', '366 kJ/mol', '0,8272', '2,96', '11,66', '−66,8', '−86,9',
        '89,8', '85,0', '807 kg/m³', '3,307', '2,71', '221', '204', '130', '68,85',
        '47,6', '124,3', '8,77', '1,49', '1,38', '7,0 mol/L', '−11 °C', '126 °C',
        '−8,8', '−6,8', '6,3×10⁸', '1,000', '−0,943', '−1,079', '0,0001',
        '60,40', '99,29', '99,48', '12,46', '987,5', '639,5', '395,9', '606,1',
        '344,6', '776,6', '568,0', '1344,6', '1344,5', '210,2', '98,9',
        '3,862', '2,898', '0,036', '9,884', '8,996', '2,475', '6,658', '2,547', '0,252',
        '65,5', '127,2', '27,2', '11,1', '165,5', '45,8', '123,5', '22,3', '145,8',
        '0,2000', '30,00', '0,12580', '569', '432', '295', '3,17', '−9,5', '−10,9',
        '440 000', '190', '112', '100 kt', '65 ppm', '1 mlrd t',
        '3 ppm', '2 ppm', '35 ppm', '30 ppm', '116 mg/m³', '2 ppm', 'ODP 0,6',
        'AgBr', '[PdBr₄]²⁻', 'Al₂Br₆', 'Br₃⁻', 'BrO', 'Montreal']
for x in need:
    ck(x in ms_all, 'matnda «%s» bor' % x)
ck(ms.count('[[EQ:') == 11 and ms.count('[[TABLE:') == 11, 'matnda 11 formula va 11 jadval belgisi')
ck('amaliyot tavsiyalari berilmaydi' in ms, 'xavfsizlik chegarasi matnda bayon qilindi')
ck('o‘lchov ma’lumoti emas' in ms, 'model cheklovi matnda ko‘rsatilgan')
ck('H₂ + Br₂' in ms and 'SO₂ + 2H₂O' in ms and '2NaBr + Cl₂' in ms, 'asosiy tengliklar matnda')
ck('4HBr + O₂' in ms and 'H₂O₂' in ms, 'redoks tengliklari matnda')
ck('vodorod bog‘i' in ms and 'passivatsiya' in ms and 'embrittlement' in ms,
   'mavzu kalit tushunchalari matnda')
ck('hydrogen bog‘i' not in ms.lower(), 'terminologiya: vodorod bog‘i')
ck('$' not in ms, 'matnda dollar belgisi yo‘q')

print('5) DOCX tuzilmasi')
doc = Document(DOCX)
xml = doc.element.xml
def collect(container):
    out = []
    for t in container.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
                out.extend(collect(c))
    return out


text_all = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                text_all += '\n' + p.text
            text_all += '\n' + '\n'.join(collect(c))
ck(len(doc.inline_shapes) == 17, 'rasmlar soni: %d' % len(doc.inline_shapes))
ck(xml.count('<m:oMath') == 61, 'OMML satrlari (11 guruh, 61 qator): %d' % xml.count('<m:oMath'))
ck(xml.count('<m:f>') >= 1, 'kasrlar: %d' % xml.count('<m:f>'))
ck(xml.count('<m:sSub>') >= 30, 'pastki indekslar: %d' % xml.count('<m:sSub>'))
ck(xml.count('sSup') >= 1, 'yuqori indekslar: %d' % xml.count('sSup'))
ck(text_all.count('-jadval.') == 11, 'jadval sarlavhalari: %d' % text_all.count('-jadval.'))
ck(text_all.count('-sxema.') == 7, 'sxema sarlavhalari: %d' % text_all.count('-sxema.'))
ck(text_all.count('-grafik.') == 5, 'grafik sarlavhalari: %d' % text_all.count('-grafik.'))
ck(text_all.count('-rasm.') == 5, 'rasm sarlavhalari: %d' % text_all.count('-rasm.'))
for i in range(1, 12):
    ck(text_all.count('(6.4.%d)' % i) == 1, 'formula raqami (6.4.%d)' % i)
ck('_' not in text_all, 'matnda "_" qolmadi (pastki indeksga o‘girilgan)')
ck('Times New Roman' in xml, 'Times New Roman shrifti qo‘llangan')
ck('Cambria Math' in xml, 'formulalar Cambria Math shriftida')
ck(text_all.count('HISOBIY GRAFIK') == 2, 'grafik yorliqlari: %d' % text_all.count('HISOBIY GRAFIK'))
ck(text_all.count('TUSHUNTIRUVCHI SXEMA') == 7, 'sxema yorliqlari: %d' % text_all.count('TUSHUNTIRUVCHI SXEMA'))
ck(text_all.count('CHIZMA VA IZOH') == 2, 'chizma yorliqlari: %d' % text_all.count('CHIZMA VA IZOH'))
ck('NOORGANIK KISLOTALAR' in doc.sections[0].header.paragraphs[0].text, 'sarlavha matni')
ck('6-BOB' in doc.sections[0].footer.paragraphs[0].text, 'kolontitul matni')
for x in ['987,5 kg', '639,5 MJ', '2,898 g', '0,12580 mol/L', '776,6 kg', '99,29',
          '6,3×10', '165,5 kg', 'IDLH']:
    ck(x in text_all, 'DOCX matnida «%s» bor' % x)
ck('OSHA' in text_all and 'NIOSH' in text_all and 'UNEP' in text_all,
   'adabiyotlar ro‘yxati joylashdi')
ck('amaliyot tavsiyalari berilmaydi' in text_all, 'xavfsizlik chegarasi bayon qilindi')
size = DOCX.stat().st_size
ck(size > 400000, 'DOCX hajmi: %d bayt' % size)

print()
if fails:
    print('XATO: %d tekshiruv bajarilmadi' % len(fails))
    for f in fails: print('  -', f)
    raise SystemExit(1)
print('BARCHA TEKSHIRUVLAR O‘TDI (6.4).')
