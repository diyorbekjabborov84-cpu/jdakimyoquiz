"""QA: model identities, native mathematics, layout and figure inventory for 4.4."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm
import math,csv,re,json
from complex_04_04 import (KSP,BETA2,PKA_NH4,KSP_M,BETA4,CM0,alpha_nh3,
    conditional_beta,complex_fraction,saturation_s,amph_csat,amph_minimum,
    separation_onset,separation_quantitative,worked)
D=Path(__file__).resolve().parents[1];A=D/'rasmlar/kompleks'
def near(a,b):assert math.isclose(a,b,rel_tol=2e-9,abs_tol=1e-14),(a,b)
def rows(name):return [{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/(name+'.csv')).open())]
points=0
for name,n in [('species',401),('proton',401),('solubility',401),('amph',401)]:
    rr=rows(name);assert len(rr)==n;points+=n
    for r in rr:
        if name=='species':
            near(r['fraction_free_Ag']+r['fraction_complex'],1.0)
            L=r['free_L_M'];near(L,10**r['log_free_L'])
            near(r['fraction_complex'],complex_fraction(L))
            near(r['fraction_free_Ag'],1-complex_fraction(L))
            assert 0.0<=r['fraction_complex']<=1.0
        elif name=='proton':
            near(r['alpha_NH3'],alpha_nh3(r['pH']))
            near(r['log_beta_conditional'],math.log10(conditional_beta(r['pH'])))
            assert 0.0<r['alpha_NH3']<1.0
        elif name=='solubility':
            L=r['free_L_M'];near(L,10**r['log_free_L'])
            near(r['saturated_s_M'],saturation_s(L));near(r['log_s'],math.log10(r['saturated_s_M']))
        else:
            near(r['Csat_M'],amph_csat(r['pH']));near(r['log_Csat'],math.log10(r['Csat_M']))
            near(r['Cdiss_M'],min(r['Csat_M'],CM0))
# Model identities that the prose asserts.
near(math.sqrt(KSP),1.2649110640673517e-05)
near(saturation_s(1e-2),5.216895625628197e-04)
near(complex_fraction(1/math.sqrt(BETA2)),0.5)
x,pH,cs=amph_minimum();near(pH,5.619280313679916);near(cs,5.549055267050422e-05)
near(BETA4*x**4,3.0)
near(alpha_nh3(9.25),0.5);near(alpha_nh3(7.25),0.009900990099009901);near(alpha_nh3(11.25),0.99009901)
# Separation model quoted in the text and tables.
near(separation_onset(1e-18,0.010),6.0)
near(separation_onset(1e-12,0.010),9.0)
near(separation_quantitative(1e-18,0.010,fraction=0.999),7.5)
near(separation_onset(1e-18,0.010,F=1e4),8.0)
near(separation_quantitative(1e-18,0.010,F=1e4,fraction=0.999),9.5)
# Worked ligand-balanced example.
w=worked()
near(w['alpha'],0.8490204427886767)
near(w['free_NH3'],0.007799796153051602)
near(w['s'],0.0004069842140419101)
near(w['free_Ag'],3.931356413335571e-07)
near(w['complex'],0.00040659107840057654)
near(w['unbound'],0.009186817843198847)
near(w['bound'],0.0008131821568011531)
near(w['total'],0.01)
near(w['complex_percent'],99.90340272969578)
near(w['naive_s'],0.0004429755227304077)
near(w['naive_error_percent'],8.843416389804082)
# Independent atom/charge balances for the printed reactions.
spec={'AgCl':({'Ag':1,'Cl':1},0),'Ag':({'Ag':1},0),'AgI':({'Ag':1},1),'Cl':({'Cl':1},-1),'AgNH3_2':({'Ag':1,'N':2,'H':6},1),
      'NH3':({'N':1,'H':3},0),'HSO4':({'H':1,'S':1,'O':4},-1),'SO4':({'S':1,'O':4},-2),'H':({'H':1},1),
      'S':({'S':1},-2),'H2S':({'H':2,'S':1},0),'CaSO4':({'Ca':1,'S':1,'O':4},0),'Ca':({'Ca':1},2),
      'AlOH3':({'Al':1,'O':3,'H':3},0),'Al3':({'Al':1},3),'H2O':({'H':2,'O':1},0),'OH':({'O':1,'H':1},-1),
      'AlOH4':({'Al':1,'O':4,'H':4},-1),'NH4':({'N':1,'H':4},1)}
rx=[({'AgCl':1},{'AgI':1,'Cl':1}),
    ({'AgI':1,'NH3':2},{'AgNH3_2':1}),
    ({'HSO4':1},{'H':1,'SO4':1}),
    ({'S':1,'H':2},{'H2S':1}),
    ({'CaSO4':1},{'Ca':1,'SO4':1}),
    ({'AlOH3':1,'H':3},{'Al3':1,'H2O':3}),
    ({'AlOH3':1,'OH':1},{'AlOH4':1})]
def total(side):
    c=Counter();charge=0
    for k,n in side.items():
        if not n:continue
        a,z=spec[k];charge+=n*z
        for atom,v in a.items():c[atom]+=n*v
    return c,charge
for left,right in rx:
    assert total(left)==total(right),(left,right)
p=D/'chiqish/Noorganik-kislotalar-4.4-Kompleks-va-chokma-hosil-bolish.docx'
doc=Document(p)
def paras(o):
    yield from o.paragraphs
    for t in o.tables:
        for row in t.rows:
            for c in row.cells:yield from paras(c)
ps=list(paras(doc));txt='\n'.join(x.text for x in ps)
assert len(doc.inline_shapes)==16
for label,n in [('jadval.',11),('grafik.',4),('sxema.',7),('rasm.',5)]:assert sum(label in x.text for x in ps)==n,(label,)
assert all(txt.count(f'(4.4.{i})')==1 for i in range(1,12))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt and '4.3.' not in txt
assert all(r.font.name=='Times New Roman' for x in ps for r in x.runs if r.text)
outer=[t for t in doc.tables if len(t.columns)==2];assert len(outer)==11
for t in outer:assert abs(t.columns[0].width-Cm(13.5))<1000 and abs(t.columns[1].width-Cm(4.5))<1000
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(p) as z:assert z.testzip() is None;tree=etree.fromstring(z.read('word/document.xml'))
omml=len(tree.xpath('//m:oMath',namespaces=ns));fracs=len(tree.xpath('//m:f',namespaces=ns))
assert omml==23,(omml,)
assert fracs==13,(fracs,)
assert len(tree.xpath('//w:tblHeader',namespaces=ns))==11
alltext=' '.join(tree.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
without_urls=re.sub(r'https?://\S+','',alltext)
assert '_' not in without_urls,[m for m in re.findall(r'\S*_\S*',without_urls)][:8]
assert len(tree.xpath('//m:sSub',namespaces=ns))>25 and len(tree.xpath('//m:sSup',namespaces=ns))>10
for key in ['K_sp','Q_sp','α_L','C_M,0','β_2','F_M','n_s','C_sat','K_a,c']:
    assert key not in alltext,key
assert 'β 6 =7×10' in alltext and 'AlF 6' in alltext
notes=json.loads((D/'matn/04_04_yonlavhalar.json').read_text());assert sum(map(len,notes))==22
tables=json.loads((D/'matn/04_04_jadvallar.json').read_text());assert len(tables)==11
print(f'PASS: {points} model points; speciation, protonation, solubility, amphoteric minimum, separation and ligand-balance identities; {len(rx)} atom/charge balances; {omml} OMML lines / {fracs} fractions; 11 tables, 16 figures, 22 notes, 3:1 layout.')
print('Words with math:',len(alltext.split()),'Bytes:',p.stat().st_size)
rt=D.parent/'google_drive/Noorganik kislotalar — 4.4. Kompleks va cho‘kma hosil qilish — tekshirilgan.roundtrip.docx'
if rt.exists():
    with ZipFile(rt) as z:gt=etree.fromstring(z.read('word/document.xml'))
    assert len(gt.xpath('//m:oMath',namespaces=ns))==omml
    assert len(gt.xpath('//m:f',namespaces=ns))==fracs
    gs=' '.join(gt.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
    assert gs.count('jadval.')==11 and all(f'(4.4.{i})' in gs for i in range(1,12))
    print('Google Docs→DOCX roundtrip PASS: native math lines and fractions retained.')
