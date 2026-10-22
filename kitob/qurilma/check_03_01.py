"""Reproducible numerical and DOCX integrity checks for topic 3.1."""
from pathlib import Path
from math import isclose,sqrt,log10
import csv,zipfile
from docx import Document
from equilibria_03_01 import quadratic,with_water,diprotic,KC,KWC
ROOT=Path(__file__).resolve().parents[1]
for j in range(321):
 c=10**(-10+j*.025);x=quadratic(c)
 assert 0<x<c and isclose(x*x/(c-x),KC,rel_tol=1e-8)
 h,a=with_water(c);an=c*a;oh=KWC/h
 assert 0<a<1
 assert isclose(h,an+oh,rel_tol=2e-14)
 assert isclose(h*an/(c-an),KC,rel_tol=1e-12)
 assert isclose((c-an)+an,c,rel_tol=1e-15)
 assert h>=sqrt(KWC)
for j in range(481):
 ph=j*.025;a=diprotic(ph);h=10**(-ph)
 assert isclose(sum(a),1,abs_tol=2e-15)
 assert isclose(h*a[1]/a[0],1e-3,rel_tol=1e-13)
 assert isclose(h*a[2]/a[1],1e-7,rel_tol=1e-13)
for c,ph in [(.01,3.5068665126),(.0001,4.568385042)]:assert isclose(-log10(quadratic(c)),ph,abs_tol=1e-9)
h,a=with_water(1e-8);assert isclose(-log10(h),6.9785198198,abs_tol=1e-9)
assert isclose(a,.98960221486,abs_tol=1e-10)
p=ROOT/'chiqish'/'Noorganik-kislotalar-3.1-Ka-pKa-alfa.docx';doc=Document(p)
def paras(obj):
 yield from obj.paragraphs
 for t in obj.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(doc));txt='\n'.join(p.text for p in ps)
assert len(doc.inline_shapes)==13
assert sum('jadval.' in p.text for p in ps)==9
assert sum('rasm.' in p.text for p in ps)==4
assert '[[EQ:' not in txt and '[[TABLE:' not in txt
assert '2.5.' not in txt and '2-BOB' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode()
 assert 'subscript' in xml and 'superscript' in xml
print('PASS: 321 monoacid cases, 481 diprotic cases, examples, DOCX ZIP, fonts, indices, 13 images and 9 tables.')
print('Recursive words:',len(txt.split()),'Bytes:',p.stat().st_size)
