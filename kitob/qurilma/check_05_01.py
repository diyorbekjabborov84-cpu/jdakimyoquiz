#!/usr/bin/env python3
"""5.1 moduli uchun QA: taxminiy hisoblarning matn va hujjat bilan mosligi."""
from pathlib import Path
import csv, math, re, json, zipfile
from docx import Document
from docx.oxml.ns import qn
import prep_05_01 as M

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'chiqish' / 'Noorganik-kislotalar-5.1-Olishning-umumiy-usullari.docx'
ASSETS = ROOT / 'rasmlar' / 'olish'
MS = ROOT / 'matn' / '05_01.md'
fails = []


def ck(cond, msg):
    if cond:
        print('  OK  ', msg)
    else:
        print('  FAIL', msg); fails.append(msg)


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol


print('1) Taxminiy hisob: termodinamik mezon')
dh_so2, ds_so2 = M.dH(*M.RX_SO2), M.dS(*M.RX_SO2)
ck(close(dh_so2, -98.89, 1e-6), 'SO₂ oksidlanishi ΔH° = −98,89 kJ/mol')
ck(close(ds_so2, -94.04, 1e-6), 'SO₂ oksidlanishi ΔS° = −94,04 J/(mol·K)')
ck(close(M.dG(dh_so2, ds_so2), -70.85, 1e-2), 'SO₂ oksidlanishi ΔG°(298,15) = −70,85 kJ/mol')
ck(close(M.lgK(dh_so2, ds_so2), 12.413, 1e-3), 'SO₂ oksidlanishi lg K = 12,413')
dh_d, ds_d = M.dH(*M.RX_DIMER), M.dS(*M.RX_DIMER)
ck(close(dh_d, 55.30, 1e-6) and close(ds_d, 175.80, 1e-6), 'N₂O₄ ⇌ 2NO₂: ΔH° = +55,30; ΔS° = +175,80')
ck(close(M.dG(dh_d, ds_d), 2.89, 1e-2), 'N₂O₄ ⇌ 2NO₂: ΔG°(298,15) = +2,89 kJ/mol')
ck(close(M.lgK(dh_d, ds_d), -0.5055, 1e-3), 'N₂O₄ ⇌ 2NO₂: lg K = −0,506')
rx_hcl = ({'H2(g)': 1.0, 'Cl2(g)': 1.0}, {'HCl(g)': 2.0})
dh_h, ds_h = M.dH(*rx_hcl), M.dS(*rx_hcl)
ck(close(dh_h, -184.61, 1e-2) and close(ds_h, 20.00, 1e-6), 'H₂ + Cl₂ → 2HCl: ΔH° = −184,61; ΔS° = +20,00')
ck(close(M.dG(dh_h, ds_h), -190.58, 1e-2) and close(M.lgK(dh_h, ds_h), 33.388, 1e-2),
   'H₂ + Cl₂ → 2HCl: ΔG° = −190,58; lg K = 33,39')
rx_nh3 = ({'NH3(g)': 4.0, 'O2(g)': 5.0}, {'NO(g)': 4.0, 'H2O(l)': 6.0})
dh_n, ds_n = M.dH(*rx_nh3), M.dS(*rx_nh3)
ck(close(dh_n, -1170.38, 1e-2) and close(ds_n, -534.00, 1e-6), 'NH₃ oksidlanishi: ΔH° = −1170,38; ΔS° = −534,00')
ck(close(M.dG(dh_n, ds_n), -1011.17, 1e-1) and close(M.lgK(dh_n, ds_n), 177.15, 1e-1),
   'NH₃ oksidlanishi: ΔG° = −1011,17; lg K = 177,15')

print('2) Taxminiy hisob: muvozanat, van ’t Hoff, kaskad, Fenske')
for T, x1, x10, x100 in [(250.0, 0.0326, 0.0103, 0.0033), (298.15, 0.2691, 0.0880, 0.0279),
                         (350.0, 0.8247, 0.4188, 0.1443), (400.0, 0.9788, 0.8341, 0.4314)]:
    ck(close(M.eq_conversion(T, 1.0), x1, 2e-3) and close(M.eq_conversion(T, 10.0), x10, 2e-3)
       and close(M.eq_conversion(T, 100.0), x100, 2e-3),
       'N₂O₄ konversiyasi %g K: %.3f / %.3f / %.3f' % (T, x1, x10, x100))
lh = M.vanthoff(dh_d, 300.0, 400.0)
ck(close(lh, 5.5426, 1e-3) and close(math.exp(lh), 255.3, 0.5), 'van ’t Hoff 300→400 K: K 255 marta')
for n, x in [(1, 0.60), (2, 0.84), (5, 0.9898), (10, 0.9999)]:
    ck(close(M.cascade(0.6, n), x, 1e-3), 'Kaskad x₁ = 0,60, N = %d → %.3f' % (n, x))
for a, nmin in [(1.1, 61.79), (1.5, 14.52), (2.0, 8.50), (3.0, 5.36), (5.0, 3.66)]:
    ck(close(M.fenske(a), nmin, 5e-3), 'Fenske α = %.1f → %.2f pog‘ona' % (a, nmin))

print('3) Taxminiy hisob: pyrit zanjiri va resurs ko‘rsatkichlari')
r = M.pyrite_chain(1000.0, 0.96, 0.98)
for k, v, tol in [('n_fes2', 8334.7225, 1e-3), ('n_so2', 16669.4449, 1e-3),
                  ('n_so3', 16002.6671, 1e-3), ('n_so2_left', 666.7778, 1e-3),
                  ('n_h2so4', 15682.6138, 1e-3), ('n_fe2o3', 4167.3612, 1e-3),
                  ('m_h2so4', 1538.1508, 1e-3), ('m_theoretical', 1634.9392, 1e-3),
                  ('m_residue', 665.4859, 1e-3), ('m_so2_left', 42.7138, 1e-3),
                  ('m_o2', 256.0427, 1e-3), ('m_h2o_used', 282.5223, 1e-3),
                  ('yield', 0.9408, 1e-4), ('E_factor', 0.4604, 1e-4)]:
    ck(close(r[k], v, tol), 'pyrit zanjiri: %s = %.4f' % (k, v))
w = M.acid_water_96(r['m_h2so4'])
ck(close(w['m_solution'], 1602.24, 1e-2) and close(w['m_water_in_product'], 64.09, 1e-2),
   '96 % kislota: eritma 1602,24 t; suv 64,09 t')
ck(close(M.efactor(0.90), 0.5355, 1e-4) and close(M.efactor(0.99), 0.4263, 1e-4),
   'Chiqindi omili: X = 0,90 → 0,5355; X = 0,99 → 0,4263')

print('4) Taxminiy hisob: CSV fayllar')
for name, n in [('mivozanat', 401), ('kaskad', 401), ('chiqindi', 401), ('buglanish', 401), ('dG', 251)]:
    rows = list(csv.DictReader((ASSETS / (name + '.csv')).open()))
    ck(len(rows) == n, '%s.csv: %d qator' % (name, n))

print('5) Matn: taxminiy hisob raqamlari matnda bor')
ms = MS.read_text()
tables_txt = json.dumps(json.loads((ROOT / 'matn' / '05_01_jadvallar.json').read_text()), ensure_ascii=False)
ms_all = ms + tables_txt
need = ['−98,89', '−94,04', '−70,85', '12,41', '+2,89', '−0,506', '−1011,17', '177,15',
        '−190,58', '33,39', '0,033', '0,979', '0,269', '0,088', '0,028', '255 marta',
        '8334,72', '16 669,44', '16 002,67', '666,78', '42,71', '15 682,61', '1538,15',
        '1634,94', '665,49', '256,04', '282,52', '346,61', '94,08', '0,4604', '0,54', '0,43',
        '61,8', '8,5', '3,7', '0,60', '0,84', '0,990']
for s in need:
    ck(s in ms_all, 'matnda «%s» bor' % s)
ck(ms.count('[[EQ:') == 11 and ms.count('[[TABLE:') == 11, 'matnda 11 formula va 11 jadval belgisi')

print('6) DOCX tuzilmasi')
doc = Document(DOCX)
xml = doc.element.xml
text_all = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                text_all += '\n' + p.text
ck(len(doc.inline_shapes) == 17, 'rasmlar soni: %d' % len(doc.inline_shapes))
ck(xml.count('<m:oMath') == 33, 'OMML guruhlari (11 ta formula, 33 qator): %d' % xml.count('<m:oMath'))
ck(xml.count('<m:f>') >= 18, 'kasrlar: %d' % xml.count('<m:f>'))
ck(xml.count('<m:sSub>') >= 40, 'pastki indekslar: %d' % xml.count('<m:sSub>'))
ck(xml.count('<m:sSup>') >= 5, 'yuqori indekslar: %d' % xml.count('<m:sSup>'))
ck(text_all.count('-jadval.') == 11, 'jadval sarlavhalari: %d' % text_all.count('-jadval.'))
ck(text_all.count('sxema.') == 7, 'sxema sarlavhalari: %d' % text_all.count('sxema.'))
ck(text_all.count('grafik.') == 5, 'grafik sarlavhalari: %d' % text_all.count('grafik.'))
ck(text_all.count('rasm.') == 5, 'rasm sarlavhalari: %d' % text_all.count('rasm.'))
for i in range(1, 12):
    ck(text_all.count('(5.1.%d)' % i) == 1, 'formula raqami (5.1.%d)' % i)
ck('_' not in text_all, 'matnda "_" qolmadi (pastki indeksga o‘girilgan)')
ck('Times New Roman' in xml, 'Times New Roman shrifti qo‘llangan')
ck('Cambria Math' in xml, 'formulalar Cambria Math shriftida')
ck(text_all.count('HISOBIY GRAFIK') == 5, 'grafik yorliqlari: %d' % text_all.count('HISOBIY GRAFIK'))
ck(text_all.count('TUSHUNTIRUVCHI SXEMA') == 6, 'sxema yorliqlari: %d' % text_all.count('TUSHUNTIRUVCHI SXEMA'))
hdr = doc.sections[0].header.paragraphs[0].text
ftr = doc.sections[0].footer.paragraphs[0].text
ck('NOORGANIK KISLOTALAR' in hdr, 'sarlavha matni')
ck('5-BOB' in ftr and '5.1' in text_all, 'kolontitul matni')
for s in ['8334,72', '1538,15', '1634,94', '665,49', '346,61', '0,4604', '−70,85', '12,41']:
    ck(s in text_all, 'DOCX matnida «%s» bor' % s)
ck('OpenStax' in text_all and 'Gold Book' in text_all, 'adabiyotlar ro‘yxati joylashdi')
ck('amaliyot tavsiyalari berilmaydi' in text_all, 'xavfsizlik chegarasi bayon qilindi')
size = DOCX.stat().st_size
ck(size > 400000, 'DOCX hajmi: %d bayt' % size)

print()
if fails:
    print('XATO: %d tekshiruv bajarilmadi' % len(fails))
    for f in fails: print('  -', f)
    raise SystemExit(1)
print('BARCHA TEKSHIRUVLAR O‘TDI (5.1).')
