#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. Alohida mavzu moduli."""
from book_equations import add_equation
from pathlib import Path
import re, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'chiqish'/'Noorganik-kislotalar-4.3-Passivlanish.docx'
ASSETS=ROOT/'rasmlar'/'passivlanish_sxema'; ASSETS.mkdir(parents=True,exist_ok=True)
NAVY='0B2C4A'; BRONZE='C28E3B'; PALE='EDF3F7'; INK='243746'
fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
def font(n):return ImageFont.truetype(fontpath,n)

def scheme(index,nodes):
    special={1:'holatlar.png',2:'parda.png',3:'qutblanish.png',4:'aralash.png',5:'qalinlik.png',6:'xlorid.png',7:'yuqori.png',8:'dalil.png',9:'hisob.png'}
    return ROOT/'rasmlar'/'passivlanish'/special[index]

panels=[
('Holatlarni ajrating',[],'4.3.1-sxema. Immun, faol va passiv holatlar; tashqi tok bilan ichki qisman toklarning farqi.'),
('Parda — dinamik to‘siq',[],'4.3.2-sxema. Tarkib, uzluksizlik, transport va hosil bo‘lish–erish raqobati.'),
('Qutblanishni o‘qing',[],'4.3.1-grafik. Shartli anod egri chizig‘i; i_crit=10⁻³, i_pass=10⁻⁶ A cm⁻². Ixtiyoriy o‘quv etaloni, tajriba emas.'),
('Oksidlovchi va kesishish',[],'4.3.2-grafik. Bir xil sun’iy anod egri chizig‘iga ikki katod yo‘li. Musbat |i_c| lar berilgan; kesishishlar faol va passiv holatlarda.'),
('Qalinlik va vaqt',[],'4.3.3-grafik. Cheklangan o‘sish va doimiy erish bo‘yicha shartli hisob; 50 s vaqt ko‘lami, uchta boshqa limit. Eksperimental ma’lumot emas.'),
('Mahalliy buzilish',[],'4.3.4-grafik. Cl⁻ faolligiga shartli E_pit va E_rp bog‘lanishi. Real material chegarasi emas; boshqa shartli etalon bilan birlashtirilmaydi.'),
('Yuqori potensial yo‘llari',[],'4.3.3-sxema. Transpassiv erish, O₂ ajralishi va mahalliy buzilishning turli ma’nolari.'),
('Dalil va usul',[],'4.3.4-sxema. Elektrokimyo, sirt tarkibi, morfologiya va metall yo‘qotilishini birga talqin qilish.'),
('Massa va mahalliy chuqurlik',[],'4.3.5-sxema. Bir xil jami metall yo‘qotilishining bir tekis va mahalliy taqsimlanish ekvivalentlari.')
]

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for sty in doc.styles:
    if sty.type==1:
        sty.font.name='Times New Roman'
        rp=sty.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
        rp.append(rf)
sub=dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ','0123456789abmj'));sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡','0123456789+−mn‡'))
variables={
'E_cell':('E','cell'),'i_net':('i','net'),'i_corr':('i','corr'),'i_crit':('i','crit'),'i_pass':('i','pass'),
'i_a':('i','a'),'i_c':('i','c'),'i_M':('i','M'),'E_corr':('E','corr'),'E_rel':('E','rel'),'E_ref':('E','ref'),
'E_pit,ref':('E','pit,ref'),'E_rp,ref':('E','rp,ref'),'E_pit':('E','pit'),'E_rp':('E','rp'),
'a_Cl,ref':('a','Cl,ref'),'a_Cl':('a','Cl'),'L_inf':('L','∞'),'L_lim':('L','lim'),
'k_g':('k','g'),'v_diss':('v','diss'),'R_p':('R','p'),'b_a':('b','a'),'b_c':('b','c'),
'h_uniform':('h','uniform'),'h_local':('h','local')}

def runs(p,text,size=11,bold=False,color=INK):
    pattern='('+'|'.join(re.escape(v) for v in sorted(variables,key=len,reverse=True))+r'|10⁻pH|[₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡]+)'
    def put(t,lower=False,upper=False):
        r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if lower:r.font.subscript=True
        if upper:r.font.superscript=True
    for token in re.split(pattern,text):
        if not token:continue
        if token=='10⁻pH':put('10');put('−pH',upper=True)
        elif token in variables:
            main,low=variables[token];put(main);put(low,lower=True)
        elif all(c in sub for c in token):put(''.join(sub[c] for c in token),lower=True)
        elif all(c in sup for c in token):put(''.join(sup[c] for c in token),upper=True)
        else:put(token)

def shade(cell,fill):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(x)
def cell_margin(cell,**kwargs):
    x=OxmlElement('w:tcMar')
    for k,v in kwargs.items():
        y=OxmlElement('w:'+k);y.set(qn('w:w'),str(v));y.set(qn('w:type'),'dxa');x.append(y)
    cell._tc.get_or_add_tcPr().append(x)
def border(cell,edge,color,size):
    bs=OxmlElement('w:tcBorders');b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:color'),color);b.set(qn('w:sz'),str(size));bs.append(b);cell._tc.get_or_add_tcPr().append(bs)
def band(label,fill,color,size,accent=False):
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(18)
    c=t.cell(0,0);shade(c,fill);cell_margin(c,top=140,bottom=140,left=200,right=160)
    if accent:border(c,'left',BRONZE,24)
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    runs(p,label,size,True,color)
    # Prevent the banner row separating from its following paragraph.
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.keep_with_next=True

comparisons=json.loads((ROOT/'matn'/'04_03_jadvallar.json').read_text())
comparison_count=0

def add_comparison_table(cell,spec):
    global comparison_count
    comparison_count+=1
    p=cell.add_paragraph();p.paragraph_format.space_before=Pt(8)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    runs(p,spec['caption'],9.5,True,NAVY)
    t=cell.add_table(rows=1,cols=len(spec['headers']));t.autofit=False
    for col,w in zip(t.columns,spec['widths']):col.width=Cm(w)
    for i,label in enumerate(spec['headers']):
        c=t.rows[0].cells[i];c.width=Cm(spec['widths'][i]);shade(c,NAVY)
        cell_margin(c,top=70,bottom=70,left=85,right=85)
        p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        runs(p,label,9.5,True,'FFFFFF')
    rp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rp.append(repeat)
    for ri,row in enumerate(spec['rows']):
        r=t.add_row();trp=r._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        for ci,value in enumerate(row):
            c=r.cells[ci];c.width=Cm(spec['widths'][ci])
            cell_margin(c,top=70,bottom=70,left=85,right=85)
            if ri%2==0:shade(c,PALE)
            border(c,'bottom','D6DFE7',3)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            runs(p,value,9.5,ci==0,INK)
    # Word requires a trailing paragraph after a nested table.
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(3)

source=(ROOT/'matn'/'04_03.md').read_text()
sections=[]
for raw in source.split('## ')[1:]:
    title,body=raw.split('\n',1);sections.append((title,body.strip().split('\n\n')))
assert len(sections)==9
side_notes=json.loads((ROOT/'matn'/'04_03_yonlavhalar.json').read_text())
assert len(side_notes)==len(sections)
counter=0
for idx,((title,paras),(panel,nodes,note)) in enumerate(zip(sections,panels)):
    if idx:
        pb=doc.add_page_break();pb.paragraph_format.space_after=Pt(0);pb.paragraph_format.space_before=Pt(0);pb.paragraph_format.line_spacing=Pt(1)
        for r in pb.runs:r.font.size=Pt(1)
    if idx==0:
        band('4-BOB  /  UMUMIY KIMYOVIY XOSSALAR',NAVY,'FFFFFF',11)
        band('4.3. Passivlanish',PALE,NAVY,18,True)
    else:
        band('4.3  /  PASSIVLANISH',PALE,NAVY,11,True)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=True
    runs(p,title,14,True,NAVY)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    # 18 cm printable width: 13.5 cm narrative + 4.5 cm outer visual strip.
    table.columns[0].width=Cm(13.5);table.columns[1].width=Cm(4.5)
    left,right=table.rows[0].cells
    left.width=Cm(13.5);right.width=Cm(4.5)
    for c in [left,right]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left,top=0,bottom=0,left=0,right=220)
    cell_margin(right,top=140,bottom=130,left=160,right=100)
    shade(right,'F5F8FA');border(right,'left','D2DCE5',6)
    # Allow long content to flow rather than shrink type or clip a fixed-height row.
    first=True
    for text in paras:
        if text.startswith('[[TABLE:'):
            key=text[len('[[TABLE:'):-2]
            add_comparison_table(left,comparisons[key])
            continue
        p=left.paragraphs[0] if first else left.add_paragraph()
        j=0 if first else 1
        first=False
        if text.startswith('[[EQ:') and text.endswith(']]'):
            text=text[5:-2]
            counter+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            add_equation(p,left,'4.3',counter,runs)
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Cm(.45 if j else 0)
            runs(p,text)
    p=right.paragraphs[0];p.paragraph_format.space_after=Pt(9)
    runs(p,panel,10,True,NAVY)
    p=right.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(scheme(idx+1,nodes)),width=Cm(3.92))
    p=right.add_paragraph();p.paragraph_format.space_after=Pt(10)
    runs(p,'HISOBIY GRAFIK' if idx in [2,3,4,5] else 'TUSHUNTIRUVCHI SXEMA',9,True,NAVY)
    p=right.add_paragraph();runs(p,note,9)
    extra_models={
      1:[('kesim.png','4.3.1-rasm. Metall–parda–eritma kesimining shartli tasviri.')],
      4:[('tiklanish.png','4.3.2-rasm. Ochiq joyning mos muhitda qayta qoplanishi mumkin.')],
      5:[('pit_kesim.png','4.3.3-rasm. Mahalliy anod va chuqurchadagi boshqa eritma tarkibi.')],
      6:[('trans_kesim.png','4.3.4-rasm. Cr(III) qatlamdan eruvchan yuqori valentli turga formal o‘tish.')]}
    for image,caption in extra_models.get(idx,[]):
        p=right.add_paragraph();p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'rasmlar'/'passivlanish'/image),width=Cm(3.92))
        p=right.add_paragraph();runs(p,caption+' Rang va o‘lchamlar shartli; mikrofoto emas.',9)

    for label,body in side_notes[idx]:
        p=right.add_paragraph()
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        # Delicate bronze rule groups each short aside without heavy boxes.
        pp=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr')
        rule=OxmlElement('w:top');rule.set(qn('w:val'),'single')
        rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),BRONZE)
        rule.set(qn('w:space'),'5');bd.append(rule);pp.append(bd)
        runs(p,label,9,True,NAVY)
        p=right.add_paragraph()
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        runs(p,body,9)
    if idx==8:
        p=left.add_paragraph();p.paragraph_format.space_before=Pt(8);runs(p,'Adabiyotlar',10,True,NAVY)
        refs=[
        '[1] IUPAC Gold Book. Passivation. DOI: 10.1351/goldbook.P04443. https://goldbook.iupac.org/terms/view/P04443 . Faol holatdan passiv parda orqali o‘tish tushunchasi; mustaqil bayon.',
        '[2] Scully, J. R.; Lutton, K. Polarization Behavior of Active Passive Metals and Alloys. University of Virginia mualliflarining ochiq bob qo‘lyozmasi. https://www.numis.northwestern.edu/MURI/Publications/Papers/Scully_IFCC.pdf . Faol–passiv kinetika, parda tarkibi, aralash potensial va transpassivlik.',
        '[3] IUPAC Gold Book. Pitting corrosion. DOI: 10.1351/goldbook.P04679. https://goldbook.iupac.org/terms/view/P04679 . Mahalliy korroziya tushunchasi.',
        '[4] Gamry Instruments. Basics of Electrochemical Corrosion Measurements. https://www.gamry.com/application-notes/corrosion-coatings/basics-of-electrochemical-corrosion-measurements/ . Tafel/LPR talqini va Faraday qonuni; tayyor uskunaviy tartib takrorlanmagan.',
        '[5] Jones, D. A. Principles and Prevention of Corrosion. 2-nashr. Prentice Hall, 1996. Umumiy bibliografik tayanch; aniq sahifalar tekshirilmagan.',
        'Barcha grafik, sxema va rasmlar original o‘quv ishlanmalari. Sonli egri chiziqlar real metall/kislota ma’lumotlari emas. Kesimlar mikrofoto, atom tuzilishi yoki masshtabli qatlam tasviri emas. Nazariy misollar xavfli kislotalarda tajriba yoki idish materiali tanlash uchun ko‘rsatma emas.'
        ]
        for ref in refs:
            p=left.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3);runs(p,ref,9)

p=sec.header.paragraphs[0];runs(p,'NOORGANIK KISLOTALAR',8,False,NAVY)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runs(p,'4-BOB   ·   ',8,False,NAVY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Noorganik kislotalar · 4.3 · Passivlanish'
doc.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
assert len(doc.inline_shapes)==13
assert counter==8
assert comparison_count==9
assert sum(len(notes) for notes in side_notes)==18
print(OUT)
print('9 sections; 9 tables; 5 schemes; 4 calculated graphs; 4 process illustrations; 8 equations; 18 margin notes.')
