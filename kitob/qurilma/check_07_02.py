#!/usr/bin/env python3
"""7.2 moduli uchun QA: taxminiy hisoblarning matn va hujjat bilan mosligi."""
from pathlib import Path
import csv, math, re, json
from docx import Document
import prep_07_02 as M

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'chiqish' / 'Noorganik-kislotalar-7.2-Sulfit-kislota.docx'
ASSETS = ROOT / 'rasmlar' / 'sulfit'
MS = ROOT / 'matn' / '07_02.md'
fails = []


def ck(cond, msg):
    if cond:
        print('  OK  ', msg)
    else:
        print('  FAIL', msg); fails.append(msg)


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol


print('1) Taxminiy hisob: kislotalik va spetsiatsiya')
ss = M.saturated_solution()
ck(close(ss['c'], 1.4674, 1e-4) and close(ss['ph'], 0.845, 0.01),
   'to‘yingan eritma: c = 1,4674 mol/L;  pH = 0,845')
ck(close(ss['h'], 0.1428, 1e-4) and close(ss['pct'], 9.73, 0.02),
   'to‘yingan eritma: h = 0,1428 mol/L;  α = 9,73 %')
sp = M.speciation(2.0)
ck(close(sp['so2'], 39.37, 0.01) and close(sp['hso3'], 60.63, 0.01),
   'pH 2: SO₂(aq) 39,37 %;  HSO₃⁻ 60,63 %')
sp = M.speciation(7.0)
ck(close(sp['hso3'], 61.35, 0.01) and close(sp['so3'], 38.65, 0.01),
   'pH 7: HSO₃⁻ 61,35 %;  SO₃²⁻ 38,65 %')
sp = M.speciation(10.0)
ck(close(sp['so3'], 99.84, 0.01), 'pH 10: SO₃²⁻ 99,84 %')
p = M.ph_of(0.100)
ck(close(p['h'], 0.032291, 1e-6) and close(p['ph'], 1.491, 1e-2)
   and close(p['pct'], 32.29, 0.02), '0,100 M: h = 0,03229 M;  pH = 1,49;  α = 32,3 %')
p = M.ph_of(1.000)
ck(close(p['h'], 0.116635, 1e-5) and close(p['ph'], 0.933, 1e-2),
   '1,000 M: h = 0,11664 M;  pH = 0,93')
ck(close(M.KA1, 1.54e-2) and close(M.KA2, 6.3e-8),
   'K_a1 = 1,54×10⁻²;  K_a2 = 6,3×10⁻⁸')
ck(close(M.KW / M.KA2, 1.587e-7, 1e-9), 'K_b(SO₃²⁻) = K_w/K_a2 ≈ 1,59×10⁻⁷')

print('2) Taxminiy hisob: fizik, eruvchanlik, tahlil, neytralizatsiya, kontakt')
ck(close(M.vapor_pressure(298.15), 3.8541, 1e-3), 'p(25 °C) = 3,85 bar (Clausius–Clapeyron)')
ck(close(M.vapor_pressure_table_25() / 101.325, 3.8536, 1e-3),
   'jadval interpolatsiyasi = 3,85 bar')
ck(close(100.0 * (M.vapor_pressure(298.15) - M.vapor_pressure_table_25() / 101.325)
          / (M.vapor_pressure_table_25() / 101.325), 0.01, 0.05),
   'C–C va jadval farqi: 0,01 %')
hd = M.henry_constant()
ck(close(hd['kh'], 1.242) and close(hd['c'], 1.242, 1e-4)
   and close(hd['m_g_per_l'], 79.56, 0.02) and close(hd['h'], 0.8052, 1e-4),
   'Henry: k_H = 1,242 mol/(L·atm);  c = 1,242 M = 79,6 g/L;  H = 0,805 atm·L/mol')
ck(close(M.gas_density(273.15), 2.8580, 1e-3) and close(M.gas_density(298.15), 2.6184, 1e-3),
   'ρ_gaz: 273,15 K da 2,858 g/L;  298,15 K da 2,618 g/L')
io = M.iodometric()
ck(close(io['n_so3'], 0.0025, 1e-9) and close(io['m_i2'], 0.6345, 1e-4)
   and close(io['m_thio'], 1.2409, 1e-4),
   'yodometriya: n = 0,002500 mol;  I₂ 0,6345 g;  tiosulfat 1,2409 g')
nn = M.neutralisation()
ck(close(nn['naoh_nahs3'], 62.437, 0.01) and close(nn['naoh_na2so3'], 124.874, 0.01)
   and close(nn['caoh2'], 115.662, 0.01) and close(nn['na2s2o5'], 148.377, 0.01),
   'neytralizatsiya: NaOH 62,4 / 124,9 kg;  Ca(OH)₂ 115,7 kg;  Na₂S₂O₅ 148,4 kg')
eq = M.equivalents()
ck(close(eq['h2so3'], 41.04, 1e-6) and close(eq['so2_redox'], 32.03, 1e-6),
   'ekvivalent massa: 41,04 g/eq (kislota);  32,03 g/eq (redoks)')
cc = M.contact()
ck(close(cc['m_so3'], 1249.766, 0.01) and close(cc['m_o2'], 249.75, 0.01)
   and close(cc['m_h2so4'], 1531.065, 0.01),
   'kontakt usuli: SO₃ 1249,8 kg;  O₂ 249,8 kg;  H₂SO₄ 1531,1 kg')
ck(close(abs(cc['balance']), 0.094, 0.01), 'kontakt moddiy balansi: 0,09 kg')
ck(close(M.air_load(5.0)['mg'], 13.1, 0.05) and close(M.air_load(100.0)['mg'], 262.0, 0.1),
   'havo yuklamasi: 5 ppm → 13,1 mg/m³;  100 ppm → 262 mg/m³')

print('3) CSV fayllar va rasmlar')
for name, n in [('kislotalik', 57), ('fizik', 21), ('eritma', 11), ('tuzlar', 8),
                ('ekvivalent', 4), ('yodometriya', 5), ('neytralizatsiya', 5),
                ('kontakt', 6), ('xavfsizlik', 8), ('sanoat', 10)]:
    rows = list(csv.DictReader((ASSETS / (name + '.csv')).open()))
    ck(len(rows) == n, '%s.csv: %d qator' % (name, len(rows)))
pngs = sorted(p.stem for p in ASSETS.glob('*.png'))
ck(len(pngs) == 17, 'rasmlar soni: %d' % len(pngs))
for n in ['kislotalik', 'fizik', 'eritma', 'tuzlar', 'ekvivalent',
          'fizikx', 'eritmax', 'gidratx', 'redoksx', 'tuzlarx', 'tahlilx', 'xavfsizlikx',
          'molekulax', 'hosilax', 'tabiiyx', 'tuzlarilx', 'zavodx']:
    ck(n in pngs, 'rasm %s.png bor' % n)

print('4) Matn: taxminiy hisob raqamlari va atamalar')
ms = MS.read_text()
tables_txt = json.dumps(json.loads((ROOT / 'matn' / '07_02_jadvallar.json').read_text()), ensure_ascii=False)
notes_txt = json.dumps(json.loads((ROOT / 'matn' / '07_02_yonlavhalar.json').read_text()), ensure_ascii=False)
ms_all = ms + tables_txt + notes_txt
need = ['82,08', '64,06', '119,47', '143,10', '1,63', '12,30',
        '263,15', '200,75', '430,80', '78,79', '24,90', '3,85', '0,01',
        '94 g/L', '1,4674', '0,1428', '0,845', '9,7', '1,242', '0,805', '79,6',
        '1,54×10⁻²', '6,3×10⁻⁸', '1,59×10⁻⁷', '1,81', '7,2',
        '39,4', '60,6', '61,3', '38,7', '99,8', '1,49', '32,3',
        '0,002500', '0,6345', '1,2409', '62,4', '124,9', '115,7', '148,4',
        '41,04', '32,03', '1249,8', '249,8', '1531,1', '2,62',
        '+0,17 V', '+0,45 V', '−0,93 V', '0,3', '3', '25', '0,2', '0,75', '30',
        '5 ppm', '2 ppm', '0,25', 'IDLH 100', '0,7 mg/kg', '92–99', '95 %', '75 ppb', '10 ppb']
for x in need:
    ck(x in ms_all, 'matnda «%s» bor' % x)
ck(ms.count('[[EQ:') == 11 and ms.count('[[TABLE:') == 11, 'matnda 11 formula va 11 jadval belgisi')
ck(len(re.findall(r'^## ', ms, flags=re.M)) == 11, 'matnda 11 bo‘lim')
ck('amaliyot tavsiyalari berilmaydi' in ms, 'xavfsizlik chegarasi matnda bayon qilindi')
ck('o‘lchov ma’lumoti emas' in ms, 'taxminiy hisob cheklovi matnda ko‘rsatilgan')
ck('taxminiy hisob' in ms.lower(), 'terminologiya: «taxminiy hisob» ishlatilgan')
ck('model' not in ms_all.lower(), 'terminologiya: «model» yo‘q')
ck('sulfit kislota' in ms and 'sulfit ioni' in ms and 'Clausius–Clapeyron' in ms
   and 'sulfit kislota' in ms.lower(), 'mavzu kalit tushunchalari matnda')
ck('$' not in ms, 'matnda dollar belgisi yo‘q')
ck(not re.search(r'[Ѐ-ӿ]', ms_all), 'matnda kirill harflari yo‘q')
ck(not re.search(r'[　-鿿]', ms_all), 'matnda CJK belgilar yo‘q')

print('5) DOCX tuzilmasi')
doc = Document(DOCX)
xml = doc.element.xml


def collect(container):
    out = []
    for t in container.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
                out.extend(collect(c))
    return out


text_all = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                text_all += '\n' + p.text
            text_all += '\n' + '\n'.join(collect(c))
ck(len(doc.inline_shapes) == 17, 'rasmlar soni: %d' % len(doc.inline_shapes))
import equations_07_02 as EQ
n_eq = sum(len(g) for g in EQ.sulfite_equations())
ck(xml.count('<m:oMath') == n_eq, 'OMML satrlari: %d ( kutilgan %d)' % (xml.count('<m:oMath'), n_eq))
ck(xml.count('<m:f>') >= 2, 'kasrlar: %d' % xml.count('<m:f>'))
ck(xml.count('<m:sSub>') >= 20, 'pastki indekslar: %d' % xml.count('<m:sSub>'))
ck(xml.count('sSup') >= 1, 'yuqori indekslar: %d' % xml.count('sSup'))
ck(text_all.count('-jadval.') == 11, 'jadval sarlavhalari: %d' % text_all.count('-jadval.'))
ck(text_all.count('-sxema.') == 7, 'sxema sarlavhalari: %d' % text_all.count('-sxema.'))
ck(text_all.count('-grafik.') == 5, 'grafik sarlavhalari: %d' % text_all.count('-grafik.'))
ck(text_all.count('-rasm.') == 5, 'rasm sarlavhalari: %d' % text_all.count('-rasm.'))
for i in range(1, 12):
    ck(text_all.count('(7.2.%d)' % i) == 1, 'formula raqami (7.2.%d)' % i)
ck('_' not in text_all, 'matnda "_" qolmadi (pastki indeksga o‘girilgan)')
ck('**' not in text_all, 'matnda "**" markerlari qolmadi')
omml_txt = ''.join(re.findall(r'<m:t[^>]*>([^<]*)</m:t>', xml))
ck('_' not in omml_txt, 'OMML matnida "_" qolmadi (pastki indeks native)')
ck(omml_txt.count('C') > 0 and 'C s' not in omml_txt, 'OMML simvol izohlari toza')
ck('Times New Roman' in xml, 'Times New Roman shrifti qo‘llangan')
ck('Cambria Math' in xml, 'formulalar Cambria Math shriftida')
ck(text_all.count('HISOBIY GRAFIK') == 5, 'grafik yorliqlari: %d' % text_all.count('HISOBIY GRAFIK'))
ck(text_all.count('TUSHUNTIRUVCHI SXEMA') == 3, 'sxema yorliqlari: %d' % text_all.count('TUSHUNTIRUVCHI SXEMA'))
ck(text_all.count('CHIZMA VA IZOH') == 3, 'chizma yorliqlari: %d' % text_all.count('CHIZMA VA IZOH'))
ck('NOORGANIK KISLOTALAR' in doc.sections[0].header.paragraphs[0].text, 'sarlavha matni')
ck('7-BOB' in doc.sections[0].footer.paragraphs[0].text, 'kolontitul matni')
for x in ['82,08', '119,47', '1,242', '0,002500', '148,4', 'IDLH', '41,04']:
    ck(x in text_all, 'DOCX matnida «%s» bor' % x)
ck('OSHA' in text_all and 'NIOSH' in text_all and 'NIST' in text_all,
   'adabiyotlar ro‘yxati joylashdi')
ck('amaliyot tavsiyalari berilmaydi' in text_all, 'xavfsizlik chegarasi bayon qilindi')
size = DOCX.stat().st_size
ck(size > 400000, 'DOCX hajmi: %d bayt' % size)

print()
if fails:
    print('XATO: %d tekshiruv bajarilmadi' % len(fails))
    for f in fails: print('  -', f)
    raise SystemExit(1)
print('BARCHA TEKSHIRUVLAR O‘TDI (7.2).')
