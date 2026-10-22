#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. 4.5 "Parchalanish" moduli."""
from book_equations import add_equation
from pathlib import Path
import re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'chiqish' / 'Noorganik-kislotalar-4.5-Parchalanish.docx'
ASSETS = ROOT / 'rasmlar' / 'parchalanish'
NAVY = '0B2C4A'; BRONZE = 'C28E3B'; PALE = 'EDF3F7'; INK = '243746'

panels = [
 ('Parchalanish turlari', [], '4.5.1-sxema. Termik, fotokimyoviy, katalitik, dismutatsion va kondensatsion yo‘nalishlar.'),
 ('Termodinamik chegara', [], '4.5.2-grafik. ΔG° ning haroratga bog‘liqligi; HNO₃ va H₂SO₄ parchalanishi. Model egri chizig‘i.'),
 ('K va harorat', [], '4.5.3-grafik. Sulfur kislota parchalanishining muvozanat konstantasi; van ’t Hoff chizig‘i.'),
 ('Kinetik blok', [], '4.5.4-grafik. Arrhenius modeli bo‘yicha tezlikning haroratga bog‘liqligi; aniq kislotga tegishli emas.'),
 ('Fotoskimyoviy zanjir', [], '4.5.5-sxema. Boshlanish, davom, tarmoqlanish va tugatish bosqichlari.'),
 ('Dismutatsiya', [], '4.5.6-sxema. Bitta elementning oksidlanish-qaytarilish yo‘nalishlari.'),
 ('Kondensatsiya zanjiri', [], '4.5.7-sxema. Monomerdan oligomer va polimer gacha suv ajratib o‘tish.'),
 ('pH ta’siri', [], '4.5.8-grafik. S(IV) ning uch shakli ulushlari pH ga bog‘liqligi; hisob egri chizig‘i.'),
 ('Saqlash omillari', [], '4.5.9-sxema. Harorat, yorug‘lik, ifloslik va hujjatlashtirish.'),
 ('Tahlil usullari', [], '4.5.10-sxema. Qolgan kislota, rangli mahsulot, uchuvchi mahsulot va tarkib.'),
 ('Hisob natijasi', [], '4.5.11-sxema. Termodinamik ko‘rsatkichlar va kinetik model natijalari.')]

extra = {
 1: [('profil.png', '4.5.1-rasm. Termodinamik qulay, lekin kinetik blok bilan parchalanish profili.')],
 5: [('zanjir.png', '4.5.5-rasm. Radikal zanjirining ko‘payishi va tugatish sxemasi.')],
 7: [('kondzanjir.png', '4.5.7-rasm. Kondensatsiya: monomerdan oligomerga o‘tish sxemasi.')],
 8: [('phshakli.png', '4.5.8-rasm. pH ga qarab sulfit tizimining uch shakli.')],
 9: [('sargayish.png', '4.5.9-rasm. Mahsulot to‘shishining rang belgisi sifatidagi ifodasi.')]}

doc = Document(); sec = doc.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)
sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.5)
sec.header_distance = Cm(.65); sec.footer_distance = Cm(.65)
normal = doc.styles['Normal']; normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.widow_control = True
for sty in doc.styles:
    if sty.type == 1:
        sty.font.name = 'Times New Roman'
        rp = sty.element.get_or_add_rPr(); rf = OxmlElement('w:rFonts')
        for a in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            rf.set(qn('w:' + a), 'Times New Roman')
        rp.append(rf)
sub = dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ', '0123456789abmj'))
sup = dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡', '0123456789+−mn‡'))
TOKEN = re.compile(r'([A-Za-zΑ-Ωα-ωβγδΔ′][A-Za-zΑ-Ωα-ωβγδΔ′₀-₉]*)_([A-Za-z0-9,′]+)')


def runs(p, text, size=11, bold=False, color=INK):
    def put(t, lower=False, upper=False):
        r = p.add_run(t); r.font.name = 'Times New Roman'; r.font.size = Pt(size)
        r.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
        if lower: r.font.subscript = True
        if upper: r.font.superscript = True
    def emit(chunk):
        if not chunk: return
        buf = ''
        for ch in chunk:
            if ch in sub:
                put(buf); buf = ''; put(sub[ch], lower=True)
            elif ch in sup:
                put(buf); buf = ''; put(sup[ch], upper=True)
            else:
                buf += ch
        put(buf)
    pos = 0
    for m in TOKEN.finditer(text):
        emit(text[pos:m.start()])
        put(m.group(1)); put(m.group(2), lower=True)
        pos = m.end()
    emit(text[pos:])


def shade(cell, fill):
    x = OxmlElement('w:shd'); x.set(qn('w:fill'), fill); cell._tc.get_or_add_tcPr().append(x)


def cell_margin(cell, **kwargs):
    x = OxmlElement('w:tcMar')
    for k, v in kwargs.items():
        y = OxmlElement('w:' + k); y.set(qn('w:w'), str(v)); y.set(qn('w:type'), 'dxa'); x.append(y)
    cell._tc.get_or_add_tcPr().append(x)


def border(cell, edge, color, size):
    bs = OxmlElement('w:tcBorders'); b = OxmlElement('w:' + edge)
    b.set(qn('w:val'), 'single'); b.set(qn('w:color'), color); b.set(qn('w:sz'), str(size))
    bs.append(b); cell._tc.get_or_add_tcPr().append(bs)


def band(label, fill, color, size, accent=False):
    t = doc.add_table(rows=1, cols=1); t.autofit = False; t.columns[0].width = Cm(18)
    c = t.cell(0, 0); shade(c, fill); cell_margin(c, top=140, bottom=140, left=200, right=160)
    if accent: border(c, 'left', BRONZE, 24)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
    runs(p, label, size, True, color)
    for p in c.paragraphs: p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = Pt(3)
    p.paragraph_format.keep_with_next = True


comparisons = json.loads((ROOT / 'matn' / '04_05_jadvallar.json').read_text())
comparison_count = 0


def add_comparison_table(cell, spec):
    global comparison_count
    comparison_count += 1
    p = cell.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5); p.paragraph_format.keep_with_next = True
    runs(p, spec['caption'], 9.5, True, NAVY)
    t = cell.add_table(rows=1, cols=len(spec['headers'])); t.autofit = False
    for col, w in zip(t.columns, spec['widths']): col.width = Cm(w)
    for i, label in enumerate(spec['headers']):
        c = t.rows[0].cells[i]; c.width = Cm(spec['widths'][i]); shade(c, NAVY)
        cell_margin(c, top=50, bottom=50, left=85, right=85)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        runs(p, label, 9.5, True, 'FFFFFF')
    rp = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); rp.append(repeat)
    for ri, row in enumerate(spec['rows']):
        r = t.add_row(); trp = r._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
        for ci, value in enumerate(row):
            c = r.cells[ci]; c.width = Cm(spec['widths'][ci])
            cell_margin(c, top=50, bottom=50, left=85, right=85)
            if ri % 2 == 0: shade(c, PALE)
            border(c, 'bottom', 'D6DFE7', 3)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            runs(p, value, 9.5, ci == 0, INK)
    p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(3)


source = (ROOT / 'matn' / '04_05.md').read_text()
sections = []
for raw in source.split('## ')[1:]:
    title, body = raw.split('\n', 1); sections.append((title, body.strip().split('\n\n')))
assert len(sections) == 11
side_notes = json.loads((ROOT / 'matn' / '04_05_yonlavhalar.json').read_text())
assert len(side_notes) == len(sections)
counter = 0
for idx, ((title, paras), (panel, nodes, note)) in enumerate(zip(sections, panels)):
    if idx:
        pb = doc.add_page_break(); pb.paragraph_format.space_after = Pt(0)
        pb.paragraph_format.space_before = Pt(0); pb.paragraph_format.line_spacing = Pt(1)
        for r in pb.runs: r.font.size = Pt(1)
    if idx == 0:
        band('4-BOB  /  UMUMIY KIMYOVIY XOSSALAR', NAVY, 'FFFFFF', 11)
        band('4.5. Parchalanish', PALE, NAVY, 18, True)
    else:
        band('4.5  /  PARCHALANISH', PALE, NAVY, 11, True)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9); p.paragraph_format.keep_with_next = True
    runs(p, title, 14, True, NAVY)
    table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    table.columns[0].width = Cm(13.5); table.columns[1].width = Cm(4.5)
    left, right = table.rows[0].cells
    left.width = Cm(13.5); right.width = Cm(4.5)
    for c in [left, right]: c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left, top=0, bottom=0, left=0, right=220)
    cell_margin(right, top=140, bottom=130, left=160, right=100)
    shade(right, 'F5F8FA'); border(right, 'left', 'D2DCE5', 6)
    first = True
    for text in paras:
        if text.startswith('[[TABLE:'):
            key = text[len('[[TABLE:'):-2]
            add_comparison_table(left, comparisons[key])
            continue
        p = left.paragraphs[0] if first else left.add_paragraph()
        j = 0 if first else 1
        first = False
        if text.startswith('[[EQ:') and text.endswith(']]'):
            text = text[5:-2]
            counter += 1; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(8)
            add_equation(p, left, '4.5', counter, runs)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(.45 if j else 0)
            runs(p, text)
    p = right.paragraphs[0]; p.paragraph_format.space_after = Pt(9)
    runs(p, panel, 10, True, NAVY)
    graphic = {0: 'turlar.png', 4: 'fotokimyo.png', 5: 'dismutatsiya.png',
               6: 'kondensatsiya.png', 8: 'saqlash.png', 9: 'tahlil.png', 10: 'hisob.png'}
    graph = {1: 'dG.png', 2: 'K.png', 3: 'kinetics.png', 7: 'sulfite.png'}
    is_graph = idx in graph
    pic = graphic.get(idx) or graph.get(idx)
    p = right.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / pic), width=Cm(3.92))
    p = right.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    runs(p, 'HISOBIY GRAFIK' if is_graph else 'TUSHUNTIRUVCHI SXEMA', 9, True, NAVY)
    p = right.add_paragraph(); runs(p, note, 9)
    for image, caption in extra.get(idx + 1, []):
        p = right.add_paragraph(); p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ASSETS / image), width=Cm(3.92))
        p = right.add_paragraph(); runs(p, caption + ' Rang, o‘lcham va geometriya shartli.', 9)
    for label, body in side_notes[idx]:
        p = right.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        pp = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr')
        rule = OxmlElement('w:top'); rule.set(qn('w:val'), 'single')
        rule.set(qn('w:sz'), '4'); rule.set(qn('w:color'), BRONZE)
        rule.set(qn('w:space'), '5'); bd.append(rule); pp.append(bd)
        runs(p, label, 9, True, NAVY)
        p = right.add_paragraph()
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True
        runs(p, body, 9)
    if idx == 10:
        p = left.add_paragraph(); p.paragraph_format.space_before = Pt(8); runs(p, 'Adabiyotlar', 10, True, NAVY)
        refs = [
        '[1] OpenStax. Chemistry 2e. Appendix G — Standard Thermodynamic Properties for Selected Substances. https://openstax.org/books/chemistry-2e/pages/g-standard-thermodynamic-properties-for-selected-substances . ΔfH°, ΔfG° va S° qiymatlari (25 °C).',
        '[2] OpenStax. Chemistry 2e. Appendix H — Ionization Constants of Weak Acids. https://openstax.org/books/chemistry-2e/pages/h-ionization-constants-of-weak-acids . H₂SO₃, HNO₂, H₂O₂ va boshqa kislotalarning K_a qiymatlari.',
        '[3] OpenStax. Chemistry 2e. 16.4 Free Energy. https://openstax.org/books/chemistry-2e/pages/16-4-free-energy . ΔG°=ΔH°−TΔS° va uni muvozanat konstantasi bilan bog‘lash.',
        '[4] CRC Handbook of Chemistry and Physics, 97th ed. CRC Press, 2016. — moddalarning qaynab chiqish haroratlari, faza ma’lumotlari va termodinamik xossalari.',
        '[5] IUPAC Gold Book. Disproportionation. DOI: 10.1351/goldbook.D01799. https://goldbook.iupac.org/terms/view/D01799 . Bitta elementning bir qismi oksidlanishi, ikkinchi qismi qaytarilishi; «dismutation» atamasi.',
        '[6] IUPAC Gold Book. Comproportionation. DOI: 10.1351/goldbook.C01217. https://goldbook.iupac.org/terms/view/C01217 . Disproportionationning teskari yo‘nalishi.'
        ]
        for ref in refs:
            p = left.add_paragraph(); p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3); runs(p, ref, 9)
        p = left.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(3)
        runs(p, 'Barcha grafik, sxema va rasmlar original o‘quv ishlanmalari. Sonli egri chiziqlar model hisoblaridir, tajriba yoki o‘lchov ma’lumotlari emas. Rasmlar mikrofoto, atom tuzilmasi yoki masshtabli tuzilma tasviri emas. Moddalar konsentratsiyasi, harorat va pH oralig‘i real eksperiment yoki idish materiali tanlash uchun ko‘rsatma emas; xavfli kislotalarning tayyorlanishi, konsentratsiyalanishi va saqlanishi bo‘yicha amaliyot tavsiyalari berilmaydi.', 9)

p = sec.header.paragraphs[0]; runs(p, 'NOORGANIK KISLOTALAR', 8, False, NAVY)
p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
runs(p, '4-BOB   ·   ', 8, False, NAVY)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); p._p.append(f)
doc.core_properties.title = 'Noorganik kislotalar · 4.5 · Parchalanish'
doc.core_properties.author = ''
OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT)
assert len(doc.inline_shapes) == 16
assert counter == 11
assert comparison_count == 11
assert sum(len(notes) for notes in side_notes) == 22
print(OUT)
print('11 bo‘lim; 11 jadval; 4 hisoblangan grafik; 7 sxema; 5 original illyustratsiya; 11 formula; 22 yon lavha.')
