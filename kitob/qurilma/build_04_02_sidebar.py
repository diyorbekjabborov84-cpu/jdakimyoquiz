#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. Alohida mavzu moduli."""
from book_equations import add_equation
from pathlib import Path
import re, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'chiqish'/'Noorganik-kislotalar-4.2-Oksidlovchilik-va-qaytaruvchilik.docx'
ASSETS=ROOT/'rasmlar'/'redoks_sxema'; ASSETS.mkdir(parents=True,exist_ok=True)
NAVY='0B2C4A'; BRONZE='C28E3B'; PALE='EDF3F7'; INK='243746'
fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
def font(n):return ImageFont.truetype(fontpath,n)

def scheme(index,nodes):
    special={1:'elektron.png',2:'energiya.png',3:'ph.png',4:'balans.png',5:'nitrat.png',6:'sigim.png',7:'fosfor.png',8:'peroksid.png',9:'hisob.png'}
    return ROOT/'rasmlar'/'redoks'/special[index]

panels=[
('Proton va elektron',[],'4.2.1-sxema. Proton ko‘chishi va elektron ko‘chishining boshqa-boshqa vazifalari.'),
('Potensial va energiya',[],'4.2.1-grafik. ΔG=−nFE bog‘lanishi; turli n, bir xil potensial. Nazariy bog‘lanish, tok–kuchlanish tajribasi emas.'),
('pH ning ta’siri',[],'4.2.2-grafik. 25 °C da boshqa nisbiy faolliklar birga teng; shartli yarimreaksiya chiziqlari, to‘liq faza xaritasi emas.'),
('Ion–elektron balansi',[],'4.2.2-sxema. MnO₄⁻/Mn²⁺ va Fe³⁺/Fe²⁺ uchun atom, zaryad va elektron nazorati.'),
('Nitratning ikki vazifasi',[],'4.2.3-sxema. NO₂ va NO nishonlari; kuzatuvchi nitrat hamda havoda gazning o‘zgarishi.'),
('Nishon va sig‘im',[],'4.2.3-grafik. Bir N yoki S markaziga elektron talabi; mahsulotlarning real tanlanishini bashorat qilmaydi.'),
('Tuzilish va ekvivalent',[],'4.2.4-sxema. Fosforli kislotalarda P–H, O–H hamda elektron ekvivalentlarini ajratish.'),
('Ikki rolni farqlang',[],'4.2.4-grafik. H₂O₂/H₂O va O₂/H₂O₂ qaytarilish juftlari, 25 °C; boshqa faolliklar birga teng.'),
('Qoldiq redoks jufti',[],'4.2.5-sxema. Stexiometrik sarfdan qolgan Fe³⁺/Fe²⁺ nisbatiga va SHE ga nisbatan potensialga o‘tish.')
]

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for sty in doc.styles:
    if sty.type==1:
        sty.font.name='Times New Roman'
        rp=sty.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
        rp.append(rf)
sub=dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ','0123456789abmj'));sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡','0123456789+−mn‡'))
variables={
'E°_cell':('E°','cell'),'E_cell':('E','cell'),'E_kat':('E','kat'),'E_an':('E','an'),
'aᵢ':('a','i'),'cᵢ':('c','i'),'E_a':('E','a'),'K_w,c':('K','w,c'),'K_c':('K','c'),'K_w':('K','w')}

def runs(p,text,size=11,bold=False,color=INK):
    pattern='('+'|'.join(re.escape(v) for v in sorted(variables,key=len,reverse=True))+r'|10⁻pH|[₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡]+)'
    def put(t,lower=False,upper=False):
        r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if lower:r.font.subscript=True
        if upper:r.font.superscript=True
    for token in re.split(pattern,text):
        if not token:continue
        if token=='10⁻pH':put('10');put('−pH',upper=True)
        elif token in variables:
            main,low=variables[token];put(main);put(low,lower=True)
        elif all(c in sub for c in token):put(''.join(sub[c] for c in token),lower=True)
        elif all(c in sup for c in token):put(''.join(sup[c] for c in token),upper=True)
        else:put(token)

def shade(cell,fill):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(x)
def cell_margin(cell,**kwargs):
    x=OxmlElement('w:tcMar')
    for k,v in kwargs.items():
        y=OxmlElement('w:'+k);y.set(qn('w:w'),str(v));y.set(qn('w:type'),'dxa');x.append(y)
    cell._tc.get_or_add_tcPr().append(x)
def border(cell,edge,color,size):
    bs=OxmlElement('w:tcBorders');b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:color'),color);b.set(qn('w:sz'),str(size));bs.append(b);cell._tc.get_or_add_tcPr().append(bs)
def band(label,fill,color,size,accent=False):
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(18)
    c=t.cell(0,0);shade(c,fill);cell_margin(c,top=140,bottom=140,left=200,right=160)
    if accent:border(c,'left',BRONZE,24)
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    runs(p,label,size,True,color)
    # Prevent the banner row separating from its following paragraph.
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.keep_with_next=True

comparisons=json.loads((ROOT/'matn'/'04_02_jadvallar.json').read_text())
comparison_count=0

def add_comparison_table(cell,spec):
    global comparison_count
    comparison_count+=1
    p=cell.add_paragraph();p.paragraph_format.space_before=Pt(8)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    runs(p,spec['caption'],9.5,True,NAVY)
    t=cell.add_table(rows=1,cols=len(spec['headers']));t.autofit=False
    for col,w in zip(t.columns,spec['widths']):col.width=Cm(w)
    for i,label in enumerate(spec['headers']):
        c=t.rows[0].cells[i];c.width=Cm(spec['widths'][i]);shade(c,NAVY)
        cell_margin(c,top=70,bottom=70,left=85,right=85)
        p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        runs(p,label,9.5,True,'FFFFFF')
    rp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rp.append(repeat)
    for ri,row in enumerate(spec['rows']):
        r=t.add_row();trp=r._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        for ci,value in enumerate(row):
            c=r.cells[ci];c.width=Cm(spec['widths'][ci])
            cell_margin(c,top=70,bottom=70,left=85,right=85)
            if ri%2==0:shade(c,PALE)
            border(c,'bottom','D6DFE7',3)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            runs(p,value,9.5,ci==0,INK)
    # Word requires a trailing paragraph after a nested table.
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(3)

source=(ROOT/'matn'/'04_02.md').read_text()
sections=[]
for raw in source.split('## ')[1:]:
    title,body=raw.split('\n',1);sections.append((title,body.strip().split('\n\n')))
assert len(sections)==9
side_notes=json.loads((ROOT/'matn'/'04_02_yonlavhalar.json').read_text())
assert len(side_notes)==len(sections)
counter=0
for idx,((title,paras),(panel,nodes,note)) in enumerate(zip(sections,panels)):
    if idx:
        pb=doc.add_page_break();pb.paragraph_format.space_after=Pt(0);pb.paragraph_format.space_before=Pt(0);pb.paragraph_format.line_spacing=Pt(1)
        for r in pb.runs:r.font.size=Pt(1)
    if idx==0:
        band('4-BOB  /  UMUMIY KIMYOVIY XOSSALAR',NAVY,'FFFFFF',11)
        band('4.2. Oksidlovchilik va qaytaruvchilik',PALE,NAVY,18,True)
    else:
        band('4.2  /  OKSIDLOVCHILIK VA QAYTARUVCHILIK',PALE,NAVY,11,True)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=True
    runs(p,title,14,True,NAVY)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    # 18 cm printable width: 13.5 cm narrative + 4.5 cm outer visual strip.
    table.columns[0].width=Cm(13.5);table.columns[1].width=Cm(4.5)
    left,right=table.rows[0].cells
    left.width=Cm(13.5);right.width=Cm(4.5)
    for c in [left,right]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left,top=0,bottom=0,left=0,right=220)
    cell_margin(right,top=140,bottom=130,left=160,right=100)
    shade(right,'F5F8FA');border(right,'left','D2DCE5',6)
    # Allow long content to flow rather than shrink type or clip a fixed-height row.
    first=True
    for text in paras:
        if text.startswith('[[TABLE:'):
            key=text[len('[[TABLE:'):-2]
            add_comparison_table(left,comparisons[key])
            continue
        p=left.paragraphs[0] if first else left.add_paragraph()
        j=0 if first else 1
        first=False
        if text.startswith('[[EQ:') and text.endswith(']]'):
            text=text[5:-2]
            counter+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            add_equation(p,left,'4.2',counter,runs)
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Cm(.45 if j else 0)
            runs(p,text)
    p=right.paragraphs[0];p.paragraph_format.space_after=Pt(9)
    runs(p,panel,10,True,NAVY)
    p=right.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(scheme(idx+1,nodes)),width=Cm(3.2 if idx==6 else 3.92))
    p=right.add_paragraph();p.paragraph_format.space_after=Pt(10)
    runs(p,'HISOBIY GRAFIK' if idx in [1,2,5,7] else 'TUSHUNTIRUVCHI SXEMA',9,True,NAVY)
    p=right.add_paragraph();runs(p,note,9)
    extra_models={
      0:[('nitrat_model.png','4.2.1-rasm. NO₃⁻ ning tekis ulanish modeli; N(+5), O(−2), umumiy zaryad −1.')],
      6:[('fosfit_yon.png','4.2.2-rasm. H₃PO₃ da bitta P–H va ikkita O–H.'),('gipofosfit_yon.png','4.2.3-rasm. H₃PO₂ da ikkita P–H va bitta O–H.')],
      7:[('peroksid_model.png','4.2.4-rasm. H₂O₂: peroksidning O–O ulanishi; O(−1).')]}
    for image,caption in extra_models.get(idx,[]):
        p=right.add_paragraph();p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'rasmlar'/'redoks'/image),width=Cm(3.92))
        p=right.add_paragraph();runs(p,caption+('' if idx==6 else ' Sxematik ulanish; optimallashtirilgan geometriya yoki to‘liq bog‘ tartibi emas.'),9)

    for label,body in side_notes[idx]:
        p=right.add_paragraph()
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        # Delicate bronze rule groups each short aside without heavy boxes.
        pp=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr')
        rule=OxmlElement('w:top');rule.set(qn('w:val'),'single')
        rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),BRONZE)
        rule.set(qn('w:space'),'5');bd.append(rule);pp.append(bd)
        runs(p,label,9,True,NAVY)
        p=right.add_paragraph()
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        runs(p,body,9)
    if idx==8:
        p=left.add_paragraph();p.paragraph_format.space_before=Pt(8);runs(p,'Adabiyotlar',10,True,NAVY)
        refs=[
        '[1] OpenStax. Chemistry 2e. 4.2: Classifying Chemical Reactions. Oksidlanish darajasi va redoks/kislota–asos tasnifi; matn mustaqil bayon qilindi.',
        '[2] OpenStax. Chemistry 2e. 17.3: Electrode and Cell Potentials. https://openstax.org/books/chemistry-2e/pages/17-3-electrode-and-cell-potentials . SHE va qaytarilish potensiallari.',
        '[3] OpenStax. Chemistry 2e. 17.4: Potential, Free Energy, and Equilibrium. https://openstax.org/books/chemistry-2e/pages/17-4-potential-free-energy-and-equilibrium . Nernst, ΔG va K bog‘lanishlari.',
        '[4] OpenStax. Chemistry 2e. Appendix L: Standard Electrode (Half-Cell) Potentials. https://openstax.org/books/chemistry-2e/pages/l-standard-electrode-half-cell-potentials . 25 °C qiymatlari: Cu²⁺/Cu +0,337 V, Fe³⁺/Fe²⁺ +0,771 V, nitrat/NO +0,957 V, permanganat/Mn²⁺ +1,507 V va peroksid juftlari.',
        '[5] Housecroft, C. E.; Sharpe, A. G. Inorganic Chemistry. 5-nashr. Pearson, 2018. N, P va S birikmalariga umumiy bibliografik tayanch; aniq sahifalar tekshirilmagan.',
        '[6] Atkins, P.; de Paula, J.; Keeler, J. Atkins’ Physical Chemistry. 11-nashr. OUP, 2018. Faollik, elektrokimyo va termodinamik potensialga bibliografik tayanch; aniq sahifalar tekshirilmagan.',
        'Grafiklar va sxemalar original hisobiy ishlanma; tajriba emas. Rasmlar ulanish modellari, atom koordinatalari yoki elektron zichlik tasviri emas. Nazariy redoks tenglamalari zaharli gazlarni olish tajribasiga ko‘rsatma emas.'
        ]
        for ref in refs:
            p=left.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3);runs(p,ref,9)

p=sec.header.paragraphs[0];runs(p,'NOORGANIK KISLOTALAR',8,False,NAVY)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runs(p,'4-BOB   ·   ',8,False,NAVY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Noorganik kislotalar · 4.2 · Oksidlovchilik va qaytaruvchilik'
doc.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
assert len(doc.inline_shapes)==13
assert counter==8
assert comparison_count==9
assert sum(len(notes) for notes in side_notes)==18
print(OUT)
print('9 sections; 9 tables; 5 schemes; 4 calculated graphs; 4 molecular images; 8 equations; 18 margin notes.')
