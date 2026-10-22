"""Independent kinetic, thermodynamic, stoichiometric and DOCX checks for 3.6."""
from pathlib import Path
import math,csv,zipfile,json
from docx import Document
from kinetics_03_06 import *
RDIR=Path(__file__).resolve().parents[1];A=RDIR/'rasmlar/kinetika_kataliz'
def near(a,b,rel=2e-11):assert math.isclose(a,b,rel_tol=rel,abs_tol=1e-14),(a,b)
def load(name):return [{k:float(v) for k,v in row.items()} for row in csv.DictReader((A/(name+'.csv')).open())]
counts={}
rows=load('muvozanat');counts['muvozanat']=len(rows)
for r in rows:
 t=r['t_s'];s=r['S_base_fraction'];p=r['P_base_fraction'];sf=r['S_fast_fraction'];pf=r['P_fast_fraction']
 near(s+p,1);near(sf+pf,1);assert 0<=p<=.8 and 0<=pf<=.8
 near(.8-p,.8*math.exp(-.025*t));near(.8-pf,.8*math.exp(-.25*t))
 near(pf,reversible(10*t)[1])
for initial in [0,.8,1]:
 for fac in [1,10]:
  s,p=reversible(10000,factor=fac,p0=initial);near(p,.8);near(.02*s,.005*p)
near(reversible(math.log(2)/.025)[1],.4)
rows=load('energiya');counts['energiya']=len(rows)
for r in rows:
 x=r['coordinate'];near(r['G_uncatalyzed_kJmol'],profile(x));near(r['G_catalyzed_kJmol'],profile(x,True))
near(rows[0]['G_uncatalyzed_kJmol'],rows[0]['G_catalyzed_kJmol']);near(rows[-1]['G_uncatalyzed_kJmol'],-20);near(rows[-1]['G_catalyzed_kJmol'],-20)
near(max(r['G_uncatalyzed_kJmol'] for r in rows),80);near(max(r['G_catalyzed_kJmol'] for r in rows),55)
for x,g in [(.28,55),(.48,10),(.72,45)]:near(profile(x,True),g)
rows=load('harorat');counts['harorat']=len(rows)
for r in rows:
 t=r['T_K'];near(r['k_per_s'],arrhenius(t));near(r['1000K_over_T'],1000/t);near(r['ln_k_over_kref'],-50000/R*(1/t-1/T0))
a,b=rows[0],rows[-1];ea=-R*(math.log(b['k_per_s'])-math.log(a['k_per_s']))/(1/b['T_K']-1/a['T_K']);near(ea,50000)
rows=load('ph_tezlik');counts['ph_tezlik']=len(rows)
for r in rows:
 p=r['pH'];h=10**(-p);near(r['kH_h_per_s'],10*h);near(r['kOH_OH_per_s'],100*KW/h)
 near(r['kobs_per_s'],r['k0_per_s']+r['kH_h_per_s']+r['kOH_OH_per_s']);near(r['log_kobs_over_1s'],math.log10(r['kobs_per_s']))
near(min(rows,key=lambda r:r['kobs_per_s'])['pH'],6.5)
near(ph_channels(6.5)[1],ph_channels(6.5)[2])
rows=load('bufer');counts['bufer']=len(rows)
for r in rows:
 near(r['HA_M']+r['A_M'],r['CT_M']);near(r['HA_M'],r['A_M']);near(r['k_specific_per_s'],.0001101)
 near(r['k_with_general_per_s'],.0001101+.02*r['HA_M'])
near((rows[-1]['k_with_general_per_s']-rows[0]['k_with_general_per_s'])/.1,.01)
assert counts=={'muvozanat':481,'energiya':1001,'harorat':241,'ph_tezlik':401,'bufer':301}
near(eyring_first_order(70000)/eyring_first_order(80000),speedup())
near(speedup(),56.4838385860927);near(-R*T0*math.log(4)/1000,-3.436564151596304)
near(math.log(2)/180,.003850817669777474)
near(math.log(2)/.0002101,3299.129845597074)
near(math.log(2)/.0011101,624.4006671110218)
# Independent atom balance of the displayed formal NO/NO2 redox cycle.
atoms={'NO':{'N':1,'O':1},'O2':{'O':2},'NO2':{'N':1,'O':2},'SO2':{'S':1,'O':2},'SO3':{'S':1,'O':3}}
steps=[{'NO':-2,'O2':-1,'NO2':2},{'NO2':-2,'SO2':-2,'NO':2,'SO3':2}]
for step in steps:
 for element in ['N','S','O']:assert sum(v*atoms[m].get(element,0) for m,v in step.items())==0
net={m:sum(step.get(m,0) for step in steps) for m in atoms};assert net=={'NO':0,'O2':-1,'NO2':0,'SO2':-2,'SO3':2}
p=RDIR/'chiqish/Noorganik-kislotalar-3.6-Kinetika-va-kataliz.docx';d=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(d));txt='\n'.join(p.text for p in ps)
assert len(d.inline_shapes)==13
for label,count in [('jadval.',9),('grafik.',5),('sxema.',4),('rasm.',4)]:assert sum(label in p.text for p in ps)==count,(label,count)
assert all(f'(3.6.{i})' in txt for i in range(1,9))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt and '3.5.' not in txt
assert 'BUFERLAR VA GIDROLIZ' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
assert len(d.element.xpath('//m:sSup'))>=5
assert len(d.element.xpath('//m:f'))>=8
notes=json.loads((RDIR/'matn/03_06_yonlavhalar.json').read_text());assert len(notes)==9 and sum(map(len,notes))==18
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert xml.count('w:tblHeader')==9
 assert 'subscript' in xml and 'superscript' in xml
print('PASS: 2425 model points, mass/equilibrium, Arrhenius slope, Eyring ratio, pH minimum, buffer slope, atom balance and DOCX.')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size)
