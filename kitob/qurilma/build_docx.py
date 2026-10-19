#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
JDA Kitob — DOCX qurilchi (Markdown -> chiroyli Word hujjati).

Manbalar:
    kitob/kitob.json          — kitob metama'lumotlari (sarlavha, muallif, muqova, boblar)
    kitob/boblar/*.md         — boblar (fayl nomi tartibini belgilaydi)

Chiqish:
    kitob/chiqish/<sarlavha>.docx

Ishlatish:
    python kitob/qurilma/build_docx.py            # to'liq kitobni quradi
    python kitob/qurilma/build_docx.py --check     # so'zlar hisoboti (qurmaydi)
    python kitob/qurilma/build_docx.py --bob 03   # faqat bitta bobni quradi
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("python-docx o'rnatilmagan. `pip install python-docx` yoki venvdan foydalaning.")

ROOT = Path(__file__).resolve().parents[1]
BOBLAR = ROOT / "boblar"
CHIQISH = ROOT / "chiqish"
META_FILE = ROOT / "kitob.json"

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)

BODY_PT = 12
TABLE_PT = 10.5

# --------------------------------------------------------------------------- #
# Yordamchilar
# --------------------------------------------------------------------------- #

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*"                       # qalin
    r"|__.+?__"                          # qalin (alt)
    r"|\*[^*\n]+?\*"                     # kursiv
    r"|_[^_\n]+?_"                       # kursiv (alt)
    r"|`[^`\n]+?`"                       # kod
    r"|\$[^$\n]+?\$"                     # formula
    r"|\[[^\]]+\]\([^)\s]+\))"           # havola
)
IMAGE_RE = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)(?:\s+\"([^\"]*)\")?\)\s*$")
FRONT_MATTER_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n?", re.S)


def set_font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def shade(paragraph_or_cell, fill="F2F4F7"):
    el = paragraph_or_cell._element
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def keep_with_next(paragraph, on=True):
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines"):
        el = ppr.find(qn(tag))
        if on and el is None:
            ppr.append(OxmlElement(tag))
        elif not on and el is not None:
            ppr.remove(el)


def add_field(paragraph, instr, size=9, color=GREY):
    run = paragraph.add_run()
    set_font(run, BODY_FONT, size=size, color=color)
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)
    return run


def border_bottom(paragraph, size=6, color="1F3B63"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


# --------------------------------------------------------------------------- #
# Stillar
# --------------------------------------------------------------------------- #

def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)

    heads = {
        "Heading 1": (16, True, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 18),
        "Heading 2": (14, True, False, WD_ALIGN_PARAGRAPH.LEFT, 16, 8),
        "Heading 3": (12.5, True, True, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
        "Heading 4": (12, True, False, WD_ALIGN_PARAGRAPH.LEFT, 10, 4),
        "Heading 5": (12, False, True, WD_ALIGN_PARAGRAPH.LEFT, 8, 4),
    }
    for name, (size, bold, italic, align, before, after) in heads.items():
        st = styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = italic
        st.font.color.rgb = ACCENT if name in ("Heading 1", "Heading 2") else RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), BODY_FONT)
        f = st.paragraph_format
        f.alignment = align
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        f.first_line_indent = Cm(0)
        f.line_spacing = 1.15
        f.keep_with_next = True

    def mk(name, kind=WD_STYLE_TYPE.PARAGRAPH, **kw):
        try:
            st = styles.add_style(name, kind)
        except ValueError:
            st = styles[name]
        if kind == WD_STYLE_TYPE.PARAGRAPH:
            st.font.name = kw.get("font", BODY_FONT)
            st.font.size = Pt(kw.get("size", BODY_PT))
            if kw.get("bold"):
                st.font.bold = True
            if kw.get("italic"):
                st.font.italic = True
            if kw.get("color"):
                st.font.color.rgb = kw["color"]
            f = st.paragraph_format
            f.line_spacing = kw.get("line_spacing", 1.5)
            f.space_before = Pt(kw.get("before", 0))
            f.space_after = Pt(kw.get("after", 0))
            f.first_line_indent = Cm(kw.get("indent", 1.25))
            f.alignment = kw.get("align", WD_ALIGN_PARAGRAPH.JUSTIFY)
        else:
            st.font.name = kw.get("font", BODY_FONT)
            st.font.size = Pt(kw.get("size", BODY_PT))
            if kw.get("bold"):
                st.font.bold = True
            if kw.get("italic"):
                st.font.italic = True
            if kw.get("color"):
                st.font.color.rgb = kw["color"]
        return st

    mk("Kod blok", size=10, font=MONO_FONT, line_spacing=1.05, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    mk("Jadval matni", size=TABLE_PT, line_spacing=1.05, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    mk("Rasm sarlavhasi", size=10.5, italic=True, line_spacing=1.0, indent=0,
       align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=12)
    mk("Jadval sarlavhasi", size=10.5, bold=True, line_spacing=1.0, indent=0,
       align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=2)
    mk("Formula", size=12, italic=True, line_spacing=1.2, indent=0,
       align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6)
    mk("Atama izohi", size=11.5, line_spacing=1.3, indent=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
       before=6, after=6)
    mk("Sarlavha qator", size=10, italic=True, color=GREY, line_spacing=1.0, indent=0,
       align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6)
    mk("Variant", size=12, line_spacing=1.3, indent=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)

    mk("Atama", kind=WD_STYLE_TYPE.CHARACTER, bold=True)
    mk("Kursiv", kind=WD_STYLE_TYPE.CHARACTER, italic=True)
    mk("Kod", kind=WD_STYLE_TYPE.CHARACTER, font=MONO_FONT, size=10.5)


def setup_page(doc, meta):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)
    sec.footer_distance = Cm(1.25)
    sec.header_distance = Cm(1.25)

    # Kolontitul
    footer = sec.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(meta.get("kolontitul", meta.get("sarlavha", "")))
    set_font(run, BODY_FONT, size=9, color=GREY)
    add_field(p, " PAGE ", size=9)

    # Kolontitul sarlavhasi
    header = sec.header
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.text = ""
    hp.paragraph_format.first_line_indent = Cm(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(meta.get("sarlavha", ""))
    set_font(hr, BODY_FONT, size=9, italic=True, color=GREY)
    border_bottom(hp, size=4, color="C8CDD6")


# --------------------------------------------------------------------------- #
# Inline matn
# --------------------------------------------------------------------------- #


GREEK = {
    "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
    "epsilon": "\u03b5", "zeta": "\u03b6", "eta": "\u03b7", "theta": "\u03b8",
    "iota": "\u03b9", "kappa": "\u03ba", "lambda": "\u03bb", "mu": "\u03bc",
    "nu": "\u03bd", "xi": "\u03be", "pi": "\u03c0", "rho": "\u03c1",
    "sigma": "\u03c3", "tau": "\u03c4", "upsilon": "\u03c5", "phi": "\u03c6",
    "varphi": "\u03c6", "chi": "\u03c7", "psi": "\u03c8", "omega": "\u03c9",
    "Delta": "\u0394", "Sigma": "\u03a3", "Omega": "\u03a9", "Phi": "\u03a6",
    "Psi": "\u03a8", "Gamma": "\u0393", "Theta": "\u0398", "Lambda": "\u039b",
    "Pi": "\u03a0",
}

SYMBOLS = {
    "times": "\u00d7", "cdot": "\u00b7", "pm": "\u00b1", "mp": "\u2213",
    "approx": "\u2248", "neq": "\u2260", "leq": "\u2264", "geq": "\u2265",
    "leftarrow": "\u2190", "rightarrow": "\u2192", "to": "\u2192",
    "leftrightarrow": "\u2194", "sum": "\u2211", "prod": "\u220f",
    "int": "\u222b", "infty": "\u221e", "partial": "\u2202", "nabla": "\u2207",
    "degree": "\u00b0", "in": "\u2208", "notin": "\u2209", "subset": "\u2282",
    "cup": "\u222a", "cap": "\u2229", "forall": "\u2200", "exists": "\u2203",
    "prime": "\u2032", "ldots": "\u2026", "dots": "\u2026", "circ": "\u00b0",
    "equiv": "\u2261", "propto": "\u221d", "angle": "\u2220", "perp": "\u22a5",
    "parallel": "\u2225", "therefore": "\u2234", "because": "\u2235",
    "aleph": "\u2135", "hbar": "\u210f", "ell": "\u2113", "Re": "\u211c",
    "Im": "\u2111", "emptyset": "\u2205", "varnothing": "\u2205",
    "quad": " ", "qquad": "  ", " ": " ", ",": ", ", ";": ";",
}


def latex_to_text(s: str) -> str:
    """LaTeX ifodalarini oddiy Unicode matnga aylantiradi."""
    s = s.replace("\\\\", " ")
    for _ in range(4):
        s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1/\2)", s)
        s = re.sub(r"\\sqrt\{([^{}]*)\}", "\u221a(\\1)", s)
    s = re.sub(r"\\(?:text|mathrm|mathbf|mathit|mathcal|operatorname)\{([^{}]*)\}",
               r"\1", s)
    s = s.replace("\\left", "").replace("\\right", "")
    s = re.sub(r"\\([A-Za-z]+)", lambda m: GREEK.get(m.group(1), SYMBOLS.get(m.group(1), m.group(1))), s)
    s = s.replace("\\{", "{").replace("\\}", "}").replace("\\,", " ").replace("\\;", " ")
    s = s.replace("\\", "")
    s = re.sub(r"\{2,}", " ", s)
    return s.strip()


def _match_brace(text: str, start: int) -> int:
    """{ ... } juftlashuvchi qavsning yopilish indeksini topadi."""
    depth = 0
    for k in range(start, len(text)):
        if text[k] == "{":
            depth += 1
        elif text[k] == "}":
            depth -= 1
            if depth == 0:
                return k
    return -1


def render_math(paragraph, content: str, size=None):
    """Formulani Word run'lariga yozadi: _{...} -> subscript, ^{...} -> superscript."""
    i = 0
    while i < len(content):
        ch = content[i]
        if ch in "_^" and i + 1 < len(content):
            if content[i + 1] == "{":
                end = _match_brace(content, i + 1)
                if end == -1:
                    run = paragraph.add_run(content[i:])
                    set_font(run, BODY_FONT, size=size, italic=True)
                    return
                inner = content[i + 2:end]
                i = end + 1
            else:
                inner = content[i + 1]
                i += 2
            run = paragraph.add_run(latex_to_text(inner))
            set_font(run, BODY_FONT, size=size, italic=True)
            run.font.subscript = (ch == "_")
            run.font.superscript = (ch == "^")
            continue
        j = i
        while j < len(content) and content[j] not in "_^":
            j += 1
        run = paragraph.add_run(latex_to_text(content[i:j]))
        set_font(run, BODY_FONT, size=size, italic=True)
        i = j


def add_inline(paragraph, text, base_size=None, base_bold=False, base_italic=False):
    """Markdown inline formatlashni Word run'lariga aylantiradi."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_font(run, BODY_FONT, size=base_size, bold=base_bold or None, italic=base_italic or None)
        tok = m.group(0)
        if tok.startswith("**") or tok.startswith("__"):
            run = paragraph.add_run(tok[2:-2])
            set_font(run, BODY_FONT, size=base_size, bold=True, italic=base_italic or None)
        elif tok.startswith("*") or tok.startswith("_"):
            run = paragraph.add_run(tok[1:-1])
            set_font(run, BODY_FONT, size=base_size, italic=True, bold=base_bold or None)
        elif tok.startswith("`"):
            run = paragraph.add_run(tok[1:-1])
            set_font(run, MONO_FONT, size=(base_size or BODY_PT) - 1.5)
        elif tok.startswith("$"):
            render_math(paragraph, tok.strip("$"), base_size)
        else:
            lm = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", tok)
            if lm:
                add_hyperlink(paragraph, lm.group(2), lm.group(1))
            else:
                run = paragraph.add_run(tok)
                set_font(run, BODY_FONT, size=base_size)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, BODY_FONT, size=base_size, bold=base_bold or None, italic=base_italic or None)


# --------------------------------------------------------------------------- #
# Markdown bloki -> Word
# --------------------------------------------------------------------------- #

@dataclass
class Ctx:
    doc: object
    chapter_no: int | None = None
    eq_no: int = 0
    fig_no: int = 0
    tbl_no: int = 0
    sec2: int = 0
    sec3: int = 0
    toc: list = field(default_factory=list)
    words: int = 0
    seen: set = field(default_factory=set)


def word_count(text):
    return len(re.findall(r"[A-Za-zЀ-ӿ0-9’'-]+", text))


def add_heading(ctx, level, text):
    doc = ctx.doc
    if level == 1:
        ctx.fig_no = 0
        ctx.tbl_no = 0
        ctx.eq_no = 0
        ctx.sec2 = 0
        ctx.sec3 = 0
    if level == 2:
        ctx.sec2 += 1
        ctx.sec3 = 0
        if ctx.chapter_no:
            text = f"{ctx.chapter_no}.{ctx.sec2}. {text}"
    elif level == 3:
        ctx.sec3 += 1
        if ctx.chapter_no:
            text = f"{ctx.chapter_no}.{ctx.sec2}.{ctx.sec3}. {text}"
    elif level == 4 and ctx.chapter_no:
        text = f"{ctx.chapter_no}.{ctx.sec2}.{ctx.sec3}. {text}"

    p = doc.add_paragraph(style=f"Heading {min(level, 5)}")
    if level == 1:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    set_font(run, BODY_FONT,
             size={1: 16, 2: 14, 3: 12.5, 4: 12, 5: 12}[min(level, 5)],
             bold=True, italic=(level == 3))
    ctx.toc.append((level, text))
    return p


def add_paragraph(ctx, text, style=None, size=None, bold=False, italic=False, align=None, indent=None):
    p = ctx.doc.add_paragraph(style=style) if style else ctx.doc.add_paragraph()
    if align is not None:
        p.paragraph_format.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    add_inline(p, text, base_size=size, base_bold=bold, base_italic=italic)
    ctx.words += word_count(text)
    return p


def add_list(ctx, items, ordered, level):
    style = "List Number" if ordered else "List Bullet"
    for it in items:
        p = ctx.doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Cm(0.75 + 0.6 * level)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        text = it
        if text.startswith("[ ] "):
            text = "☐  " + text[4:]
        elif text.startswith("[x] ") or text.startswith("[X] "):
            text = "☑  " + text[4:]
        add_inline(p, text)
        ctx.words += word_count(text)


def add_quote(ctx, lines):
    if lines and lines[0].strip().lower().startswith("[!"):
        m = re.match(r"\[!(\w+)\]\s*(.*)", lines[0].strip())
        label = {"NOTE": "Eslatma", "TIP": "Maslahat", "IMPORTANT": "Muhim",
                 "WARNING": "Diqqat", "EXAMPLE": "Misol"}.get(m.group(1).upper(), m.group(1))
        body = [m.group(2)] + lines[1:]
        p = ctx.doc.add_paragraph(style="Atama izohi")
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.2)
        shade(p, "EEF3FA")
        run = p.add_run(f"{label}. ")
        set_font(run, BODY_FONT, size=11.5, bold=True, color=ACCENT)
        add_inline(p, " ".join(x for x in body if x.strip()))
        ctx.words += word_count(" ".join(body))
        return
    for i, ln in enumerate(lines):
        p = ctx.doc.add_paragraph(style="Atama izohi")
        p.paragraph_format.left_indent = Cm(0.8)
        if i == 0:
            p.add_run("— ")
        add_inline(p, ln)
        ctx.words += word_count(ln)


def add_code(ctx, lines):
    p = ctx.doc.add_paragraph(style="Kod blok")
    shade(p, "F4F5F7")
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(6)
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(ln)
        set_font(run, MONO_FONT, size=10)
    ctx.words += word_count(" ".join(lines))


def add_equation(ctx, text):
    p = ctx.doc.add_paragraph(style="Formula")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    ctx.eq_no += 1
    tag = ""
    m = re.search(r"\\tag\{([^}]*)\}", text)
    if m:
        tag = m.group(1)
        text = text.replace(m.group(0), "").strip()
    if not tag and ctx.chapter_no:
        tag = f"{ctx.chapter_no}.{ctx.eq_no}"
    p.add_run("\t")
    render_math(p, text, 12)
    if tag:
        p.add_run("\t")
        run = p.add_run(f"({tag})")
        set_font(run, BODY_FONT, size=12)
    ctx.words += word_count(text)


def add_table(ctx, rows):
    ncols = len(rows[0])
    t = ctx.doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = ctx.doc.styles["Jadval matni"]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_inline(p, val, base_size=TABLE_PT, base_bold=(ri == 0))
            if ri == 0:
                shade(cell, "E7ECF3")
    ctx.words += word_count(" ".join(" ".join(r) for r in rows))


def add_image(ctx, path, caption, alt, base_dir):
    p = ctx.doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    fp = (base_dir / path)
    if not fp.exists():
        fp = Path(path)
    if not fp.exists():
        run = p.add_run(f"[rasm topilmadi: {path}]")
        set_font(run, BODY_FONT, size=10.5, italic=True, color=RGBColor(0xB0, 0x20, 0x20))
    else:
        p.add_run().add_picture(str(fp), width=Cm(14.5))
    ctx.fig_no += 1
    tag = f"{ctx.chapter_no}.{ctx.fig_no}" if ctx.chapter_no else str(ctx.fig_no)
    cp = ctx.doc.add_paragraph(style="Rasm sarlavhasi")
    run = cp.add_run(f"{tag}-rasm. {caption or alt}")
    set_font(run, BODY_FONT, size=10.5, italic=True)
    ctx.words += word_count(caption or alt)


def render(ctx, text, base_dir):
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Sahifa uzilishi
        if stripped in ("<!-- pagebreak -->", "<!-- pagebreak-->"):
            ctx.doc.add_page_break()
            i += 1
            continue

        # Sarlavhalar
        m = re.match(r"^(#{1,5})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).rstrip("#").strip()
            add_heading(ctx, level, title)
            i += 1
            continue

        # Kod bloki
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code(ctx, buf)
            continue

        # Formula bloki
        if stripped.startswith("$$"):
            if stripped.endswith("$$") and len(stripped) > 4:
                add_equation(ctx, stripped[2:-2].strip())
                i += 1
                continue
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().endswith("$$"):
                buf.append(lines[i])
                i += 1
            if i < len(lines):
                buf.append(lines[i].strip().rstrip("$"))
                i += 1
            add_equation(ctx, " ".join(x.strip() for x in buf if x.strip()))
            continue

        # Rasm
        m = IMAGE_RE.match(stripped)
        if m:
            add_image(ctx, m.group(2), m.group(3) or "", m.group(1), base_dir)
            i += 1
            continue

        # Jadval
        if stripped.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            rows = []
            for r in buf:
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    continue
                rows.append(cells)
            if rows:
                ctx.tbl_no += 1
                tag = f"{ctx.chapter_no}.{ctx.tbl_no}" if ctx.chapter_no else str(ctx.tbl_no)
                cp = ctx.doc.add_paragraph(style="Jadval sarlavhasi")
                run = cp.add_run(f"{tag}-jadval.")
                set_font(run, BODY_FONT, size=10.5, bold=True)
                add_table(ctx, rows)
            continue

        # Ajratuvchi chiziq
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = ctx.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Cm(0)
            border_bottom(p, size=4, color="C8CDD6")
            i += 1
            continue

        # Kvota / eslatma
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_quote(ctx, buf)
            continue

        # Javob variantlari: A) ... B) ...
        m = re.match(r"^(\s*)([A-Ea-e])[.)]\s+(.*)$", raw)
        if m:
            while i < len(lines):
                mm = re.match(r"^(\s*)([A-Ea-e])[.)]\s+(.*)$", lines[i])
                if not mm:
                    break
                p = ctx.doc.add_paragraph(style="Variant")
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.9)
                txt = f"{mm.group(2)})  {mm.group(3).strip()}"
                add_inline(p, txt)
                ctx.words += word_count(txt)
                i += 1
            continue

        # Ro'yxat
        m = re.match(r"^(\s*)([-*+])\s+(.*)$", raw)
        if m:
            level = len(m.group(1)) // 2
            items = []
            while i < len(lines):
                mm = re.match(r"^(\s*)([-*+])\s+(.*)$", lines[i])
                if not mm or len(mm.group(1)) // 2 != level:
                    break
                items.append(mm.group(3).strip())
                i += 1
            add_list(ctx, items, ordered=False, level=level)
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if m:
            level = len(m.group(1)) // 2
            items = []
            while i < len(lines):
                mm = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", lines[i])
                if not mm or len(mm.group(1)) // 2 != level:
                    break
                items.append(mm.group(3).strip())
                i += 1
            add_list(ctx, items, ordered=True, level=level)
            continue

        # Oddiy paragraf (qatorlar birlashtiriladi)
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or ns.startswith(("#", ">", "|", "```", "$$", "!["))
                    or re.match(r"^(\s*)([-*+]|\d+[.)])\s+", nxt)
                    or ns in ("<!-- pagebreak -->", "<!-- pagebreak-->")
                    or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", ns)):
                break
            buf.append(ns)
            i += 1
        add_paragraph(ctx, " ".join(buf))


# --------------------------------------------------------------------------- #
# Muqova / mundarija
# --------------------------------------------------------------------------- #

def build_cover(doc, meta):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    if meta.get("seriya"):
        run = p.add_run(meta["seriya"].upper())
        set_font(run, BODY_FONT, size=12, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(meta.get("sarlavha", ""))
    set_font(run, BODY_FONT, size=26, bold=True)
    if meta.get("tagline"):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(meta["tagline"])
        set_font(run, BODY_FONT, size=14, italic=True, color=GREY)
    if meta.get("muqova_tavsifi"):
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(2.5)
        p.paragraph_format.right_indent = Cm(2.5)
        p.paragraph_format.space_before = Pt(24)
        add_inline(p, meta["muqova_tavsifi"], base_size=12, base_italic=True)
    for _ in range(3):
        doc.add_paragraph()
    for line in [meta.get("muallif"), meta.get("muallif_izoh"), meta.get("nashr"), meta.get("yil")]:
        if not line:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line)
        set_font(run, BODY_FONT, size=12.5)


def build_toc(doc, entries, meta):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("MUNDARIJA")
    set_font(run, BODY_FONT, size=16, bold=True, color=ACCENT)

    for level, text in entries:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        if level == 1:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.left_indent = Cm(0)
            run = p.add_run(text)
            set_font(run, BODY_FONT, size=12.5, bold=True)
        elif level == 2:
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(text)
            set_font(run, BODY_FONT, size=12)
        else:
            p.paragraph_format.left_indent = Cm(1.6)
            run = p.add_run(text)
            set_font(run, BODY_FONT, size=11.5, color=GREY)


# --------------------------------------------------------------------------- #
# Asosiy qurilish
# --------------------------------------------------------------------------- #

def read_front_matter(text):
    m = FRONT_MATTER_RE.match(text)
    if not m:
        return {}, text
    meta = {}
    for line in m.group(1).split("\n"):
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.strip().strip('"').strip("'")
    return meta, text[m.end():]


def collect_sources(meta):
    files = []
    for f in sorted(BOBLAR.glob("*.md")):
        if f.name.startswith("_"):
            continue
        files.append(f)
    return files


def build(meta, out_path: Path, only=None):
    doc = Document()
    configure_styles(doc)
    setup_page(doc, meta)

    props = doc.core_properties
    props.title = meta.get("sarlavha", "")
    props.author = meta.get("muallif", "")
    props.subject = meta.get("tagline", "")
    props.language = "uz-UZ"
    props.comments = meta.get("muqova_tavsifi", "")[:255]

    ctx = Ctx(doc=doc)

    # Muqova
    build_cover(doc, meta)

    # Mundarija (ikkinchi marta quriladi, chunki sarlavhalar boblardan keladi)
    toc_entries = []
    for f in collect_sources(meta):
        if only and only not in f.stem:
            continue
        fmeta, body = read_front_matter(f.read_text(encoding="utf-8"))
        title = fmeta.get("title", "")
        number = fmeta.get("number")
        if fmeta.get("kind") == "front":
            toc_entries.append((1, title or f.stem))
            continue
        toc_entries.append((1, f"{number}-bob. {title}" if number else (title or f.stem)))
        for m in re.finditer(r"^(#{2,3})\s+(.+)$", body, re.M):
            lvl = len(m.group(1))
            t = m.group(2).strip()
            if lvl == 2 and number:
                t = f"{number}.{toc_entries.count((1, t)) + 1}. {t}"  # taxminiy
            toc_entries.append((lvl, t))

    # Haqiqiy raqamlashni qurishda hal qilamiz
    toc_entries = []
    for f in collect_sources(meta):
        if only and only not in f.stem:
            continue
        fmeta, body = read_front_matter(f.read_text(encoding="utf-8"))
        title = fmeta.get("title", "")
        number = fmeta.get("number")
        if fmeta.get("kind") == "front":
            toc_entries.append((1, title or f.stem))
            continue
        toc_entries.append((1, f"{number}-bob. {title}" if number else (title or f.stem)))
        c2 = 0
        c3 = 0
        for m in re.finditer(r"^(#{2,3})\s+(.+)$", body, re.M):
            lvl = len(m.group(1))
            t = m.group(2).strip()
            if lvl == 2:
                c2 += 1
                c3 = 0
                if number:
                    t = f"{number}.{c2}. {t}"
            else:
                c3 += 1
                if number:
                    t = f"{number}.{c2}.{c3}. {t}"
            toc_entries.append((lvl, t))

    build_toc(doc, toc_entries, meta)

    stats = []
    for f in collect_sources(meta):
        if only and only not in f.stem:
            continue
        fmeta, body = read_front_matter(f.read_text(encoding="utf-8"))
        before = ctx.words
        ctx.chapter_no = fmeta.get("number")
        if ctx.chapter_no is not None:
            try:
                ctx.chapter_no = int(ctx.chapter_no)
            except ValueError:
                pass
        ctx.sec2 = 0
        ctx.sec3 = 0
        ctx.fig_no = 0
        ctx.tbl_no = 0
        ctx.eq_no = 0
        kind = fmeta.get("kind", "chapter")
        if kind == "chapter":
            title = fmeta.get("title", f.stem)
            num = ctx.chapter_no
            add_heading(ctx, 1, f"{num}-bob. {title}" if num else title)
        render(ctx, body, f.parent)
        stats.append((f.name, ctx.words - before))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return stats, sum(w for _, w in stats), len(toc_entries)


def main():
    ap = argparse.ArgumentParser(description="Kitobni DOCX formatida qurish")
    ap.add_argument("--check", action="store_true", help="faqat so'zlar hisoboti")
    ap.add_argument("--bob", help="faqat ko'rsatilgan faylni qurish (masalan 03)")
    args = ap.parse_args()

    if not META_FILE.exists():
        sys.exit(f"Metama'lumot fayli topilmadi: {META_FILE}")
    meta = json.loads(META_FILE.read_text(encoding="utf-8"))

    if args.check:
        total = 0
        for f in collect_sources(meta):
            fmeta, body = read_front_matter(f.read_text(encoding="utf-8"))
            w = word_count(body)
            total += w
            print(f"{f.name:40s} {w:7d} so'z  | {fmeta.get('title','')}")
        print("-" * 70)
        print(f"{'JAMI':40s} {total:7d} so'z")
        return

    out = CHIQISH / (re.sub(r'[\\/:*?"<>|]', "", meta.get("sarlavha", "kitob")) + ".docx")
    stats, total, toc = build(meta, out, args.bob)
    print(f"Tayyor: {out}")
    print(f"Mundarija bandlari: {toc} | So'zlar: {total}")
    for name, w in stats:
        print(f"  - {name}: {w} so'z")


if __name__ == "__main__":
    main()
