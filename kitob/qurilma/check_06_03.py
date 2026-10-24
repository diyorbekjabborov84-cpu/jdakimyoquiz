#!/usr/bin/env python3
"""6.3 moduli uchun QA: taxminiy hisoblarning matn va hujjat bilan mosligi."""
from pathlib import Path
import csv, math, re, json
from docx import Document
import prep_06_03 as M

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'chiqish' / 'Noorganik-kislotalar-6.3-HCl-xlorid-kislota.docx'
ASSETS = ROOT / 'rasmlar' / 'hcl'
MS = ROOT / 'matn' / '06_03.md'
fails = []


def ck(cond, msg):
    if cond:
        print('  OK  ', msg)
    else:
        print('  FAIL', msg); fails.append(msg)


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol


print('1) Taxminiy hisob: spetsiatsiya')
s = M.speciation_series()
ck(close(s[0]['ph'], 1.000, 1e-3) and close(s[0]['w_cl'], 100.0, 0.01),
   '0,100 M: Cl⁻ 100,00 %; pH 1,000')
ck(close(s[2]['ph'], 0.000, 1e-3) and close(s[2]['w_cl'], 100.0, 0.01),
   '1,000 M: Cl⁻ 100,00 %; pH 0,000')
ck(close(s[8]['ph'], -1.079, 1e-3) and close(s[8]['w_cl'], 99.82, 0.01)
   and close(s[8]['w_cl2'], 0.09, 0.01), '12,000 M: Cl⁻ 99,82 %; HCl₂⁻ 0,09 %; pH −1,079')
ck(close(M.KA_HCL, 7.9e5, 1e4) and close(M.K_CL2, 10 ** 0.7, 1e-6),
   'Ka = 7,9×10⁵; lgK(HCl₂⁻) = 0,7')

print('2) Taxminiy hisob: ishlab chiqarish, metallar, chiqindi, tiklash, standartlash')
mn = M.mannheim(1000.0, 0.98)
ck(close(mn['m_nacl'], 1602.8, 0.2) and close(mn['m_h2so4'], 1345.0, 0.2)
   and close(mn['m_na2so4'], 1947.8, 0.2), 'Mannheim: NaCl 1602,8 kg; H₂SO₄ 1345,0 kg; Na₂SO₄ 1947,8 kg')
ck(close(mn['dh_mj'], 884.8, 0.5) and close(mn['dh_kj_per_mol'], 32.26, 1e-2),
   'Mannheim: ΔrH 884,8 MJ (32,3 kJ/mol HCl)')
sy = M.synthesis(1000.0)
ck(close(sy['m_h2'], 27.6, 0.1) and close(sy['m_cl2'], 972.4, 0.2)
   and close(sy['dh_mj'], 2531.7, 1.0), 'sintetik: H₂ 27,6 kg; Cl₂ 972,4 kg; ΔrH −2531,7 MJ')
md = {x['metal']: x for x in M.metal_dissolution(1.0)}
ck(close(md['Fe']['m_salt'], 2.270, 1e-3) and close(md['Fe']['m_hcl'], 1.306, 1e-3)
   and close(md['Fe']['m_h2'], 0.036, 1e-3), '1 g Fe → 2,270 g FeCl₂; HCl 1,306 g; H₂ 0,036 g')
ck(close(md['Al']['m_salt'], 4.942, 1e-3) and close(md['Al']['m_hcl'], 4.054, 1e-3),
   '1 g Al → 4,942 g AlCl₃; HCl 4,054 g')
ck(close(md['Zn']['m_hcl'], 1.115, 1e-3) and close(md['Mg']['m_hcl'], 3.000, 1e-3),
   '1 g Zn → HCl 1,115 g; 1 g Mg → HCl 3,000 g')
ck(close(md['Cu']['m_hcl'], 1.148, 1e-3) and close(md['Cu']['m_o2'], 0.252, 1e-3)
   and md['Cu']['need_o2'], '1 g Cu → HCl 1,148 g + O₂ 0,252 g')
n1 = M.neutralisation(1000.0, 10.0, 'Na2CO3')
ck(close(n1['m_agent'], 145.3, 0.1) and close(n1['m_salt'], 160.3, 0.1)
   and close(n1['m_gas'], 60.4, 0.1) and close(n1['m_h2o'], 24.7, 0.1),
   'neytralizatsiya: Na₂CO₃ 145,3 kg; NaCl 160,3 kg; CO₂ 60,4 kg; H₂O 24,7 kg')
ck(close(n1['m_input'], n1['m_output'], 0.1), 'neytralizatsiya mass balansi: 245,3 = 245,3 kg')
n2 = M.neutralisation(1000.0, 10.0, 'Ca(OH)2')
ck(close(n2['m_agent'], 101.6, 0.1) and close(n2['m_salt'], 152.2, 0.1)
   and close(n2['m_h2o'], 49.4, 0.1) and close(n2['m_input'], n2['m_output'], 0.1),
   'Ca(OH)₂ bilan: 101,6 / 152,2 / 49,4 kg; kirish = chiqish 201,6 kg')
rg = M.regeneration(1000.0)
ck(close(rg['m_fe2o3'], 629.9, 0.2) and close(rg['m_hcl'], 575.3, 0.2)
   and close(rg['m_h2o'], 142.1, 0.2) and close(rg['m_o2'], 63.1, 0.2),
   'regeneratsiya: Fe₂O₃ 629,9 kg; HCl 575,3 kg; H₂O 142,1 kg; O₂ 63,1 kg')
ck(close(rg['m_input'], rg['m_output'], 0.2), 'tiklash mass balansi: 1205,2 ≈ 1205,2 kg')
pk = M.pickling_balance(1.0)
ck(close(pk['m_hcl'], 7.6, 1e-6) and close(pk['m_fe'], 5.8, 1e-6)
   and close(pk['m_cl'], 7.4, 1e-6) and close(pk['m_fecl2'], 13.2, 0.1),
   'pickling: HCl 7,6 kg; Fe 5,8 kg; Cl⁻ 7,4 kg; FeCl₂ 13,2 kg (1 t po‘lat)')
t = M.titration_standardisation(0.2500, 24.85)
ck(close(t['c_hcl'], 0.18984, 1e-5), 'standartlash: 0,2500 g Na₂CO₃; 24,85 mL → 0,18984 mol/L')
ck(close(M.PHYS['tb'], -85.05) and close(M.PHYS['tm'], -114.22) and close(M.PHYS['tc'], 51.4)
   and close(M.PHYS['pc'], 83.1) and close(M.PHYS['aze_pct'], 20.2)
   and close(M.PHYS['aze_tb'], 108.6) and close(M.PHYS['sol'], 720.0),
   'fizik konstantalar: −85,05 / −114,22 °C; 51,4 °C; 83,1 bar; azeotrop 20,2 % / 108,6 °C')

print('3) CSV fayllar va rasmlar')
for name, n in [('fizik_hcl', 16), ('eritma', 7), ('spetsiatsiya', 9),
                ('ishlabchiqarish', 12), ('metallar', 5), ('chiqindi', 13)]:
    rows = list(csv.DictReader((ASSETS / (name + '.csv')).open()))
    ck(len(rows) == n, '%s.csv: %d qator' % (name, len(rows)))
pngs = sorted(p.stem for p in ASSETS.glob('*.png'))
ck(len(pngs) == 17, 'rasmlar soni: %d' % len(pngs))
for n in ['molekulax', 'fizikx', 'spetsiatsiya', 'olishsx', 'kloridlarsx', 'metalsx',
          'korroziyax', 'suvsizsx', 'qollanishsx', 'tahlilsx', 'hisob',
          'eritma', 'ishlabchiqarish', 'metallar', 'teshikx', 'idishx', 'chiqindi']:
    ck(n in pngs, 'rasm %s.png bor' % n)

print('4) Matn: taxminiy hisob raqamlari va atamalar')
ms = MS.read_text()
tables_txt = json.dumps(json.loads((ROOT / 'matn' / '06_03_jadvallar.json').read_text()), ensure_ascii=False)
notes_txt = json.dumps(json.loads((ROOT / 'matn' / '06_03_yonlavhalar.json').read_text()), ensure_ascii=False)
ms_all = ms + tables_txt + notes_txt
need = ['1,2746', '432 kJ/mol', '1,08 D', '−85,1', '−114,2', '51,4', '83,1', '450 kg/m³',
        '20,2', '108,6', '720 g/L', '1,627', '1,27', '3,16', '68 %', '51 %', '41 %',
        '−5,9', '7,9×10⁵', '1,048', '1,098', '1,149', '1,189', '−1,1', '48 °C',
        '1,000', '0,000', '−1,079', '0,09', '99,82', '0,7',
        '1602,8', '1345,0', '1947,8', '884,8', '32,3', '27,6', '972,4', '2531,7',
        '2,270', '1,306', '0,036', '4,942', '4,054', '1,115', '3,000', '1,148', '0,252',
        '2,084', '3,917', '145,3', '160,3', '60,4', '24,7', '245,3', '101,6', '152,2',
        '49,4', '201,6', '629,9', '575,3', '142,1', '63,1', '1205,2',
        '7,6 kg', '5,8 kg', '7,4 kg', '13,2', '0,2500', '24,85', '0,18984',
        '−11,3', '−12,7', '−18', '−15,1', '5 ppm', '2,98 ppm', '50 ppm', '45 ppm',
        '20 Mt/y', '3 Mt/y', '18 %', '11–13', '2–4', '1,27 bar', '20,2 % da eng katta']
for x in need:
    ck(x in ms_all, 'matnda «%s» bor' % x)
ck(ms.count('[[EQ:') == 11 and ms.count('[[TABLE:') == 11, 'matnda 11 formula va 11 jadval belgisi')
ck('amaliyot tavsiyalari berilmaydi' in ms, 'xavfsizlik chegarasi matnda bayon qilindi')
ck('o‘lchov ma’lumoti emas' in ms, 'taxminiy hisob cheklovi matnda ko‘rsatilgan')
ck('2NaCl + H₂SO₄' in ms and 'H₂ + Cl₂' in ms, 'asosiy tengliklar matnda')
ck('4FeCl₂ + 4H₂O + O₂' in ms, 'pyrohidroliz tengligi matnda')
ck('vodorod bog‘i' in ms and 'passivatsiya' in ms and 'embrittlement' in ms,
   'mavzu kalit tushunchalari matnda')
ck('vodorod bog‘i' in ms and 'hydrogen bog‘i' not in ms.lower(), 'terminologiya: vodorod bog‘i')
ck('$' not in ms, 'matnda dollar belgisi yo‘q')

print('5) DOCX tuzilmasi')
doc = Document(DOCX)
xml = doc.element.xml
text_all = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                text_all += '\n' + p.text
ck(len(doc.inline_shapes) == 17, 'rasmlar soni: %d' % len(doc.inline_shapes))
ck(xml.count('<m:oMath') == 50, 'OMML satrlari (11 guruh, 50 qator): %d' % xml.count('<m:oMath'))
ck(xml.count('<m:f>') >= 1, 'kasrlar: %d' % xml.count('<m:f>'))
ck(xml.count('<m:sSub>') >= 30, 'pastki indekslar: %d' % xml.count('<m:sSub>'))
ck(xml.count('sSup') >= 1, 'yuqori indekslar: %d' % xml.count('sSup'))
ck(text_all.count('-jadval.') == 11, 'jadval sarlavhalari: %d' % text_all.count('-jadval.'))
ck(text_all.count('-sxema.') == 7, 'sxema sarlavhalari: %d' % text_all.count('-sxema.'))
ck(text_all.count('-grafik.') == 5, 'grafik sarlavhalari: %d' % text_all.count('-grafik.'))
ck(text_all.count('-rasm.') == 5, 'rasm sarlavhalari: %d' % text_all.count('-rasm.'))
for i in range(1, 12):
    ck(text_all.count('(6.3.%d)' % i) == 1, 'formula raqami (6.3.%d)' % i)
ck('_' not in text_all, 'matnda "_" qolmadi (pastki indeksga o‘girilgan)')
ck('Times New Roman' in xml, 'Times New Roman shrifti qo‘llangan')
ck('Cambria Math' in xml, 'formulalar Cambria Math shriftida')
ck(text_all.count('HISOBIY GRAFIK') == 1, 'grafik yorliqlari: %d' % text_all.count('HISOBIY GRAFIK'))
ck(text_all.count('TUSHUNTIRIVCHI SXEMA') == 0, 'typo yo‘q')
ck(text_all.count('TUSHUNTIRUVCHI SXEMA') == 7, 'sxema yorliqlari: %d' % text_all.count('TUSHUNTIRUVCHI SXEMA'))
ck(text_all.count('CHIZMA VA IZOH') == 3, 'chizma yorliqlari: %d' % text_all.count('CHIZMA VA IZOH'))
ck('NOORGANIK KISLOTALAR' in doc.sections[0].header.paragraphs[0].text, 'sarlavha matni')
ck('6-BOB' in doc.sections[0].footer.paragraphs[0].text, 'kolontitul matni')
for x in ['1602,8 kg', '884,8 MJ', '1,306 g', '0,18984 mol/L', '575,3 kg', '−15,1',
          '7,9×10', '245,3 kg', 'IDLH']:
    ck(x in text_all, 'DOCX matnida «%s» bor' % x)
ck('OSHA' in text_all and 'NIOSH' in text_all and 'JACS' in text_all,
   'adabiyotlar ro‘yxati joylashdi')
ck('amaliyot tavsiyalari berilmaydi' in text_all, 'xavfsizlik chegarasi bayon qilindi')
size = DOCX.stat().st_size
ck(size > 400000, 'DOCX hajmi: %d bayt' % size)

print()
if fails:
    print('XATO: %d tekshiruv bajarilmadi' % len(fails))
    for f in fails: print('  -', f)
    raise SystemExit(1)
print('BARCHA TEKSHIRUVLAR O‘TDI (6.3).')
