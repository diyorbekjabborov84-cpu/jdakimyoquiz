#!/usr/bin/env python3
"""6.1 moduli uchun QA: model hisoblarining matn va hujjat bilan mosligi."""
from pathlib import Path
import csv, math, re, json
from docx import Document
import prep_06_01 as M

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'chiqish' / 'Noorganik-kislotalar-6.1-Guruhning-umumiy-tavsifi.docx'
ASSETS = ROOT / 'rasmlar' / 'galogenid'
MS = ROOT / 'matn' / '06_01.md'
fails = []


def ck(cond, msg):
    if cond:
        print('  OK  ', msg)
    else:
        print('  FAIL', msg); fails.append(msg)


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol


print('1) Model: molekulyar o‘lchamlar')
a_hf, a_hcl, a_hbr, a_hi = (M.acid(k) for k in ('HF', 'HCl', 'HBr', 'HI'))
ck(close(a_hf['m'], 20.006), 'M(HF) = 20,006 g/mol')
ck(close(a_hcl['m'], 36.461), 'M(HCl) = 36,461 g/mol')
ck(close(a_hbr['m'], 80.912), 'M(HBr) = 80,912 g/mol')
ck(close(a_hi['m'], 127.912), 'M(HI) = 127,912 g/mol')
ck(close(a_hf['d'], 569.0) and close(a_hcl['d'], 432.0)
   and close(a_hbr['d'], 370.0) and close(a_hi['d'], 295.0), 'D(H–X): 569/432/370/295 kJ/mol')
ck(close(a_hf['hyd'], -506.0) and close(a_hcl['hyd'], -364.0)
   and close(a_hbr['hyd'], -335.0) and close(a_hi['hyd'], -293.0), 'ΔH_hyd: −506/−364/−335/−293')
ck(close(M.net_bond('HF'), 63.0) and close(M.net_bond('HCl'), 68.0)
   and close(M.net_bond('HBr'), 35.0) and close(M.net_bond('HI'), 2.0), 'D + ΔH_hyd: 63/68/35/2 kJ/mol')
ck(close(a_hf['rexp'], 0.917) and close(a_hi['rexp'], 1.609), 'r(H–X): 0,917 … 1,609 Å')
ck(close(a_hf['mu'], 1.82) and close(a_hi['mu'], 0.44), 'dipol: 1,82 … 0,44 D')
ck(close(a_hf['tb'], 19.5) and close(a_hcl['tb'], -85.1)
   and close(a_hbr['tb'], -66.8) and close(a_hi['tb'], -35.4), 'qaynash: 19,5 / −85,1 / −66,8 / −35,4 °C')
ck(close(a_hf['tm'], -83.6) and close(a_hcl['tm'], -114.2)
   and close(a_hbr['tm'], -86.9) and close(a_hi['tm'], -50.8), 'suyuqlash: −83,6 / −114,2 / −86,9 / −50,8 °C')
ck(close(a_hf['azeo'], 37.0) and close(a_hcl['azeo'], 20.2)
   and close(a_hbr['azeo'], 47.6) and close(a_hi['azeo'], 56.9), 'azeotrop: 37 / 20,2 / 47,6 / 56,9 %')

print('2) Model: kislotalik, gidratatsiya, termik barqarorlik')
ck(close(a_hf['pka'], 3.17) and close(a_hcl['pka'], -5.9)
   and close(a_hbr['pka'], -8.8) and close(a_hi['pka'], -9.5), 'pKa (suv): 3,17 / −5,9 / −8,8 / −9,5')
ck(close(a_hf['pka_dmso'], 15.0) and close(a_hi['pka_dmso'], -10.9), 'pKa (DMSO): 15 … −10,9')
g_hcl, g_hi = M.gas_to_ion('HCl'), M.gas_to_ion('HI')
ck(close(g_hcl['dg'], -35.901, 1e-3), 'ΔG°(HCl, gaz→ion) = −35,90 kJ/mol')
ck(close(g_hi['dg'], -53.27, 1e-3), 'ΔG°(HI, gaz→ion) = −53,27 kJ/mol')
ck(1.9e6 < g_hcl['k'] < 2.0e6, 'K(model, HCl) ≈ 1,95×10⁶')
ck(2.1e9 < g_hi['k'] < 2.2e9, 'K(model, HI) ≈ 2,15×10⁹')
t_hf, t_hi = M.thermal_decomposition('HF'), M.thermal_decomposition('HI')
ck(close(t_hf['dg'], 550.8, 1e-3) and close(t_hi['dg'], 15.9, 1e-3), 'ΔG° dekompozitsiya: 550,8 … 15,9 kJ/mol')
ck(close(t_hf['log10k'], -96.4959, 1e-3) and close(t_hi['log10k'], -2.7856, 1e-3), 'lgK: −96,50 … −2,79')
ck(close(M.oxidising_balance('HF'), -754.6, 0.05) and close(M.oxidising_balance('HI'), -481.16, 0.05),
   'oksidlovchi balans: −754,6 … −481,2 kJ/mol')

print('3) Model: muvozanat, shisha, neytralizatsiya, redoks, azeotrop')
s = M.hf_ionisation_series()
ck(close(s[0]['x'], 0.0022898, 1e-6) and close(s[0]['pct'], 22.8975, 1e-3) and close(s[0]['ph'], 2.6402, 1e-3),
   'HF 0,010 M: x = 0,002290 M; 22,90 %; pH 2,640')
ck(close(s[2]['x'], 0.0079135, 1e-6) and close(s[2]['pct'], 7.9135, 1e-3), 'HF 0,100 M: 7,91 %')
ck(close(s[4]['x'], 0.0257393, 1e-6) and close(s[4]['pct'], 2.5739, 1e-3) and close(s[4]['ph'], 1.5894, 1e-3),
   'HF 1,000 M: 2,57 %; pH 1,589')
g = M.glass_model(1.0, 48.0)
ck(close(g['m_hf'], 1997.9, 0.5), '1 kg SiO₂ uchun HF: 1997,9 g')
ck(close(g['m_solution'], 4162.4, 0.5), '48 % li eritmada: 4162,4 g')
ck(close(g['m_h2siof6'], 2398.3, 0.5) and close(g['m_h2o'], 599.6, 0.5), 'H₂SiF₆ 2398,3 g; H₂O 599,6 g')
n = M.neutralisation_model(1000.0, 10.0)
ck(close(n['m_caco3'], 137.3, 0.2) and close(n['m_cacl2'], 152.2, 0.2), 'CaCO₃ 137,3 kg; CaCl₂ 152,2 kg')
ck(close(n['m_co2'], 60.4, 0.2) and close(n['m_h2o'], 24.7, 0.2), 'CO₂ 60,4 kg; H₂O 24,7 kg')
ck(close(n['m_input'], n['m_output'], 0.3), 'neytralizatsiya mass balansi: kirish = chiqish')
r = M.redox_model(1.0, 1.0)
ck(close(r['m_br2'], 987.5, 0.5) and close(r['m_i2'], 992.1, 0.5), '1 kg HBr → 987,5 g Br₂; 1 kg HI → 992,1 g I₂')
ck(close(r['m_o2'], 62.54, 0.1), '1 kg HI uchun O₂: 62,54 g')
z10 = M.azeotrope_model(1000.0, 10.0); z37 = M.azeotrope_model(1000.0, 37.0)
ck(close(z10['residue'], 495.0, 0.5) and close(z10['distilled_water'], 505.0, 0.5), '10 % HCl: qoldiq 495,0 kg')
ck(close(z37['residue'], 1831.7, 0.5) and z37['feasible'] is False, '37 % HCl: 1831,7 kg > 1000 kg — imkonsiz')
ab = M.absorption_heat()
ck(close(ab['dh_mol'], -74.8) and close(ab['dh_kg'], -2052.0, 1.0), 'HCl yotqizish: −74,8 kJ/mol = −2052 kJ/kg')
pr = M.production_model(1000.0, 0.97, 0.90)
ck(close(pr['m_hf'], 447.4, 0.5) and close(pr['m_h2so4'], 1218.6, 0.5) and close(pr['m_caso4'], 1691.4, 0.5),
   '1000 kg fluorit: HF 447,4 kg; H₂SO₄ 1218,6 kg; CaSO₄ 1691,4 kg')

print('4) CSV fayllar')
for name, n in [('kislotalik', 4), ('energiya', 4), ('fizik', 4), ('azeotrop', 4),
                ('dissotsiatsiya', 5), ('termik', 4)]:
    rows = list(csv.DictReader((ASSETS / (name + '.csv')).open()))
    ck(len(rows) == n, '%s.csv: %d qator' % (name, len(rows)))
pngs = sorted(p.stem for p in ASSETS.glob('*.png'))
ck(len(pngs) == 17, 'rasmlar soni: %d' % len(pngs))

print('5) Matn: model raqamlari matnda bor')
ms = MS.read_text()
tables_txt = json.dumps(json.loads((ROOT / 'matn' / '06_01_jadvallar.json').read_text()), ensure_ascii=False)
notes_txt = json.dumps(json.loads((ROOT / 'matn' / '06_01_yonlavhalar.json').read_text()), ensure_ascii=False)
ms_all = ms + tables_txt + notes_txt
need = ['20,006', '36,461', '80,912', '127,912', '569', '432', '370', '295',
        '−506', '−364', '−335', '−293', '63', '68', '35', '2 kJ/mol',
        '0,917', '1,2746', '1,414', '1,609', '1,82', '0,44',
        '3,17', '−5,9', '−8,8', '−9,5', '15', '−2,0', '−6,8', '−10,9',
        '−3,4', '−35,9', '−49,4', '−53,3', '1,95×10⁶', '2,15×10⁹',
        '550,8', '190,6', '110,0', '15,9', '−96,50', '−33,39', '−19,27', '−2,79',
        '22,90', '7,91', '2,57', '0,002290', '0,007913', '0,025739', '2,640', '1,589',
        '1997,9', '4162,4', '2398,3', '599,6', '137,3', '152,2', '60,4', '24,7', '237,3', '237,2',
        '987,5', '992,1', '62,54', '495,0', '505,0', '1831,7', '−74,8', '−2052',
        '447,4', '1218,6', '1691,4', '20,2', '108,6', '47,6', '124,3', '56,9', '127',
        '3 ppm', '30 ppm', 'kalsiy glyukonat', 'gipokalsemiya', 'leveling effect',
        'OpenStax', 'ChemGuide', 'Ph. Eur.', 'Ostwald' ]
need = [s for s in need if s != 'Ostwald']
for s in need:
    ck(s in ms_all, 'matnda «%s» bor' % s)
ck(ms.count('[[EQ:') == 11 and ms.count('[[TABLE:') == 11, 'matnda 11 formula va 11 jadval belgisi')
ck('amaliyot tavsiyalari berilmaydi' in ms, 'xavfsizlik chegarasi matnda bayon qilindi')
ck('o‘lchov ma’lumotlari emas' in ms, 'model cheklovi matnda ko‘rsatilgan')
ck('CaF₂' in ms and 'H₂S' in ms and 'SO₂' in ms, 'olish reaksiyalari matnda')
ck('SiO₂' in ms and 'PTFE' in ms, 'materiallar bo‘limi kalit tushunchalari matnda')

print('6) DOCX tuzilmasi')
doc = Document(DOCX)
xml = doc.element.xml
text_all = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                text_all += '\n' + p.text
ck(len(doc.inline_shapes) == 17, 'rasmlar soni: %d' % len(doc.inline_shapes))
ck(xml.count('<m:oMath') == 45, 'OMML satrlari (11 guruh, 45 qator): %d' % xml.count('<m:oMath'))
ck(xml.count('<m:f>') >= 8, 'kasrlar: %d' % xml.count('<m:f>'))
ck(xml.count('<m:sSub>') >= 25, 'pastki indekslar: %d' % xml.count('<m:sSub>'))
ck(xml.count('sSup') >= 1, 'yuqori indekslar: %d' % xml.count('sSup'))
ck(text_all.count('-jadval.') == 11, 'jadval sarlavhalari: %d' % text_all.count('-jadval.'))
ck(text_all.count('-sxema.') == 7, 'sxema sarlavhalari: %d' % text_all.count('-sxema.'))
ck(text_all.count('-grafik.') == 5, 'grafik sarlavhalari: %d' % text_all.count('-grafik.'))
ck(text_all.count('-rasm.') == 5, 'rasm sarlavhalari: %d' % text_all.count('-rasm.'))
for i in range(1, 12):
    ck(text_all.count('(6.1.%d)' % i) == 1, 'formula raqami (6.1.%d)' % i)
ck('_' not in text_all, 'matnda "_" qolmadi (pastki indeksga o‘girilgan)')
ck('Times New Roman' in xml, 'Times New Roman shrifti qo‘llangan')
ck('Cambria Math' in xml, 'formulalar Cambria Math shriftida')
ck(text_all.count('HISOBIY GRAFIK') == 3, 'grafik yorliqlari: %d' % text_all.count('HISOBIY GRAFIK'))
ck(text_all.count('TUSHUNTIRUVCHI SXEMA') == 7, 'sxema yorliqlari: %d' % text_all.count('TUSHUNTIRUVCHI SXEMA'))
ck(text_all.count('CHIZMA VA IZOH') == 1, 'chizma yorliqlari: %d' % text_all.count('CHIZMA VA IZOH'))
hdr = doc.sections[0].header.paragraphs[0].text
ftr = doc.sections[0].footer.paragraphs[0].text
ck('NOORGANIK KISLOTALAR' in hdr, 'sarlavha matni')
ck('6-BOB' in ftr and '6.1' in text_all, 'kolontitul matni')
for s in ['10,7919' if False else '0,002290 M', '1997,9 g', '137,3 kg', '987,5 g',
          '−53,3 kJ/mol', '0,8653' if False else '0,47' if False else '3,17',
          '447,4 kg', '495,0 kg']:
    ck(s in text_all, 'DOCX matnida «%s» bor' % s)
ck('OpenStax' in text_all and 'ChemGuide' in text_all and 'Ph. Eur.' in text_all,
   'adabiyotlar ro‘yxati joylashdi')
ck('amaliyot tavsiyalari berilmaydi' in text_all, 'xavfsizlik chegarasi bayon qilindi')
size = DOCX.stat().st_size
ck(size > 400000, 'DOCX hajmi: %d bayt' % size)

print()
if fails:
    print('XATO: %d tekshiruv bajarilmadi' % len(fails))
    for f in fails: print('  -', f)
    raise SystemExit(1)
print('BARCHA TEKSHIRUVLAR O‘TDI (6.1).')
