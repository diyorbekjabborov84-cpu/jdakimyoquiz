"""Independent balance, thermodynamic identity and native-equation layout checks."""
from pathlib import Path
from collections import Counter
import csv,math,re,json
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm
from redox_04_02 import R,T,F,S,worked
D=Path(__file__).resolve().parents[1];A=D/'rasmlar/redoks'
def near(a,b):assert math.isclose(a,b,rel_tol=1e-10,abs_tol=2e-13),(a,b)
points=0
for name,size in [('energiya',281),('ph',201),('sigim',6),('peroksid',201)]:
 rows=list(csv.DictReader((A/(name+'.csv')).open()));assert len(rows)==size;points+=size
 for row in rows:
  if name=='energiya':
   e=float(row['E_cell_V'])
   for n in [1,2,5]:near(float(row[f'dG_n{n}_kJ'])*1000,-n*F*e)
  elif name=='ph':
   ph=float(row['pH'])
   for col,e,m,n in [('H_V',0,2,2),('Mn_V',1.507,8,5),('NO_V',.957,4,3)]:
    # Q=10^(m*pH), other ratios=1: natural log definition.
    near(float(row[col]),e-R*T/(n*F)*math.log(10**(m*ph)))
  elif name=='peroksid':
   ph=float(row['pH']);near(float(row['water_V']),1.776-R*T/F*math.log(10)*ph)
   near(float(row['oxygen_V']),.695-R*T/F*math.log(10)*ph)
   near(float(row['water_V'])-float(row['oxygen_V']),1.081)
# ASCII formula parser, with separate explicit charges; not inferred from typography.
def atoms(formula):
 stack=[Counter()];tokens=re.findall(r'[A-Z][a-z]?|\d+|[()]',formula);i=0
 while i<len(tokens):
  t=tokens[i]
  if t=='(':stack.append(Counter());i+=1;continue
  if t==')':
   c=stack.pop();i+=1;n=int(tokens[i]) if i<len(tokens) and tokens[i].isdigit() else 1
   if i<len(tokens) and tokens[i].isdigit():i+=1
   for a,v in c.items():stack[-1][a]+=n*v
   continue
  assert t.isalpha();i+=1;n=int(tokens[i]) if i<len(tokens) and tokens[i].isdigit() else 1
  if i<len(tokens) and tokens[i].isdigit():i+=1
  stack[-1][t]+=n
 assert len(stack)==1;return stack[0]
species={k:(atoms(f),z) for k,f,z in [
('Zn','Zn',0),('Zn2','Zn',2),('H','H',1),('H2','H2',0),('e','',-1),('Mn4','MnO4',-1),('Mn2','Mn',2),('MnO2','MnO2',0),('manganat','MnO4',-2),('OH','OH',-1),('water','H2O',0),('Fe2','Fe',2),('Fe3','Fe',3),('Cu','Cu',0),('Cu2','Cu',2),('HNO3','HNO3',0),('CuN','Cu(NO3)2',0),('NO2','NO2',0),('NO','NO',0),('NO3','NO3',-1),('N2O','N2O',0),('N2','N2',0),('NH4','NH4',1),('H2SO4','H2SO4',0),('CuS','CuSO4',0),('SO2','SO2',0),('SO4','SO4',-2),('H2S','H2S',0),('S','S',0),('H3PO3','H3PO3',0),('H3PO4','H3PO4',0),('H3PO2','H3PO2',0),('Ag','Ag',0),('Ag1','Ag',1),('HNO2','HNO2',0),('H2O2','H2O2',0),('O2','O2',0),('HCl','HCl',0),('NH3','NH3',0),('NH4Cl','NH4Cl',0)]}
rxns=[
({'Zn':1,'H':2},{'Zn2':1,'H2':1}),
({'Mn4':1,'H':8,'e':5},{'Mn2':1,'water':4}),
({'Fe2':1},{'Fe3':1,'e':1}),
({'Mn4':1,'H':8,'Fe2':5},{'Mn2':1,'water':4,'Fe3':5}),
({'Mn4':1,'water':2,'e':3},{'MnO2':1,'OH':4}),
({'Mn4':1,'e':1},{'manganat':1}),
({'Cu':1,'HNO3':4},{'CuN':1,'NO2':2,'water':2}),
({'Cu':3,'HNO3':8},{'CuN':3,'NO':2,'water':4}),
({'NO3':1,'H':4,'e':3},{'NO':1,'water':2}),
({'NO3':1,'H':2,'e':1},{'NO2':1,'water':1}),
({'NO3':2,'H':10,'e':8},{'N2O':1,'water':5}),
({'NO3':2,'H':12,'e':10},{'N2':1,'water':6}),
({'NO3':1,'H':10,'e':8},{'NH4':1,'water':3}),
({'Cu':1,'H2SO4':2},{'CuS':1,'SO2':1,'water':2}),
({'SO2':1,'water':2},{'SO4':1,'H':4,'e':2}),
({'H2S':1},{'S':1,'H':2,'e':2}),
({'H2S':1,'water':4},{'SO4':1,'H':10,'e':8}),
({'SO4':1,'H':4,'e':2},{'SO2':1,'water':2}),
({'SO4':1,'H':8,'e':6},{'S':1,'water':4}),
({'SO4':1,'H':10,'e':8},{'H2S':1,'water':4}),
({'H3PO3':1,'water':1},{'H3PO4':1,'H':2,'e':2}),
({'H3PO2':1,'water':2},{'H3PO4':1,'H':4,'e':4}),
({'H3PO3':1,'water':1,'Ag1':2},{'H3PO4':1,'H':2,'Ag':2}),
({'H3PO2':1,'water':2,'Ag1':4},{'H3PO4':1,'H':4,'Ag':4}),
({'HNO2':3},{'HNO3':1,'NO':2,'water':1}),
({'H2O2':1,'H':2,'e':2},{'water':2}),
({'H2O2':1},{'O2':1,'H':2,'e':2}),
({'H2O2':2},{'water':2,'O2':1}),
({'NO':2,'O2':1},{'NO2':2}),
({'HCl':1,'NH3':1},{'NH4Cl':1})]
for left,right in rxns:
 def totals(side):
  c=Counter();q=0
  for k,n in side.items():
   a,z=species[k];q+=n*z
   for element,v in a.items():c[element]+=v*n
  return c,q
 assert totals(left)==totals(right),(left,right,totals(left),totals(right))
w=worked();near(w['xi'],.0004);near(w['Fe2']*.07,.0005);near(w['Fe3']*.07,.002);near(w['H_consumed'],.0032)
near(w['Fe2']+w['Fe3'],.0025/.07);near(w['Mn2']*.07,.0004);near(w['E'],.8066174775594098)
assert 0<w['Mn4']<1e-50
q=w['Mn2']*w['Fe3']**5/(w['Mn4']*w['Fe2']**5*.1**8)
near(.736-R*T/(5*F)*math.log(q),0)
near(-5*F*.736,-355066.0222016);near(5*.736/S,62.20487580546152)
near(-2*F*1.081,-208601.28804344)
# Layout and math objects.
p=D/'chiqish/Noorganik-kislotalar-4.2-Oksidlovchilik-va-qaytaruvchilik.docx';doc=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for cell in row.cells:yield from paras(cell)
ps=list(paras(doc));txt='\n'.join(p.text for p in ps)
assert len(doc.inline_shapes)==13
for label,n in [('jadval.',9),('grafik.',4),('sxema.',5),('rasm.',4)]:assert sum(label in p.text for p in ps)==n
assert all(txt.count(f'(4.2.{i})')==1 for i in range(1,9))
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
assert '_' not in txt and '[[EQ:' not in txt and '[[TABLE:' not in txt and '4.1.' not in txt
outer=[t for t in doc.tables if len(t.columns)==2];assert len(outer)==9
for t in outer:
 assert abs(t.columns[0].width-Cm(13.5))<1000 and abs(t.columns[1].width-Cm(4.5))<1000
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(p) as z:
 assert z.testzip() is None;tree=etree.fromstring(z.read('word/document.xml'))
assert len(tree.xpath('//m:oMath',namespaces=ns))==17
assert len(tree.xpath('//m:f',namespaces=ns))==3
assert len(tree.xpath('//w:tblHeader',namespaces=ns))==9
assert len(tree.xpath('//m:sSub',namespaces=ns))>30 and len(tree.xpath('//m:sSup',namespaces=ns))>20
alltext=' '.join(tree.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
assert '_' not in alltext
notes=json.loads((D/'matn/04_02_yonlavhalar.json').read_text());assert sum(map(len,notes))==18
print(f'PASS: {points} graph values, {len(rxns)} atom/charge reactions, worked mass/electron/proton balances, Nernst trace consistency; 17 native Equation lines / 3 fractions, 9 tables, 13 figures, 18 notes, exact 3:1 width.')
print('Words including math:',len(alltext.split()),'Bytes:',p.stat().st_size)
# Optional native Google Docs roundtrip, when the delivery export is present.
rt=D.parent/'google_drive/Noorganik kislotalar — 4.2. Oksidlovchilik va qaytaruvchilik — yakuniy.roundtrip.docx'
if rt.exists():
 with ZipFile(rt) as z:gt=etree.fromstring(z.read('word/document.xml'))
 assert len(gt.xpath('//m:oMath',namespaces=ns))==17
 assert len(gt.xpath('//m:f',namespaces=ns))==3
 gs=' '.join(gt.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
 assert gs.count('jadval.')==9 and all(f'(4.2.{i})' in gs for i in range(1,9))
 print('Google Docs→DOCX roundtrip PASS: 17 native math lines and 3 fractions retained.')
