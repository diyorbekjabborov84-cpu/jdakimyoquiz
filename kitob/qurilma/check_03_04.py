"""Scientific model and DOCX integrity checks for titration topic 3.4."""
from pathlib import Path
import math,csv,zipfile,json
from docx import Document
from titration_03_04 import *
R=Path(__file__).resolve().parents[1];A=R/'rasmlar/titrlash'
def near(a,b,rel=2e-11):assert math.isclose(a,b,rel_tol=rel,abs_tol=1e-14),(a,b)
counts={}
for name,ks in [('kuchli',None),('kuchsiz',(1e-5,)),('fosfat',PHOS)]:
 rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/(name+'.csv')).open())];counts[name]=len(rows)
 for r in rows:
  v=r['Vb_mL']/1000;p=r['pH'];h=10**(-p);vol=V0+v;c=CA*V0/vol;na=CB*v/vol
  if ks is None:charge=c
  else:
   al=fractions(h,ks);near(sum(al),1);assert all(0<=a<=1 for a in al)
   for i,k in enumerate(ks):near(h*al[i+1]/al[i],k)
   charge=c*sum(i*a for i,a in enumerate(al))
  assert abs(h+na-KW/h-charge)<1e-13*max(c,na,h,KW/h),(name,r)
 assert all(a['pH']<b['pH'] for a,b in zip(rows,rows[1:]))
rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/'indikator.csv').open())];counts['indikator']=len(rows)
for r in rows:
 near(r['fraction_In']+r['fraction_HIn'],1)
 near(r['fraction_In']/r['fraction_HIn'],10**(r['pH']-7),rel=2e-11)
rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/'hosila.csv').open())];counts['hosila']=len(rows)
for r in rows:
 v=r['Vb_mL']/1000;h=10**(-r['pH']);dD=-(CA+CB)*V0/(V0+v)**2
 exact=-dD/(math.log(10)*h*(1+KW/h**2))/1000
 near(r['dpH_per_mL'],exact,rel=2e-5)
assert abs(max(rows,key=lambda r:r['dpH_per_mL'])['Vb_mL']-25)<.00001
assert counts=={'kuchli':1001,'kuchsiz':1001,'fosfat':701,'indikator':321,'hosila':401}
for k in [None,1e-5]:
 for p in [4,7,8.3,10]:near(ph(endpoint_volume(p,k),None if k is None else (k,)),p)
near(ph(.025),7)
near(ph(.025,(1e-5,)),8.849497727228037)
near(CB*.02498/.025,.09992)
ur=math.sqrt((.03/24.98)**2+(.02/25)**2+.001**2);assert abs(100*ur-.1756)<.0001
p=R/'chiqish/Noorganik-kislotalar-3.4-Titrlash-va-indikatorlar.docx';d=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(d));txt='\n'.join(p.text for p in ps)
assert len(d.inline_shapes)==13
for label,count in [('jadval.',9),('grafik.',5),('sxema.',4),('rasm.',4)]:assert sum(label in p.text for p in ps)==count,(label,count)
assert all(f'(3.4.{i})' in txt for i in range(1,9))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt
assert '3.3.' not in txt and 'DISSOTSIATSIYA VA pH' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
notes=json.loads((R/'matn/03_04_yonlavhalar.json').read_text());assert len(notes)==9 and sum(map(len,notes))==18
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert xml.count('w:tblHeader')==9
 assert 'subscript' in xml and 'superscript' in xml
print('PASS: 3425 graph points; mass/charge/K; analytic derivative; 8 endpoint inversions; uncertainty and DOCX checks')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size,'relative standard uncertainty (%):',ur*100)
