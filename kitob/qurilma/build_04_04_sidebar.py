#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. Alohida mavzu moduli."""
from book_equations import add_equation
from pathlib import Path
import re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'chiqish'/'Noorganik-kislotalar-4.4-Kompleks-va-chokma-hosil-bolish.docx'
ASSETS=ROOT/'rasmlar'/'kompleks'
NAVY='0B2C4A'; BRONZE='C28E3B'; PALE='EDF3F7'; INK='243746'

panels=[
('Metall turlarini ajrating',[],'4.4.1-sxema. Erkin akva-ion, oddiy kompleks, gidroksokompleks va qattiq faza; modda balansiga kirish xususiyatlari.'),
('Q_sp va K_sp farqi',[],'4.4.2-sxema. Ion ko‘paytmasi holatlari, to‘yinganlik sharti va nukleatsiya to‘sig‘i.'),
('Turlar taqsimotini o‘qing',[],'4.4.3-grafik. Erkin Ag⁺ va [Ag(NH₃)₂]⁺ ulushlari erkin ligand konsentratsiyasiga bog‘liq; β₂,c=1,7×10⁷ modeli.'),
('Protonlanish ta’siri',[],'4.4.4-grafik. Erkin NH₃ ulushi va shartli hosil bo‘lish konstantasi; pK_a,c=9,25. Model egri chizig‘i.'),
('Anionning uch roli',[],'4.4.5-sxema. Kuzatuvchi ion, cho‘ktiruvchi anion va kuchli ligand; protonlanish bilan o‘zgaruvchi holat.'),
('Erishning o‘sishi',[],'4.4.6-grafik. To‘yinganlik sig‘imining erkin ligand konsentratsiyasiga bog‘liqligi; K_sp,c=1,6×10⁻¹⁰ modeli.'),
('Tanlab cho‘ktirish',[],'4.4.7-sxema. K_sp farqi, protonlanish, komplekslanish va maskalashning selektiv ajratishga ta’siri.'),
('Amfoterlik minimumi',[],'4.4.8-grafik. To‘yinganlik sig‘imining U shakli; K_sp,c=10⁻³⁰, β₄,c=10³⁴ xayoliy model. Real Al uchun emas.'),
('Kislota ta’siri',[],'4.4.9-sxema. Protonlanish, cho‘ktiruvchi anion va komplekslanishning eruvchanlikka qarama-qarshi ta’siri.'),
('Yechim tartibi',[],'4.4.10-sxema. Turlar ro‘yxati, balanslar, yechim va qattiq faza tekshiruvi; manfiy n_s — qayta yechish.'),
('Hisob natijasi',[],'4.4.11-sxema. AgCl–ammiak tizimida ligand balansining yechimi va yaxlitlash xatosi.')]
extra={
 1:[('kesim.png','4.4.1-rasm. Cho‘kma yuzasi va diffuziya qatlamining shartli kesimi.')],
 3:[('ligand.png','4.4.3-rasm. Koordinatsion qobiqning umumiy sxemasi; ligand soni metallga bog‘liq.')],
 4:[('protonatsiya.png','4.4.4-rasm. Protonlanish erkin ligand faolligini qanday kamaytiradi.')],
 5:[('gidroksokompleks.png','4.4.5-rasm. Ko‘p ligandli va gidroksokompleks turlarining umumiy sxemasi.')],
 6:[('erish.png','4.4.6-rasm. Erkin ion va kompleks o‘rtasidagi termodinamik yo‘nalish.')]}

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.widow_control=True
for sty in doc.styles:
    if sty.type==1:
        sty.font.name='Times New Roman'
        rp=sty.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
        rp.append(rf)
sub=dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ','0123456789abmj'));sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ᵐⁿ‡','0123456789+−mn‡'))
TOKEN=re.compile(r'([A-Za-zΑ-Ωα-ωβγδΔ′][A-Za-zΑ-Ωα-ωβγδΔ′₀-₉]*)_([A-Za-z0-9,′]+)')

def runs(p,text,size=11,bold=False,color=INK):
    def put(t,lower=False,upper=False):
        r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if lower:r.font.subscript=True
        if upper:r.font.superscript=True
    def emit(chunk):
        if not chunk:return
        buf=''
        for ch in chunk:
            if ch in sub:put(buf);buf='';put(sub[ch],lower=True)
            elif ch in sup:put(buf);buf='';put(sup[ch],upper=True)
            else:buf+=ch
        put(buf)
    pos=0
    for m in TOKEN.finditer(text):
        emit(text[pos:m.start()])
        put(m.group(1));put(m.group(2),lower=True)
        pos=m.end()
    emit(text[pos:])

def shade(cell,fill):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(x)
def cell_margin(cell,**kwargs):
    x=OxmlElement('w:tcMar')
    for k,v in kwargs.items():
        y=OxmlElement('w:'+k);y.set(qn('w:w'),str(v));y.set(qn('w:type'),'dxa');x.append(y)
    cell._tc.get_or_add_tcPr().append(x)
def border(cell,edge,color,size):
    bs=OxmlElement('w:tcBorders');b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:color'),color);b.set(qn('w:sz'),str(size));bs.append(b);cell._tc.get_or_add_tcPr().append(bs)
def band(label,fill,color,size,accent=False):
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(18)
    c=t.cell(0,0);shade(c,fill);cell_margin(c,top=140,bottom=140,left=200,right=160)
    if accent:border(c,'left',BRONZE,24)
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    runs(p,label,size,True,color)
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.keep_with_next=True

comparisons=json.loads((ROOT/'matn'/'04_04_jadvallar.json').read_text())
comparison_count=0

def add_comparison_table(cell,spec):
    global comparison_count
    comparison_count+=1
    p=cell.add_paragraph();p.paragraph_format.space_before=Pt(8)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    runs(p,spec['caption'],9.5,True,NAVY)
    t=cell.add_table(rows=1,cols=len(spec['headers']));t.autofit=False
    for col,w in zip(t.columns,spec['widths']):col.width=Cm(w)
    for i,label in enumerate(spec['headers']):
        c=t.rows[0].cells[i];c.width=Cm(spec['widths'][i]);shade(c,NAVY)
        cell_margin(c,top=70,bottom=70,left=85,right=85)
        p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
        runs(p,label,9.5,True,'FFFFFF')
    rp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rp.append(repeat)
    for ri,row in enumerate(spec['rows']):
        r=t.add_row();trp=r._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        for ci,value in enumerate(row):
            c=r.cells[ci];c.width=Cm(spec['widths'][ci])
            cell_margin(c,top=70,bottom=70,left=85,right=85)
            if ri%2==0:shade(c,PALE)
            border(c,'bottom','D6DFE7',3)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            runs(p,value,9.5,ci==0,INK)
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(3)

source=(ROOT/'matn'/'04_04.md').read_text()
sections=[]
for raw in source.split('## ')[1:]:
    title,body=raw.split('\n',1);sections.append((title,body.strip().split('\n\n')))
assert len(sections)==11
side_notes=json.loads((ROOT/'matn'/'04_04_yonlavhalar.json').read_text())
assert len(side_notes)==len(sections)
counter=0
for idx,((title,paras),(panel,nodes,note)) in enumerate(zip(sections,panels)):
    if idx:
        pb=doc.add_page_break();pb.paragraph_format.space_after=Pt(0);pb.paragraph_format.space_before=Pt(0);pb.paragraph_format.line_spacing=Pt(1)
        for r in pb.runs:r.font.size=Pt(1)
    if idx==0:
        band('4-BOB  /  UMUMIY KIMYOVIY XOSSALAR',NAVY,'FFFFFF',11)
        band('4.4. Kompleks va cho‘kma hosil qilish',PALE,NAVY,18,True)
    else:
        band('4.4  /  KOMPLEKS VA CHO‘KMA HOSIL BO‘LISH',PALE,NAVY,11,True)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=True
    runs(p,title,14,True,NAVY)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    table.columns[0].width=Cm(13.5);table.columns[1].width=Cm(4.5)
    left,right=table.rows[0].cells
    left.width=Cm(13.5);right.width=Cm(4.5)
    for c in [left,right]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left,top=0,bottom=0,left=0,right=220)
    cell_margin(right,top=140,bottom=130,left=160,right=100)
    shade(right,'F5F8FA');border(right,'left','D2DCE5',6)
    first=True
    for text in paras:
        if text.startswith('[[TABLE:'):
            key=text[len('[[TABLE:'):-2]
            add_comparison_table(left,comparisons[key])
            continue
        p=left.paragraphs[0] if first else left.add_paragraph()
        j=0 if first else 1
        first=False
        if text.startswith('[[EQ:') and text.endswith(']]'):
            text=text[5:-2]
            counter+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            add_equation(p,left,'4.4',counter,runs)
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Cm(.45 if j else 0)
            runs(p,text)
    p=right.paragraphs[0];p.paragraph_format.space_after=Pt(9)
    runs(p,panel,10,True,NAVY)
    graphic={0:'turlar.png',1:'toyinganlik.png',4:'anion.png',6:'selektiv.png',8:'kislota.png',9:'yechim.png',10:'hisob.png'}
    graph={2:'species.png',3:'proton.png',5:'solubility.png',7:'amph.png'}
    is_graph=idx in graph
    pic=graphic.get(idx) or graph.get(idx)
    p=right.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ASSETS/pic),width=Cm(3.92))
    p=right.add_paragraph();p.paragraph_format.space_after=Pt(10)
    runs(p,'HISOBIY GRAFIK' if is_graph else 'TUSHUNTIRUVCHI SXEMA',9,True,NAVY)
    p=right.add_paragraph();runs(p,note,9)
    for image,caption in extra.get(idx+1,[]):
        p=right.add_paragraph();p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ASSETS/image),width=Cm(3.92))
        p=right.add_paragraph();runs(p,caption+' Rang, o‘lcham va geometriya shartli.',9)
    for label,body in side_notes[idx]:
        p=right.add_paragraph()
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        pp=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr')
        rule=OxmlElement('w:top');rule.set(qn('w:val'),'single')
        rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),BRONZE)
        rule.set(qn('w:space'),'5');bd.append(rule);pp.append(bd)
        runs(p,label,9,True,NAVY)
        p=right.add_paragraph()
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        runs(p,body,9)
    if idx==10:
        p=left.add_paragraph();p.paragraph_format.space_before=Pt(8);runs(p,'Adabiyotlar',10,True,NAVY)
        refs=[
        '[1] OpenStax. Chemistry 2e. 15.1 Precipitation and Dissolution. https://openstax.org/books/chemistry-2e/pages/15-1-precipitation-and-dissolution . Erish–cho‘kish muvozanati, K_sp va stoxiometrik eruvchanlik.',
        '[2] OpenStax. Chemistry 2e. 15.2 Lewis Acids and Bases. https://openstax.org/books/chemistry-2e/pages/15-2-lewis-acids-and-bases . Kompleks ionlar, hosil bo‘lish konstantalari va komplekslanish orqali erish.',
        '[3] IUPAC Gold Book. Solubility product (S05742). DOI: 10.1351/goldbook.S05742. https://goldbook.iupac.org/terms/view/S05742 . Erigan ion faolliklarining ko‘paytmasiga ta’rif.',
        '[4] IUPAC Gold Book. Complex (C01203). DOI: 10.1351/goldbook.C01203. https://goldbook.iupac.org/terms/view/C01203 . Noorganik kimyoda «coordination entity» atamasining tavsiya etilishi.',
        '[5] OpenStax. Chemistry 2e. Appendices H, J, K. https://openstax.org/books/chemistry-2e/pages/j-solubility-products ; https://openstax.org/books/chemistry-2e/pages/k-formation-constants-for-complex-ions ; https://openstax.org/books/chemistry-2e/pages/h-ionization-constants-of-weak-acids . K_a, K_sp va β_f qiymatlarining jadvallari.',
        '[6] IUPAC Gold Book. Equilibrium constant (E02177). DOI: 10.1351/goldbook.E02177. https://goldbook.iupac.org/terms/view/E02177 . Muvozanat konstantasi tushunchasi va standart holat.',
        'Barcha grafik, sxema va rasmlar original o‘quv ishlanmalari. Sonli egri chiziqlar model hisoblaridir, tajriba yoki o‘lchov ma’lumotlari emas. Rasmlar mikrofoto, atom tuzilmasi yoki masshtabli tuzilma tasviri emas. Moddalar konsentratsiyasi va pH oralig‘i real eksperiment yoki idish materiali tanlash uchun ko‘rsatma emas; tsianid va sulfid bilan bog‘liq xavfli jarayonlar muhokama qilinadi.'
        ]
        for ref in refs:
            p=left.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3);runs(p,ref,9)

p=sec.header.paragraphs[0];runs(p,'NOORGANIK KISLOTALAR',8,False,NAVY)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runs(p,'4-BOB   ·   ',8,False,NAVY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Noorganik kislotalar · 4.4 · Kompleks va cho‘kma hosil qilish'
doc.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
assert len(doc.inline_shapes)==16
assert counter==11
assert comparison_count==11
assert sum(len(notes) for notes in side_notes)==22
print(OUT)
print('11 sections; 11 tables; 4 calculated graphs; 7 schemes; 5 process illustrations; 11 equations; 22 margin notes.')
