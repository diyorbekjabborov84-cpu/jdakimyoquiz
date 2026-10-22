"""Mass, charge, equilibrium and document integrity checks for topic 3.3."""
from pathlib import Path
import csv,math,zipfile,json
from docx import Document
from ph_03_03 import strong_h,mono_mix,first_complete_diprotic,KW
R=Path(__file__).resolve().parents[1];A=R/'rasmlar/dissotsiatsiya_ph'
def near(a,b):assert math.isclose(a,b,rel_tol=2e-11,abs_tol=1e-22),(a,b)
def charge(h,rhs):assert abs(h-rhs)<=2e-14*h,(h,rhs)
counts={}
for name in ['kuchli','kuchsiz','umumiy_ion','sulfat']:
 rows=[{k:float(v) for k,v in row.items()} for row in csv.DictReader((A/(name+'.csv')).open())]
 counts[name]=len(rows)
 for r in rows:
  h=r['h_M'];assert h>0
  if name=='kuchli':
   c=r['C_M'];near(h,strong_h(c));charge(h,c+KW/h);near(r['pH_water'],-math.log10(h));assert r['pH_water']<7
  elif name in ['kuchsiz','umumiy_ion']:
   c=r.get('C_M',.01);cs=r.get('Cs_M',0);a=r['A_M'];ha=c-a
   assert 0<a<c;near(h*a/ha,1e-5);charge(h,cs+a+KW/h)
   if name=='kuchsiz':near(r['pH_HA'],-math.log10(h));assert r['pH_HA']>=r['pH_strong']
   else:near(r['alpha_percent'],100*a/c)
  else:
   c=r['C_M'];ha=r['HSO4_M'];a=r['SO4_M'];assert ha>0 and a>0
   near(ha+a,c);near(h*a/ha,.012);charge(h,ha+2*a+KW/h);near(r['pH_equilibrium'],-math.log10(h))
 if name=='umumiy_ion':assert all(a['alpha_percent']>b['alpha_percent'] for a,b in zip(rows,rows[1:]))
assert counts=={'kuchli':321,'kuchsiz':321,'umumiy_ion':281,'sulfat':321}
# Additional two-weak-acid mixtures, outside the plotted examples.
for c in [0,1e-12,1e-8,1e-4,.01]:
 for cs in [0,1e-8,.001]:
  h,anions=mono_mix([(c,1e-5),(c/3,1e-7)],cs)
  charge(h,cs+sum(anions)+KW/h)
  for cc,k,a in [(c,1e-5,anions[0]),(c/3,1e-7,anions[1])]:
   if cc:near(h*a/(cc-a),k)
near(-math.log10(strong_h(1e-8)),6.978294313542888)
near(-math.log10(.01*.75),2.1249387366083)
cs=(.01*.01+.001*.03)/.04;near(cs,.00325);near(-math.log10(strong_h(cs)),2.48811663861096)
p=R/'chiqish/Noorganik-kislotalar-3.3-Dissotsiatsiya-va-pH.docx';d=Document(p)
def paras(obj):
 yield from obj.paragraphs
 for t in obj.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(d));txt='\n'.join(p.text for p in ps)
assert len(d.inline_shapes)==13
assert sum('jadval.' in p.text for p in ps)==9
assert sum('rasm.' in p.text for p in ps)==4
assert sum('grafik.' in p.text for p in ps)==4
assert sum('sxema.' in p.text for p in ps)==5
assert sum(f'(3.3.{i})' in txt for i in range(1,10))==9
assert '[[EQ:' not in txt and '[[TABLE:' not in txt
assert 'KUCHGA TA’SIR' not in txt and '3.2.' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
notes=json.loads((R/'matn/03_03_yonlavhalar.json').read_text());assert len(notes)==9 and sum(map(len,notes))==18
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert 'subscript' in xml and 'superscript' in xml
 assert xml.count('w:tblHeader')==9
print('PASS: 1244 graph points; charge/mass/K; 15 extra mixtures; worked examples; DOCX tables/fonts/indices/media.')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size)
