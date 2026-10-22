"""Mass/charge/equilibrium, capacity derivative and editable document checks."""
from pathlib import Path
import math,csv,zipfile,json
from docx import Document
from buffer_03_05 import *
R=Path(__file__).resolve().parents[1];A=R/'rasmlar/bufer_gidroliz'
def near(a,b,rel=2e-10):assert math.isclose(a,b,rel_tol=rel,abs_tol=1e-20),(a,b)
def mono(p,c,d,k=1e-5):
 h=10**(-p);aa=c*k/(h+k);ha=c*h/(k+h)
 assert aa>0 and ha>0;near(ha+aa,c);near(h*aa/ha,k)
 assert abs(h+d-KW/h-aa)<1e-13*max(h,abs(d),KW/h,c)
def load(name):return [{k:float(v) if v else None for k,v in row.items()} for row in csv.DictReader((A/(name+'.csv')).open())]
counts={}
for name in ['nisbat','qoshish','sigim','suyultirish','ikkala']:
 rows=load(name);counts[name]=len(rows)
 for r in rows:
  if name=='nisbat':
   mono(r['pH_full'],.1,.1*r['ratio']/(1+r['ratio']));near(r['pH_HH'],5+r['log_ratio'])
  elif name=='qoshish':
   q=r['q_M'];mono(r['pH_full'],.1,.05+q)
   if r['pH_HH_or_empty'] is not None:near(r['pH_HH_or_empty'],5+math.log10((.05+q)/(.05-q)))
  elif name=='suyultirish':mono(r['pH_full'],r['CT_M'],r['CT_M']/2)
  elif name=='sigim':
   p=r['pH'];h=10**(-p);near(r['beta_pair_C_01'],10*r['beta_pair_C_001'])
   for c,key in [(.1,'beta_total_C_01'),(.01,'beta_total_C_001')]:
    def D(p):
     hh=10**(-p);return KW/hh+c*1e-5/(hh+1e-5)-hh
    dt=1e-4;num=(D(p+dt)-D(p-dt))/(2*dt);near(num,r[key],rel=2e-8)
   near(r['beta_total_C_01'],r['beta_pair_C_01']+r['beta_water'])
  else:
   c=r['C_M'];h=10**(-r['pH_full']);a=c*K_HF/(K_HF+h);ha=c*h/(K_HF+h);bh=c*h/(h+K_NH4);b=c*K_NH4/(h+K_NH4)
   near(a+ha,c);near(b+bh,c);near(h*a/ha,K_HF);near(h*b/bh,K_NH4)
   assert abs(h+bh-KW/h-a)<1e-13*max(c,h,KW/h)
   near(r['pH_approx'],-math.log10(math.sqrt(K_HF*K_NH4)))
 if name in ['nisbat','qoshish']:assert all(a['pH_full']<b['pH_full'] for a,b in zip(rows,rows[1:]))
 if name in ['suyultirish','ikkala']:assert all(a['pH_full']>b['pH_full'] for a,b in zip(rows,rows[1:]))
assert counts=={'nisbat':321,'qoshish':481,'sigim':361,'suyultirish':361,'ikkala':321}
near(buffer(c=.01/.11,added=.001/.11),5.17622389381102)
near(buffer(c=.01/.11,added=-.001/.11),4.8242070795999545)
near(solve([(.01,(1e-5,))],spectator=.01),8.500148473648082)
near(solve(cations=[(.01,K_NH4)],spectator=-.01),5.625569818430584)
for d,expected in [(1,4.788344727622908),(2,9.52671400389817)]:
 p=solve([(.01,PHOS)],spectator=d*.01);near(p,expected);h=10**(-p);al=fractions(h,PHOS);near(sum(al),1)
 for i,k in enumerate(PHOS):near(h*al[i+1]/al[i],k)
 assert abs(h+d*.01-KW/h-.01*sum(i*a for i,a in enumerate(al)))<1e-14
near(solve(),7)
p=R/'chiqish/Noorganik-kislotalar-3.5-Buferlar-va-gidroliz.docx';d=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(d));txt='\n'.join(p.text for p in ps)
assert len(d.inline_shapes)==13
for label,count in [('jadval.',9),('grafik.',5),('sxema.',4),('rasm.',4)]:assert sum(label in p.text for p in ps)==count,(label,count)
assert all(f'(3.5.{i})' in txt for i in range(1,10))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt and '3.4.' not in txt
assert 'TITRLASH VA INDIKATORLAR' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
notes=json.loads((R/'matn/03_05_yonlavhalar.json').read_text());assert len(notes)==9 and sum(map(len,notes))==18
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert xml.count('w:tblHeader')==9
 assert 'subscript' in xml and 'superscript' in xml
print('PASS: 1845 graph points, charge/mass/K, capacity derivatives, worked examples, phosphate and DOCX integrity.')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size)
