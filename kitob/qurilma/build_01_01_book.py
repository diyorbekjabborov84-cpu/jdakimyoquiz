#!/usr/bin/env python3
"""1.1-mavzu: kitob maketi. Manba: kitob/matn/01_01.md."""
from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'chiqish'/'Noorganik-kislotalar-1.1-kitob-tahriri.docx'
ASSETS = ROOT/'rasmlar'/'ilmiy'
ASSETS.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)
# Chizmalar tajriba natijasi emas: tushunchalar orasidagi mantiqiy aloqalar.
face='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
def font(size): return ImageFont.truetype(face,size)

def evidence_diagram():
    im=Image.new('RGB',(1800,540),'white'); d=ImageDraw.Draw(im)
    xs=[45,650,1255]
    labels=[('KUZATISH','Rang o‘zgarishi','Neytrallanish'),('TARKIB','Kislorod / vodorod','Bog‘lanish tartibi'),('O‘ZARO TA’SIR','Proton almashinuvi','Elektron jufti')]
    for x,(title,a,b) in zip(xs,labels):
        d.line((x,95,x+490,95),fill='#222222',width=3)
        d.text((x,125),title,font=font(27),fill='#222222')
        d.text((x,210),a,font=font(31),fill='#222222')
        d.text((x,270),b,font=font(31),fill='#222222')
        d.line((x,345,x+490,345),fill='#999999',width=2)
    for x in [555,1160]:
        d.line((x,230,x+55,230),fill='#333333',width=3)
        d.polygon([(x+55,230),(x+41,222),(x+41,238)],fill='#333333')
    d.text((45,410),'Belgini aniqlash → tarkibni tekshirish → reaksiyaning mexanizmini izohlash',font=font(30),fill='#444444')
    path=ASSETS/'01_01_bilish.png'; im.save(path); return path

def criteria_diagram():
    im=Image.new('RGB',(1800,650),'white'); d=ImageDraw.Draw(im)
    rows=[('KISLOROD BOR','Kislota bo‘lishi shart emas','Na₂O: asosli oksid'),('KISLOROD YO‘Q','Kislota bo‘lishi mumkin','HCl: suvda kislotali muhit'),('VODOROD BOR','Kislota bo‘lishi shart emas','CH₄: oddiy suvli muhitda')]
    for i,(a,b,c) in enumerate(rows):
        y=40+i*190
        d.line((40,y+150,1750,y+150),fill='#AAAAAA',width=2)
        d.text((40,y+30),a,font=font(29),fill='#222222')
        d.text((550,y+30),b,font=font(32),fill='#222222')
        d.text((550,y+85),c,font=font(29),fill='#555555')
    path=ASSETS/'01_01_mezon.png'; im.save(path); return path

def structure_diagram():
    im=Image.new('RGB',(1800,760),'white'); d=ImageDraw.Draw(im)
    ink='#222222'
    for cx,label,lower in [(450,'H₃PO₃','H'),(1350,'H₃PO₄','OH')]:
        cy=315
        d.text((cx-85,45),label,font=font(43),fill=ink)
        d.text((cx-22,cy-25),'P',font=font(48),fill=ink)
        d.text((cx-23,cy-173),'O',font=font(43),fill=ink)
        d.line((cx-7,cy-40,cx-7,cy-112),fill=ink,width=3)
        d.line((cx+6,cy-40,cx+6,cy-112),fill=ink,width=3)
        d.line((cx-41,cy+5,cx-155,cy+5),fill=ink,width=3)
        d.line((cx+43,cy+5,cx+151,cy+5),fill=ink,width=3)
        d.text((cx-242,cy-22),'HO',font=font(43),fill=ink)
        d.text((cx+168,cy-22),'OH',font=font(43),fill=ink)
        d.line((cx,cy+51,cx,cy+124),fill=ink,width=3)
        d.text((cx-23,cy+143),lower,font=font(43),fill=ink)
        d.text((cx-187,585),'2 ta O–H guruhi' if lower=='H' else '3 ta O–H guruhi',font=font(33),fill=ink)
        d.text((cx-170,643),'ikki asosli kislota' if lower=='H' else 'uch asosli kislota',font=font(31),fill='#555555')
    path=ASSETS/'01_01_tuzilish.png';im.save(path);return path

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(17); sec.page_height=Cm(24)
sec.top_margin=Cm(1.9); sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.1); sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.8); sec.footer_distance=Cm(.85)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=Pt(14.5)
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.first_line_indent=Cm(.65)
normal.paragraph_format.widow_control=True
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    sty=doc.styles[name]; sty.font.name='Times New Roman'
    rpr=sty.element.get_or_add_rPr(); rf=OxmlElement('w:rFonts')
    for a in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+a),'Times New Roman')
    rpr.append(rf)
    sty.font.color.rgb=RGBColor.from_string('171717')
# Bosma kitob singari ichki va tashqi hoshiyalarni navbatlash.
mirror=OxmlElement('w:mirrorMargins'); doc.settings.element.append(mirror)
sec.different_first_page_header_footer=True
# Shu hujjatda shrift almashtiruvchi mavzu ranglari va dekorativ chiziqlar yo‘q.
for name in ['Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.size=Pt(12); st.font.bold=True
    st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.space_before=Pt(13)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True

sub=dict(zip('₀₁₂₃₄₅₆₇₈₉','0123456789'))
sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻','0123456789+−'))
def text_runs(p,text,size=None):
    for token in re.findall(r'[₀₁₂₃₄₅₆₇₈₉]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+|[^₀₁₂₃₄₅₆₇₈₉⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+',text):
        issub=all(c in sub for c in token); issup=all(c in sup for c in token)
        r=p.add_run(''.join(sub[c] for c in token) if issub else ''.join(sup[c] for c in token) if issup else token)
        r.font.name='Times New Roman'
        if size:r.font.size=Pt(size)
        if issub:r.font.subscript=True
        if issup:r.font.superscript=True
    return p

def figure(path,caption):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(13.1))
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_together=True
    p.paragraph_format.line_spacing=Pt(12)
    p.paragraph_format.space_after=Pt(10)
    text_runs(p,caption,10)

p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0)
p.paragraph_format.space_before=Pt(18)
p.paragraph_format.space_after=Pt(12)
p.paragraph_format.keep_with_next=True
r=p.add_run('1-BOB. KISLOTA TUSHUNCHASI VA NAZARIYALAR');r.font.size=Pt(9)
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
p.paragraph_format.space_after=Pt(16)
p.paragraph_format.line_spacing=Pt(25)
p.paragraph_format.keep_with_next=True
r=p.add_run('1.1. Kislota tushunchasining\nshakllanishi');r.bold=True;r.font.size=Pt(20)

source=(ROOT/'matn'/'01_01.md').read_text()
blocks=source.strip().split('\n\n')
eqnum=0
first_paragraph=True
for block in blocks:
    if block.startswith('# '):continue
    if block.startswith('## '):
        title=block[3:]
        if title=='Kuzatiladigan xossadan ilmiy tushunchaga':
            # Birinchi bo‘lim ochilish matni; sarlavha takrori yo‘q.
            continue
        doc.add_paragraph(title,'Heading 2');first_paragraph=True;continue
    p=doc.add_paragraph()
    if '\n' not in block and re.search(r' [→⇌] ',block):
        eqnum+=1
        p.paragraph_format.tab_stops.add_tab_stop(Cm(6.55), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(13.1), WD_TAB_ALIGNMENT.RIGHT)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.keep_together=True
        text_runs(p,'\t'+block+'\t(1.'+str(eqnum)+')')
    else:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if first_paragraph:
            p.paragraph_format.first_line_indent=Cm(0)
            first_paragraph=False
        text_runs(p,block)
    if block.startswith('Ta’m kislotalarning'):
        figure(evidence_diagram(),'1.1-rasm. Kislotalarni o‘rganishdagi uch daraja. Kuzatiladigan belgi moddaning tarkibi yoki reaksiyadagi vazifasi bilan aynan bir narsa emas.')
    if block.startswith('Vodorod atomlarining umumiy soni'):
        figure(structure_diagram(),'1.3-rasm. Fosfit va ortofosfat kislotalarning odatiy tuzilish formulalari. Bir xil umumiy vodorod soni bir xil asoslikni anglatmaydi: H₃PO₃ dagi P–H vodorodi odatiy suvli dissotsiatsiyada ajralmaydi [4].')
    if block.startswith('Vodorod xlorid va vodorod sulfid'):
        figure(criteria_diagram(),'1.2-rasm. Element tarkibiga asoslangan mezonlarning chegarasi. Kislorodning mavjudligi kislotalilik uchun na zarur, na yetarli shartdir.')

doc.add_paragraph('Mavzuga oid adabiyotlar','Heading 2')
refs=[
'Partington, J. R. A Short History of Chemistry. 3-nashr. Macmillan, 1957. Kimyoviy tushunchalar va tarixiy tajribalar taraqqiyoti.',
'Arrhenius, S. Über die Dissociation der in Wasser gelösten Stoffe. Zeitschrift für physikalische Chemie, 1, 1887, 631–648.',
'Lewis, G. N. Valence and the Structure of Atoms and Molecules. Chemical Catalog Company, 1923.',
'Housecroft, C. E.; Sharpe, A. G. Inorganic Chemistry. 5-nashr. Pearson, 2018. Kislota–asos tushunchalari va noorganik birikmalar tuzilishi.'
]
for i,t in enumerate(refs,1):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(-.4)
    p.paragraph_format.left_indent=Cm(.4)
    p.paragraph_format.line_spacing=Pt(12)
    p.paragraph_format.space_after=Pt(5)
    text_runs(p,f'{i}. {t}',10)
# Diskret kolontitullar, sahifa raqami maydon ko‘rinishida.
p=sec.header.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0)
p.paragraph_format.space_after=Pt(0)
r=p.add_run('NOORGANIK KISLOTALAR');r.font.name='Times New Roman';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
p=sec.footer.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();r.font.name='Times New Roman';r.font.size=Pt(10)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
p=sec.first_page_footer.paragraphs[0]
p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
# Hujjatning tilini belgilash.
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'uz-Latn-UZ');normal.element.get_or_add_rPr().append(lang)
doc.core_properties.title='Noorganik kislotalar — 1.1. Kislota tushunchasining shakllanishi'
doc.core_properties.subject='Ilmiy darslik mavzusi'
doc.core_properties.author=''
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
doc.save(OUT)
print(OUT)
print('Paragraphs:',len(doc.paragraphs),'Figures:',len(doc.inline_shapes))
# Formulalarning indekslari Word run xususiyatlari bilan berilganini tekshirish.
assert len(doc.inline_shapes)==3
assert any(r.font.subscript for p in doc.paragraphs for r in p.runs)
assert any(r.font.superscript for p in doc.paragraphs for r in p.runs)
assert all(r.font.name=='Times New Roman' for p in doc.paragraphs for r in p.runs if r.text)
