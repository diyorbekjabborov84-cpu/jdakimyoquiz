"""QA for equilibrium, heat, stoichiometry, ionic balance and editable layout of 4.1."""
from pathlib import Path
import csv,math,zipfile,json
from docx import Document
from docx.shared import Cm
from neutral_04_01 import *
D=Path(__file__).resolve().parents[1];A=D/'rasmlar/neytrallanish'
def near(a,b,rel=2e-10):assert math.isclose(a,b,rel_tol=rel,abs_tol=1e-20),(a,b)
rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/'yonalish.csv').open())];assert len(rows)==481
for r in rows:
 f=r['fraction'];k=r['K'];assert 0<f<1
 near((f/(1-f))**2,k);near(math.log10(k),r['delta_pKa']);near(r['dG_standard_Jmol'],-R*T*math.log(k))
rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/'fosfat.csv').open())];assert len(rows)==301
for r in rows:
 h=10**(-r['pH']);al=[r[k] for k in ['alpha_H3PO4','alpha_H2PO4','alpha_HPO4','alpha_PO4']]
 near(sum(al),1);assert all(0<a<1 for a in al)
 for j,k in enumerate(PHOS):near(h*al[j+1]/al[j],k)
 charge=.01*sum(i*a for i,a in enumerate(al));na=.01*r['OH_equivalents_per_P']
 assert abs(h+na-KW/h-charge)<1e-14
near(rows[-1]['alpha_PO4'],.24158267332262107)
rows=list(csv.DictReader((A/'tuz_ph.csv').open()));assert len(rows)==4
for r in rows:near(float(r['pH']),salt_ph()[r['salt_model']]);near(float(r['C_M']),.01)
rows=[{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/'issiqlik.csv').open())];assert len(rows)==201
for r in rows:
 xi=r['extent_mmol']/1000;near(r['q_abs_model57_J'],57000*xi);near(r['q_abs_model42_J'],42000*xi)
 qrxn,qabs,dt=calorimetry(xi);near(qrxn+qabs,0);near(dt*(200*4.18+20),qabs)
xi,h,p,na,cl=strong_mix(.15,.04,.2,.025);near(xi,.005);near(p,1.812913356624507)
near(h+na,cl+KW/h);near(xi*58.44,.2922)
# Check both sides and exact equivalence of the stable quadratic.
for ratio in [0.1,0.9,1,1.1,10]:
 xi,h,p,na,cl=strong_mix(.1,.025,.1,.025*ratio);near(h+na,cl+KW/h)
near(calorimetry()[2],.6658878504672897)
# Independent atom/charge balances for all explicit representative reactions.
atoms={'HCl':{'H':1,'Cl':1},'NaOH':{'Na':1,'O':1,'H':1},'NaCl':{'Na':1,'Cl':1},'H2O':{'H':2,'O':1},'H3O':{'H':3,'O':1},'OH':{'O':1,'H':1},'H3PO4':{'H':3,'P':1,'O':4},'NaH2PO4':{'Na':1,'H':2,'P':1,'O':4},'Na2HPO4':{'Na':2,'H':1,'P':1,'O':4},'Na3PO4':{'Na':3,'P':1,'O':4},'NH3':{'N':1,'H':3},'NH4Cl':{'N':1,'H':4,'Cl':1},'CuO':{'Cu':1,'O':1},'CuCl2':{'Cu':1,'Cl':2},'Cu':{'Cu':1},'NaHCO3':{'Na':1,'H':1,'C':1,'O':3},'CO2':{'C':1,'O':2},'HCO3':{'H':1,'C':1,'O':3}}
charge={'H3O':1,'OH':-1,'Cu':2,'HCO3':-1}
rxns=[{'H3O':-1,'OH':-1,'H2O':2},{'HCl':-1,'NaOH':-1,'NaCl':1,'H2O':1},{'H3PO4':-1,'NaOH':-1,'NaH2PO4':1,'H2O':1},{'H3PO4':-1,'NaOH':-2,'Na2HPO4':1,'H2O':2},{'H3PO4':-1,'NaOH':-3,'Na3PO4':1,'H2O':3},{'NH3':-1,'HCl':-1,'NH4Cl':1},{'CuO':-1,'HCl':-2,'CuCl2':1,'H2O':1},{'CuO':-1,'H3O':-2,'Cu':1,'H2O':3},{'NaHCO3':-1,'HCl':-1,'NaCl':1,'CO2':1,'H2O':1},{'HCO3':-1,'H3O':-1,'CO2':1,'H2O':2}]
for rxn in rxns:
 for element in ['H','O','Na','Cl','P','N','Cu','C']:assert sum(v*atoms[m].get(element,0) for m,v in rxn.items())==0
 assert sum(v*charge.get(m,0) for m,v in rxn.items())==0
# Equal weak acid/base: exact neutral model, products ~99.01%, neutrals nonzero.
pp=solve([(.01,(1e-5,))],[(.01,1e-9)]);near(pp,7);near((.01*1e-5/(1e-5+1e-7))/.01,100/101)
p=D/'chiqish/Noorganik-kislotalar-4.1-Neytrallanish-va-tuz-hosil-bolishi.docx';doc=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(doc));txt='\n'.join(p.text for p in ps)
assert len(doc.inline_shapes)==13
for label,count in [('jadval.',9),('grafik.',4),('sxema.',5),('rasm.',4)]:assert sum(label in p.text for p in ps)==count,(label,count)
assert all(f'(4.1.{i})' in txt for i in range(1,9))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt and '3-BOB' not in txt
assert '4-BOB' in txt and 'UMUMIY KIMYOVIY XOSSALAR' in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
outer=[t for t in doc.tables if len(t.columns)==2];assert len(outer)==9
for t in outer:
 assert abs(t.columns[0].width-Cm(13.5))<1000
 assert abs(t.columns[1].width-Cm(4.5))<1000
notes=json.loads((D/'matn/04_01_yonlavhalar.json').read_text());assert len(notes)==9 and sum(map(len,notes))==18
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert xml.count('w:tblHeader')==9
 assert 'subscript' in xml and 'superscript' in xml
print('PASS: 987 graph values, phosphate/K/charge, calorimetry, limiting reagent, 10 atom/charge reactions and DOCX 3+1 layout.')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size)
