"""QA: model identities, stoichiometry, native mathematics and 3:1 layout."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm
import math,csv,re,json
from passive_04_03 import *
D=Path(__file__).resolve().parents[1];A=D/'rasmlar/passivlanish'
def near(a,b):assert math.isclose(a,b,rel_tol=2e-10,abs_tol=1e-12),(a,b)
def rows(name):return [{k:float(v) for k,v in r.items()} for r in csv.DictReader((A/(name+'.csv')).open())]
points=0
for name,n in [('qutblanish',501),('aralash',501),('qalinlik',301),('xlorid',301)]:
 rr=rows(name);assert len(rr)==n;points+=n
 for r in rr:
  if name=='qutblanish':near(10**r['log_i_anodic'],r['i_anodic_Acm2']);assert -6-1e-12<=r['log_i_anodic']<=-2+1e-12
  elif name=='aralash':
   e=r['E_relative_V'];near(r['log_abs_i_weak'],-3-10*e);near(r['log_abs_i_strong'],1-10*e)
  elif name=='qalinlik':
   t=r['time_s']
   for key,v in [('L_v002_nm',.02),('L_v006_nm',.06),('L_v008_nm',.08)]:
    l=r[key];linf=4-v/.02;near(l,linf+(.5-linf)*math.exp(-.02*t));assert l>=0
    dt=.001;deriv=(thickness(t+dt,v)-thickness(t-dt,v))/(2*dt)
    near(deriv,.02*(4-l)-v)
  else:
   x=r['log_a_Cl'];near(math.log10(r['a_Cl']),x)
   near(r['E_pit_model_V'],.65-.08*(x+2));near(r['E_rp_model_V'],.35-.08*(x+2));near(r['E_pit_model_V']-r['E_rp_model_V'],.3)
near(log_anodic(.3),-3);near(log_anodic(.45),-6);near(log_anodic(.8),-6)
near(log_anodic(.15),log_cathodic(.15,-3));near(log_anodic(.7),log_cathodic(.7,1))
# Exactly one sign change per illustrative cathodic curve on the plotted interval.
for strength in [-3,1]:
 vals=[log_anodic(i/10000)-log_cathodic(i/10000,strength) for i in range(10001)]
 signs=[1 if v>1e-10 else -1 for v in vals if abs(v)>1e-10]
 assert sum(a!=b for a,b in zip(signs,signs[1:]))==1
w=worked();near(w['Q'],.864);near(w['m_mg'],.08053745494015095);near(w['h_um'],.02982868701487072)
near(w['local_h_um']/w['h_um'],1000);near(w['local_i']*w['active_area'],1e-6*10)
near(w['m_g']/M*3*F,w['Q']);near(w['h_um']*1e-4*RHO*10,w['m_g'])
b,i=stern();near(b,.026057668914195108);near(i*26057.67,b)
# Independent concrete atom/charge balances; no mechanism inferred.
spec={
 'Cr':({'Cr':1},0),'water':({'H':2,'O':1},0),'Cr2O3':({'Cr':2,'O':3},0),'H':({'H':1},1),'e':({},-1),
 'Al2O3':({'Al':2,'O':3},0),'Al3':({'Al':1},3),'OH':({'O':1,'H':1},-1),'AlOH4':({'Al':1,'O':4,'H':4},-1),
 'NO3':({'N':1,'O':3},-1),'NO2':({'N':1,'O':2},0),'Fe3':({'Fe':1},3),'FeOH2':({'Fe':1,'O':1,'H':1},2),
 'Cr2O7':({'Cr':2,'O':7},-2),'O2':({'O':2},0),'Al':({'Al':1},0)}
rx=[({'Cr':2,'water':3},{'Cr2O3':1,'H':6,'e':6}),({'Al2O3':1,'H':6},{'Al3':2,'water':3}),({'Al2O3':1,'OH':2,'water':3},{'AlOH4':2}),({'NO3':1,'H':2,'e':1},{'NO2':1,'water':1}),({'Fe3':1,'water':1},{'FeOH2':1,'H':1}),({'Cr2O3':1,'water':4},{'Cr2O7':1,'H':8,'e':6}),({'water':2},{'O2':1,'H':4,'e':4}),({'Al':1},{'Al3':1,'e':3})]
for left,right in rx:
 def total(side):
  c=Counter();charge=0
  for k,n in side.items():
   a,z=spec[k];charge+=n*z
   for atom,v in a.items():c[atom]+=n*v
  return c,charge
 assert total(left)==total(right),(left,right)
p=D/'chiqish/Noorganik-kislotalar-4.3-Passivlanish.docx';doc=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(doc));txt='\n'.join(p.text for p in ps)
assert len(doc.inline_shapes)==13
for label,n in [('jadval.',9),('grafik.',4),('sxema.',5),('rasm.',4)]:assert sum(label in p.text for p in ps)==n
assert all(txt.count(f'(4.3.{i})')==1 for i in range(1,9))
assert '[[EQ:' not in txt and '[[TABLE:' not in txt and '4.2.' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
outer=[t for t in doc.tables if len(t.columns)==2];assert len(outer)==9
for t in outer:assert abs(t.columns[0].width-Cm(13.5))<1000 and abs(t.columns[1].width-Cm(4.5))<1000
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(p) as z:assert z.testzip() is None;tree=etree.fromstring(z.read('word/document.xml'))
assert len(tree.xpath('//m:oMath',namespaces=ns))==18
assert len(tree.xpath('//m:f',namespaces=ns))==8
assert len(tree.xpath('//w:tblHeader',namespaces=ns))==9
alltext=' '.join(tree.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
without_urls=re.sub(r'https?://\S+','',alltext)
assert '_' not in without_urls
assert len(tree.xpath('//m:sSub',namespaces=ns))>20 and len(tree.xpath('//m:sSup',namespaces=ns))>10
notes=json.loads((D/'matn/04_03_yonlavhalar.json').read_text());assert sum(map(len,notes))==18
print(f'PASS: {points} model points; roots, continuity, film ODE, Faraday/area/Stern identities, {len(rx)} atom/charge balances; 18 OMML lines / 8 fractions; 9 tables, 13 figures, 18 notes, 3:1 layout.')
print('Words with math:',len(alltext.split()),'Bytes:',p.stat().st_size)
rt=D.parent/'google_drive/Noorganik kislotalar — 4.3. Passivlanish — tekshirilgan.roundtrip.docx'
if rt.exists():
 with ZipFile(rt) as z:gt=etree.fromstring(z.read('word/document.xml'))
 assert len(gt.xpath('//m:oMath',namespaces=ns))==18
 assert len(gt.xpath('//m:f',namespaces=ns))==8
 gs=' '.join(gt.xpath('//w:t/text()|//m:t/text()',namespaces=ns))
 assert gs.count('jadval.')==9 and all(f'(4.3.{i})' in gs for i in range(1,9))
 print('Google Docs→DOCX roundtrip PASS: 18 native math lines and 8 fractions retained.')
