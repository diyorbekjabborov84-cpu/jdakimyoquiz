"""Check the illustrative models and the editable topic document."""
from pathlib import Path
import math,csv,zipfile
from docx import Document
from effects_03_02 import born_factor,pka_temperature,R,T0
ROOT=Path(__file__).resolve().parents[1]
rows=list(csv.DictReader(open(ROOT/'rasmlar/kuch_omillari/born.csv')))
assert len(rows)==251
for row in rows:
 rho=float(row['rho_r_over_r0']);b2=float(row['B_epsilon_2']);b80=float(row['B_epsilon_80'])
 assert b80>b2>0
 assert math.isclose(rho*b2,.5,rel_tol=1e-14)
 assert math.isclose(rho*b80,79/80,rel_tol=1e-14)
for a,b in zip(rows,rows[1:]):assert float(a['B_epsilon_80'])>float(b['B_epsilon_80'])
rows=list(csv.DictReader(open(ROOT/'rasmlar/kuch_omillari/harorat.csv')))
assert len(rows)==241
for row in rows:
 t=float(row['T_K']);p1=float(row['pKa_dH_plus10kJ']);p2=float(row['pKa_dH_minus10kJ'])
 assert math.isclose(p1+p2,10,abs_tol=2e-15)
 for p,dh in [(p1,10000),(p2,-10000)]:
  assert math.isclose(-math.log(10)*(p-5),-dh/R*(1/t-1/T0),abs_tol=2e-15)
for dh in [10000,-10000]:
 assert pka_temperature(T0,dh)==5
 t=300;dt=.001
 derivative=-math.log(10)*(pka_temperature(t+dt,dh)-pka_temperature(t-dt,dh))/(2*dt)
 assert math.isclose(derivative,dh/(R*t*t),rel_tol=1e-8)
p=ROOT/'chiqish/Noorganik-kislotalar-3.2-Kuchga-tasir-etuvchi-omillar.docx';d=Document(p)
def paras(o):
 yield from o.paragraphs
 for t in o.tables:
  for row in t.rows:
   for c in row.cells:yield from paras(c)
ps=list(paras(d));txt='\n'.join(p.text for p in ps)
assert len(d.inline_shapes)==13
assert sum('jadval.' in p.text for p in ps)==9
assert sum('rasm.' in p.text for p in ps)==4
assert '[[EQ:' not in txt and '[[TABLE:' not in txt
assert '3.1.' not in txt and '2-BOB' not in txt
assert all(r.font.name=='Times New Roman' for p in ps for r in p.runs if r.text)
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert 'subscript' in xml and 'superscript' in xml
print('PASS: 251 Born and 241 temperature points, van’t Hoff derivatives, DOCX ZIP/fonts/indices, 13 images and 9 tables')
print('Words:',len(txt.split()),'Bytes:',p.stat().st_size)
