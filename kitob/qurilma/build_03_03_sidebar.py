#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. Alohida mavzu moduli."""
from pathlib import Path
import re, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'chiqish'/'Noorganik-kislotalar-3.3-Dissotsiatsiya-va-pH.docx'
ASSETS=ROOT/'rasmlar'/'dissotsiatsiya_ph_sxema'; ASSETS.mkdir(parents=True,exist_ok=True)
NAVY='0B2C4A'; BRONZE='C28E3B'; PALE='EDF3F7'; INK='243746'
fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
def font(n):return ImageFont.truetype(fontpath,n)

def scheme(index,nodes):
    special={1:'shkala.png',2:'balans.png',3:'kuchli.png',4:'kuchsiz.png',5:'miqdor.png',6:'umumiy_ion.png',7:'sulfat.png',8:'faollik.png',9:'nazorat.png'}
    if index in special:return ROOT/'rasmlar'/'dissotsiatsiya_ph'/special[index]
    im=Image.new('RGB',(640,1100),'#F5F8FA');d=ImageDraw.Draw(im)
    y=65
    if index==-1:
        for cy,label,lower,caption in [(235,'H₃PO₃','H','Ikki O–H guruhi'),(730,'H₃PO₄','OH','Uch O–H guruhi')]:
            cx=315;ink='#0B2C4A'
            d.text((230,cy-180),label,font=font(38),fill=ink)
            d.text((cx-16,cy-17),'P',font=font(38),fill=ink)
            d.text((cx-17,cy-107),'O',font=font(34),fill=ink)
            for dx in [-4,5]:d.line((cx+dx,cy-29,cx+dx,cy-66),fill=ink,width=3)
            d.line((cx-27,cy+5,cx-100,cy+5),fill=ink,width=3)
            d.line((cx+31,cy+5,cx+100,cy+5),fill=ink,width=3)
            d.text((cx-171,cy-17),'HO',font=font(34),fill=ink)
            d.text((cx+111,cy-17),'OH',font=font(34),fill=ink)
            d.line((cx,cy+39,cx,cy+83),fill=ink,width=3)
            d.text((cx-17,cy+96),lower,font=font(34),fill=ink)
            d.text((145,cy+173),caption,font=font(29),fill='#243746')
        path=ASSETS/f'01_01_{index:02d}.png';im.save(path);return path
    for i,(a,b) in enumerate(nodes):
        d.text((45,y),a,font=font(29),fill='#0B2C4A');y+=58
        # Hand-authored diagrams; not measured experimental data.
        lines=textwrap.wrap(b,width=25)
        for line in lines:
            d.text((45,y),line,font=font(30),fill='#243746');y+=43
        if i<len(nodes)-1:
            y+=25;d.line((320,y,320,y+65),fill='#C28E3B',width=5)
            d.polygon([(320,y+75),(309,y+56),(331,y+56)],fill='#C28E3B');y+=115
    path=ASSETS/f'01_01_{index:02d}.png';im.save(path);return path

panels=[
('Logarifm va faollik',[],'3.3.1-sxema. Faolliklar nisbati: bir pH birligi — o‘n karra farq.'),
('Balanslar bir tizim',[],'3.3.2-sxema. Faqat HA va suv uchun konsentratsion ideal tenglamalar.'),
('Suv chegarasiga yaqin',[],'3.3.1-grafik. Kuchli bir asosli kislota: suvli yechim va suvni tashlagan ifoda; 25 °C.'),
('Bir xil C, turli ajralish',[],'3.3.2-grafik. Shartli HA va to‘liq ajralgan HX; ikkala ideal model suvni saqlaydi.'),
('Miqdor va yakuniy hajm',[],'3.3.3-sxema. Ikki kuchli bir asosli kislota aralashmasining hisob yo‘li.'),
('Umumiy H⁺ ta’siri',[],'3.3.3-grafik. Kuchli kislota C_s ortganda shartli HA ning deprotonlangan ulushi.'),
('Ikkinchi proton',[],'3.3.4-grafik. H₂SO₄ ning birinchi bosqichi to‘liq deb olingan ideal model. K₂,c≈0,012 mol L⁻¹; 25 °C.'),
('Faollikni ajrating',[],'3.3.4-sxema. Shartli molal misol: γ koeffitsiyentining pH ga hissasi.'),
('Yechimning nazorati',[],'3.3.5-sxema. Tarkibdan fizik ildiz va balanslarni qayta tekshirishga o‘tish.')
]

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for sty in doc.styles:
    if sty.type==1:
        sty.font.name='Times New Roman'
        rp=sty.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
        rp.append(rf)
sub=dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ','0123456789abmj'));sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻','0123456789+−'))
variables={'K_w,c':('K','w,c'),'K_w':('K','w'),'K_c':('K','c'),'K₂,c':('K','2,c'),'C_s':('C','s'),'V_yakun':('V','yakun'),'pH_neytral':('pH','neytral')}
def runs(p,text,size=11,bold=False,color=INK):
    pattern='('+'|'.join(re.escape(v) for v in variables)+r'|[₀₁₂₃₄₅₆₇₈₉ₐᵦₘⱼ]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+)'
    def put(t,lower=False,upper=False):
        r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if lower:r.font.subscript=True
        if upper:r.font.superscript=True
    for token in re.split(pattern,text):
        if not token:continue
        if token in variables:
            main,low=variables[token];put(main);put(low,lower=True)
        elif all(c in sub for c in token):put(''.join(sub[c] for c in token),lower=True)
        elif all(c in sup for c in token):put(''.join(sup[c] for c in token),upper=True)
        else:put(token)

def shade(cell,fill):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(x)
def cell_margin(cell,**kwargs):
    x=OxmlElement('w:tcMar')
    for k,v in kwargs.items():
        y=OxmlElement('w:'+k);y.set(qn('w:w'),str(v));y.set(qn('w:type'),'dxa');x.append(y)
    cell._tc.get_or_add_tcPr().append(x)
def border(cell,edge,color,size):
    bs=OxmlElement('w:tcBorders');b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:color'),color);b.set(qn('w:sz'),str(size));bs.append(b);cell._tc.get_or_add_tcPr().append(bs)
def band(label,fill,color,size,accent=False):
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(18)
    c=t.cell(0,0);shade(c,fill);cell_margin(c,top=140,bottom=140,left=200,right=160)
    if accent:border(c,'left',BRONZE,24)
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    runs(p,label,size,True,color)
    # Prevent the banner row separating from its following paragraph.
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.keep_with_next=True

comparisons=json.loads((ROOT/'matn'/'03_03_jadvallar.json').read_text())
comparison_count=0

def add_comparison_table(cell,spec):
    global comparison_count
    comparison_count+=1
    p=cell.add_paragraph();p.paragraph_format.space_before=Pt(8)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    runs(p,spec['caption'],9.5,True,NAVY)
    t=cell.add_table(rows=1,cols=len(spec['headers']));t.autofit=False
    for col,w in zip(t.columns,spec['widths']):col.width=Cm(w)
    for i,label in enumerate(spec['headers']):
        c=t.rows[0].cells[i];c.width=Cm(spec['widths'][i]);shade(c,NAVY)
        cell_margin(c,top=70,bottom=70,left=85,right=85)
        p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        runs(p,label,9.5,True,'FFFFFF')
    rp=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');rp.append(repeat)
    for ri,row in enumerate(spec['rows']):
        r=t.add_row();trp=r._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        for ci,value in enumerate(row):
            c=r.cells[ci];c.width=Cm(spec['widths'][ci])
            cell_margin(c,top=70,bottom=70,left=85,right=85)
            if ri%2==0:shade(c,PALE)
            border(c,'bottom','D6DFE7',3)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            runs(p,value,9.5,ci==0,INK)
    # Word requires a trailing paragraph after a nested table.
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(3)

source=(ROOT/'matn'/'03_03.md').read_text()
sections=[]
for raw in source.split('## ')[1:]:
    title,body=raw.split('\n',1);sections.append((title,body.strip().split('\n\n')))
assert len(sections)==9
side_notes=json.loads((ROOT/'matn'/'03_03_yonlavhalar.json').read_text())
assert len(side_notes)==len(sections)
counter=0
for idx,((title,paras),(panel,nodes,note)) in enumerate(zip(sections,panels)):
    if idx:doc.add_page_break()
    if idx==0:
        band('3-BOB  /  KISLOTALILIK KUCHI VA MUVOZANATLAR',NAVY,'FFFFFF',11)
        band('3.3. Dissotsiatsiya va pH',PALE,NAVY,18,True)
    else:
        band('3.3  /  DISSOTSIATSIYA VA pH',PALE,NAVY,11,True)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=True
    runs(p,title,14,True,NAVY)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    # 18 cm printable width: 13.5 cm narrative + 4.5 cm outer visual strip.
    table.columns[0].width=Cm(13.5);table.columns[1].width=Cm(4.5)
    left,right=table.rows[0].cells
    left.width=Cm(13.5);right.width=Cm(4.5)
    for c in [left,right]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left,top=0,bottom=0,left=0,right=220)
    cell_margin(right,top=140,bottom=130,left=160,right=100)
    shade(right,'F5F8FA');border(right,'left','D2DCE5',6)
    # Allow long content to flow rather than shrink type or clip a fixed-height row.
    first=True
    for text in paras:
        if text.startswith('[[TABLE:'):
            key=text[len('[[TABLE:'):-2]
            add_comparison_table(left,comparisons[key])
            continue
        p=left.paragraphs[0] if first else left.add_paragraph()
        j=0 if first else 1
        first=False
        if text.startswith('[[EQ:') and text.endswith(']]'):
            text=text[5:-2]
            counter+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            runs(p,text+'   (3.3.'+str(counter)+')',11)
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Cm(.45 if j else 0)
            runs(p,text)
    p=right.paragraphs[0];p.paragraph_format.space_after=Pt(9)
    runs(p,panel,10,True,NAVY)
    p=right.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(scheme(idx+1,nodes)),width=Cm(3.92))
    p=right.add_paragraph();p.paragraph_format.space_after=Pt(10)
    runs(p,'HISOBIY GRAFIK' if idx in [2,3,5,6] else 'TUSHUNTIRUVCHI SXEMA',9,True,NAVY)
    p=right.add_paragraph();runs(p,note,9)
    extra_models={1:('gidroksoniy_model.png','3.3.1-rasm. H₃O⁺ ionining shar–tayoqcha tasviri.'),2:('hcl_model.png','3.3.2-rasm. Neytral HCl; suvdagi asosiy zarracha deb talqin qilinmaydi.'),3:('hf_model.png','3.3.3-rasm. HF tuzilishi; grafikdagi shartli HA bilan aynanlashtirilmaydi.'),6:('sulfat_model.png','3.3.4-rasm. H₂SO₄ molekulasidagi ikki O–H ulanishi.')}
    if idx in extra_models:
        image,caption=extra_models[idx]
        p=right.add_paragraph();p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'rasmlar'/'dissotsiatsiya_ph'/image),width=Cm(3.92))
        p=right.add_paragraph();runs(p,caption+' Sxematik o‘quv tasviri; tajribaviy koordinatalar emas.',9)
    for label,body in side_notes[idx]:
        p=right.add_paragraph()
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        # Delicate bronze rule groups each short aside without heavy boxes.
        pp=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr')
        rule=OxmlElement('w:top');rule.set(qn('w:val'),'single')
        rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),BRONZE)
        rule.set(qn('w:space'),'5');bd.append(rule);pp.append(bd)
        runs(p,label,9,True,NAVY)
        p=right.add_paragraph()
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        runs(p,body,9)
    if idx==8:
        p=left.add_paragraph();p.paragraph_format.space_before=Pt(8);runs(p,'Adabiyotlar',10,True,NAVY)
        refs=[
        '[1] IUPAC Gold Book. pH. DOI: 10.1351/goldbook.P04524. https://goldbook.iupac.org/terms/view/P04524 . Molal faollik ta’rifi va yakka ion konvensiyasi; mazmun mustaqil o‘zbekcha bayon qilindi.',
        '[2] Atkins, P.; de Paula, J.; Keeler, J. Atkins’ Physical Chemistry. 11-nashr. Oxford University Press, 2018. Kimyoviy muvozanat, faollik va elektrokimyo asoslari.',
        '[3] OpenStax. Chemistry 2e. Appendix H: Ionization Constants of Weak Acids. https://openstax.org/books/chemistry-2e/pages/h-ionization-constants-of-weak-acids . HSO₄⁻ uchun 25 °C dagi Kₐ₂=1,2×10⁻² qiymati.',
        'Barcha grafiklar ideal konsentratsion modellar; tajriba natijalari emas. Tuzilish tasvirlari original sxematik o‘quv modellari, optimallashtirilgan koordinatalar emas. K_c=10⁻⁵ shartli HA uchun, HF uchun emas.'
        ]
        for ref in refs:
            p=left.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3);runs(p,ref,9)

p=sec.header.paragraphs[0];runs(p,'NOORGANIK KISLOTALAR',8,False,NAVY)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runs(p,'3-BOB   ·   ',8,False,NAVY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Noorganik kislotalar · 3.3 · Dissotsiatsiya va pH'
doc.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
assert len(doc.inline_shapes)==13
assert counter==9
assert comparison_count==9
assert sum(len(notes) for notes in side_notes)==18
print(OUT)
print('9 sections; 9 tables; 5 schemes; 4 calculated graphs; 4 molecular renders; 9 equations; 18 margin notes.')
