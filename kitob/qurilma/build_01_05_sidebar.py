#!/usr/bin/env python3
"""A4: 3/4 asosiy matn + 1/4 vizual yon ustun. Alohida mavzu moduli."""
from pathlib import Path
import re, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'chiqish'/'Noorganik-kislotalar-1.5-Usanovich-Lyuks-Flud.docx'
ASSETS=ROOT/'rasmlar'/'usanovich_lux_sxema'; ASSETS.mkdir(parents=True,exist_ok=True)
NAVY='0B2C4A'; BRONZE='C28E3B'; PALE='EDF3F7'; INK='243746'
fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
def font(n):return ImageFont.truetype(fontpath,n)

def scheme(index,nodes):
    special={4:'silikat.png',5:'oksid_faolligi.png',6:'harorat.png'}
    if index in special:return ROOT/'rasmlar'/'usanovich_lux'/special[index]
    im=Image.new('RGB',(640,1100),'#F5F8FA');d=ImageDraw.Draw(im)
    y=65
    for i,(a,b) in enumerate(nodes):
        d.text((45,y),a,font=font(29),fill='#0B2C4A');y+=58
        # Hand-authored diagrams; not measured experimental data.
        lines=textwrap.wrap(b,width=25)
        for line in lines:
            d.text((45,y),line,font=font(30),fill='#243746');y+=43
        if i<len(nodes)-1:
            y+=25;d.line((320,y,320,y+65),fill='#C28E3B',width=5)
            d.polygon([(320,y+75),(309,y+56),(331,y+56)],fill='#C28E3B');y+=115
    path=ASSETS/f'01_01_{index:02d}.png';im.save(path);return path

panels=[
('Ikki xil qamrov',[('USANOVICH','Kation, anion va elektron mezoni'),('LYUKS–FLUD','Aynan oksid-ion almashinuvi'),('ASOSIY SAVOL','Qaysi zarracha berildi?')],'Keng tasnif va maxsus muhit mezonini farqlang.'),
('Na va Cl₂ misoli',[('Na · DONOR','Elektron beradi: qaytaruvchi'),('Cl₂ · AKSEPTOR','Elektron oladi: oksidlovchi'),('USANOVICH TILI','Asos va kislota vazifalari')],'Keng tasnif redoks balansini almashtirmaydi.'),
('Mos vositani tanlash',[('PROTON','Tutash juft va pKₐ'),('ELEKTRON','Redoks juft va potensial'),('KOMPLEKS','Donor–akseptor va hosil bo‘lish K')],'Bitta umumiy nom barcha miqdoriy usullar o‘rniga o‘tmaydi.'),
('Silikat tarmog‘i',[],'1.5.1-chizma. Oksid qo‘shilganda lokal Si–O–Si ko‘prigining ikki ko‘priksiz uchga o‘tish eskizi; to‘liq kristall tuzilishi emas.'),
('Oksid faolligi va tarkib',[],'1.5.1-grafik. A + O²⁻ ⇌ B ideal sharoiti; K = 100. Muayyan shlakning tajriba grafigi emas.'),
('Harorat ta’siri',[],'1.5.2-grafik. Faraziy ekzotermik muvozanat uchun lg K ning haroratga bog‘liqligi.'),
('Tahlil ketma-ketligi',[('01 · MUHIT','Suvmi yoki oksidli suyuqlanma?'),('02 · ZARRACHA','H⁺, O²⁻ yoki elektronmi?'),('03 · MAQSAD','Tasnif, tarkib yoki tezlikmi?')],'Mezonni tanlashdan oldin fizik-kimyoviy sharoitni belgilang.')
]

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.5)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.widow_control=True
for sty in doc.styles:
    if sty.type==1:
        sty.font.name='Times New Roman'
        rp=sty.element.get_or_add_rPr();rf=OxmlElement('w:rFonts')
        for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
        rp.append(rf)
sub=dict(zip('₀₁₂₃₄₅₆₇₈₉ₐᵦ','0123456789ab'));sup=dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻','0123456789+−'))
def runs(p,text,size=11,bold=False,color=INK):
    for token in re.findall(r'K_w|[₀₁₂₃₄₅₆₇₈₉ₐᵦ]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+|[^₀₁₂₃₄₅₆₇₈₉ₐᵦ⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻]+',text):
        if 'K_c' in token:
            before,after=token.split('K_c',1)
            if before:runs(p,before,size,bold,color)
            runs(p,'K',size,bold,color)
            r=p.add_run('c');r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.subscript=True;r.font.color.rgb=RGBColor.from_string(color)
            if after:runs(p,after,size,bold,color)
            continue
        if 'K_w' in token:
            before,after=token.split('K_w',1)
            if before:runs(p,before,size,bold,color)
            runs(p,'K',size,bold,color)
            r=p.add_run('w');r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.subscript=True;r.font.color.rgb=RGBColor.from_string(color)
            if after:runs(p,after,size,bold,color)
            continue
        lo=all(c in sub for c in token);hi=all(c in sup for c in token)
        t=''.join(sub[c] for c in token) if lo else ''.join(sup[c] for c in token) if hi else token
        r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
        if lo:r.font.subscript=True
        if hi:r.font.superscript=True

def shade(cell,fill):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(x)
def cell_margin(cell,**kwargs):
    x=OxmlElement('w:tcMar')
    for k,v in kwargs.items():
        y=OxmlElement('w:'+k);y.set(qn('w:w'),str(v));y.set(qn('w:type'),'dxa');x.append(y)
    cell._tc.get_or_add_tcPr().append(x)
def border(cell,edge,color,size):
    bs=OxmlElement('w:tcBorders');b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:color'),color);b.set(qn('w:sz'),str(size));bs.append(b);cell._tc.get_or_add_tcPr().append(bs)
def band(label,fill,color,size,accent=False):
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(18)
    c=t.cell(0,0);shade(c,fill);cell_margin(c,top=140,bottom=140,left=200,right=160)
    if accent:border(c,'left',BRONZE,24)
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    runs(p,label,size,True,color)
    # Prevent the banner row separating from its following paragraph.
    for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.keep_with_next=True

source=(ROOT/'matn'/'01_05.md').read_text()
sections=[]
for raw in source.split('## ')[1:]:
    title,body=raw.split('\n',1);sections.append((title,body.strip().split('\n\n')))
assert len(sections)==7
side_notes=json.loads((ROOT/'matn'/'01_05_yonlavhalar.json').read_text())
assert len(side_notes)==len(sections)
counter=0
for idx,((title,paras),(panel,nodes,note)) in enumerate(zip(sections,panels)):
    if idx:doc.add_page_break()
    if idx==0:
        band('1-BOB  /  KISLOTA TUSHUNCHASI VA NAZARIYALAR',NAVY,'FFFFFF',11)
        band('1.5. Usanovich va Lyuks–Flud nazariyalari',PALE,NAVY,18,True)
    else:
        band('1.5  /  USANOVICH VA LYUKS–FLUD NAZARIYALARI',PALE,NAVY,11,True)
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(9);p.paragraph_format.keep_with_next=True
    runs(p,title,14,True,NAVY)
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    # 18 cm printable width: 13.5 cm narrative + 4.5 cm outer visual strip.
    table.columns[0].width=Cm(13.5);table.columns[1].width=Cm(4.5)
    left,right=table.rows[0].cells
    left.width=Cm(13.5);right.width=Cm(4.5)
    for c in [left,right]:c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell_margin(left,top=0,bottom=0,left=0,right=220)
    cell_margin(right,top=140,bottom=130,left=160,right=100)
    shade(right,'F5F8FA');border(right,'left','D2DCE5',6)
    # Allow long content to flow rather than shrink type or clip a fixed-height row.
    for j,text in enumerate(paras):
        p=left.paragraphs[0] if j==0 else left.add_paragraph()
        if re.search(r' [→⇌] ',text) or text.startswith(('K_c = ', 'θ = ')):
            counter+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8)
            runs(p,text+'   (1.5.'+str(counter)+')',11)
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent=Cm(.45 if j else 0)
            runs(p,text)
    p=right.paragraphs[0];p.paragraph_format.space_after=Pt(9)
    runs(p,panel,10,True,NAVY)
    p=right.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(scheme(idx+1,nodes)),width=Cm(3.92))
    p=right.add_paragraph();p.paragraph_format.space_after=Pt(10)
    runs(p,'HISOBIY GRAFIK' if idx in [4,5] else 'KIMYOVIY CHIZMA' if idx==3 else f'1.5.{idx+1}-sxema',9,True,NAVY)
    p=right.add_paragraph();runs(p,note,9)
    for label,body in side_notes[idx]:
        p=right.add_paragraph()
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        # Delicate bronze rule groups each short aside without heavy boxes.
        pp=p._p.get_or_add_pPr();bd=OxmlElement('w:pBdr')
        rule=OxmlElement('w:top');rule.set(qn('w:val'),'single')
        rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),BRONZE)
        rule.set(qn('w:space'),'5');bd.append(rule);pp.append(bd)
        runs(p,label,9,True,NAVY)
        p=right.add_paragraph()
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        runs(p,body,9)
    if idx==6:
        p=left.add_paragraph();p.paragraph_format.space_before=Pt(8);runs(p,'Adabiyotlar',10,True,NAVY)
        refs=[
        '[1] Usanovich and Nernst colliding: inconsistencies in the all-in-one acid–base concept? Foundations of Chemistry (2023). DOI: 10.1007/s10698-023-09482-x.',
        '[2] Chemistry LibreTexts. 2.3: Acid-Base Theories and Concepts. Northern Michigan University, CH 215, 2023. Bo‘limlar: Lux–Flood va Usanovich.',
        '[3] Acid–Base Properties of Oxides Derived from Oxide Melt Solution Calorimetry (2023). Ochiq manba: PMC10304485.',
        'Qo‘shimcha o‘qish: Housecroft, C. E.; Sharpe, A. G. Inorganic Chemistry. 5-nashr. Pearson, 2018.',
        'Grafik va chizmalar shu mavzu uchun mualliflik hisobiy shartli hisoblari sifatida tayyorlangan. Tajriba ma’lumotlari emas; farazlar matn va yon lavhalarda ko‘rsatilgan.'
        ]
        for ref in refs:
            p=left.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3);runs(p,ref,9)

p=sec.header.paragraphs[0];runs(p,'NOORGANIK KISLOTALAR',8,False,NAVY)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runs(p,'1-BOB   ·   ',8,False,NAVY)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Noorganik kislotalar · 1.5 · Usanovich va Lyuks–Flud nazariyalari'
doc.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
assert len(doc.inline_shapes)==7
assert counter>=7
assert sum(len(notes) for notes in side_notes)==21
print(OUT)
print('7 modular sections; 7 original side diagrams; editable equations; width ratio 13.5:4.5 cm.')
